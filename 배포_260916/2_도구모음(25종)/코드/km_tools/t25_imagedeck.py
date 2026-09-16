# -*- coding: utf-8 -*-
"""25. 이미지 모음 PPT - 사진 폴더를 참고자료 PPT 한 개로 묶는다 (16:9).
저작권 때문에 슬라이드마다 출처를 찍는다. 파일명에 @출처 를 적어두면 자동으로 들어간다.
  예) 5성급_침대옆_BSP@우리시공_쏠비치양양.jpg
      로비_조명제어@Unsplash.jpg
출처가 없으면 슬라이드에 빨간 글씨로 '출처 미기재'가 남는다 - 대외 배포 전에 반드시 채우십시오."""
import os, datetime
from common import *

EXTS = {'.jpg', '.jpeg', '.png', '.webp'}

def parse(name):
    """파일명 -> (설명, 출처)"""
    base = os.path.splitext(os.path.basename(name))[0]
    if '@' in base:
        desc, src = base.split('@', 1)
    else:
        desc, src = base, ''
    return desc.replace('_', ' ').strip(), src.replace('_', ' ').strip()

def prep(path, tmp, maxw=1600):
    try:
        from PIL import Image, ImageOps
        im = Image.open(path)
        im = ImageOps.exif_transpose(im)
        if im.width > maxw:
            im = im.resize((maxw, int(im.height * maxw / im.width)))
        if im.mode != 'RGB':
            im = im.convert('RGB')
        out = os.path.join(tmp, safe_name(os.path.basename(path)) + '.jpg')
        im.save(out, quality=86)
        return out, im.width, im.height
    except Exception:
        return path, 0, 0

def build(src, deck_title, per_page=2, note=''):
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
        from pptx.dml.color import RGBColor
    except ImportError:
        print('python-pptx 가 없습니다. _처음_한번만_설치.bat 을 다시 눌러주십시오.')
        print('(또는 명령창에서 : pip install python-pptx)')
        return None
    pics = sorted(p for p in walk_files(src, EXTS))
    if not pics:
        print('이미지가 없습니다 : %s' % src); return None
    od = outdir('이미지모음')
    tmp = os.path.join(od, '_tmp'); os.makedirs(tmp, exist_ok=True)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    BLANK = prs.slide_layouts[6]

    s = prs.slides.add_slide(BLANK)
    tb = s.shapes.add_textbox(Inches(0.8), Inches(2.6), Inches(11.7), Inches(2.0)).text_frame
    tb.text = deck_title
    tb.paragraphs[0].runs[0].font.size = Pt(40)
    tb.paragraphs[0].runs[0].font.bold = True
    tb.paragraphs[0].runs[0].font.name = '맑은 고딕'
    p = tb.add_paragraph()
    p.text = '한국마이크로닉(주)   %s   사진 %d컷' % (today().isoformat(), len(pics))
    p.runs[0].font.size = Pt(14); p.runs[0].font.name = '맑은 고딕'
    p.runs[0].font.color.rgb = RGBColor(0x8E, 0x99, 0xA4)
    if note:
        p2 = tb.add_paragraph(); p2.text = note
        p2.runs[0].font.size = Pt(12); p2.runs[0].font.name = '맑은 고딕'

    box = {1: [(0.7, 0.9, 11.9, 5.4)],
           2: [(0.5, 1.0, 6.0, 5.0), (6.8, 1.0, 6.0, 5.0)],
           4: [(0.5, 0.8, 6.0, 2.9), (6.8, 0.8, 6.0, 2.9),
               (0.5, 4.0, 6.0, 2.9), (6.8, 4.0, 6.0, 2.9)]}[per_page]
    rows = []
    for i in range(0, len(pics), per_page):
        s = prs.slides.add_slide(BLANK)
        hd = s.shapes.add_textbox(Inches(0.5), Inches(0.18), Inches(12.3), Inches(0.5)).text_frame
        hd.text = '%s  (%d/%d)' % (deck_title, i // per_page + 1, (len(pics) - 1) // per_page + 1)
        hd.paragraphs[0].runs[0].font.size = Pt(13)
        hd.paragraphs[0].runs[0].font.name = '맑은 고딕'
        hd.paragraphs[0].runs[0].font.color.rgb = RGBColor(0x8E, 0x99, 0xA4)
        for j, path in enumerate(pics[i:i + per_page]):
            L, T, W, H = box[j]
            img, w, h = prep(path, tmp)
            try:
                if w and h and (w / h) > (W / H):
                    s.shapes.add_picture(img, Inches(L), Inches(T + (H - W * h / w) / 2), width=Inches(W))
                else:
                    s.shapes.add_picture(img, Inches(L), Inches(T), height=Inches(H))
            except Exception:
                pass
            desc, source = parse(path)
            cf = s.shapes.add_textbox(Inches(L), Inches(T + H + 0.05), Inches(W), Inches(0.6)).text_frame
            cf.word_wrap = True
            cf.text = desc[:70]
            cf.paragraphs[0].runs[0].font.size = Pt(12)
            cf.paragraphs[0].runs[0].font.bold = True
            cf.paragraphs[0].runs[0].font.name = '맑은 고딕'
            q = cf.add_paragraph()
            q.text = ('출처 : %s' % source) if source else '출처 미기재 - 대외 배포 전 확인 필요'
            q.runs[0].font.size = Pt(9); q.runs[0].font.name = '맑은 고딕'
            q.runs[0].font.color.rgb = RGBColor(0x8E, 0x99, 0xA4) if source else RGBColor(0xC0, 0x39, 0x2B)
            rows.append([os.path.basename(path), desc, source or '[미기재]',
                         datetime.date.fromtimestamp(os.path.getmtime(path)).isoformat()])
    out = os.path.join(od, '%s_%s.pptx' % (safe_name(deck_title), ymd6()))
    prs.save(out)
    csvp = os.path.join(od, '%s_출처대장_%s.csv' % (safe_name(deck_title), ymd6()))
    write_csv(csvp, rows, ['파일명', '설명', '출처', '파일날짜'])
    miss = sum(1 for r in rows if r[2] == '[미기재]')
    return out, csvp, len(pics), miss

def run():
    title('25. 이미지 모음 PPT')
    print('파일명에 @출처 를 적어두시면 슬라이드에 자동으로 들어갑니다.')
    print('  예) 5성급_침대옆_BSP@우리시공_쏠비치양양.jpg')
    print('')
    src = ask('이미지가 든 폴더 > ')
    if not os.path.isdir(src):
        print('폴더를 찾지 못했습니다.'); return
    t = ask('PPT 제목 > ', '5성급 호텔 객실 참고자료')
    per = int(ask('한 장에 몇 컷? (1/2/4, 기본 2) > ', '2') or 2)
    per = per if per in (1, 2, 4) else 2
    note = ask('표지에 넣을 한 줄 (엔터=없음) > ')
    r = build(src, t, per, note)
    if not r:
        return
    out, csvp, n, miss = r
    print('')
    print('사진 %s컷 -> %s' % (won(n), out))
    print('출처대장 : %s' % csvp)
    if miss:
        print('')
        print('[주의] 출처 미기재 %s컷. 대외 배포 전에 출처대장을 채우십시오.' % won(miss))
        print('       블로그에서 가져온 사진은 그대로 제안서에 넣으면 저작권 문제가 됩니다.')
    log('이미지모음', '%d컷 미기재 %d' % (n, miss))

if __name__ == '__main__':
    run(); pause()
