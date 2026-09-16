# -*- coding: utf-8 -*-
"""7. 단가장 검진 (경비원) - 요율 하드코딩, 검산 줄, 동명이인, 임시단가를 훑는다.
요율은 입력칸을 참조해야 하고, Claude 도 사람도 하드코딩하면 안 된다 (4.4)."""
import os, re, sys
from common import *

HARD = re.compile(r'\*\s*(1\.5|1\.8|2\.1|1\.6|0\.3)\b')
RATE_CELL = re.compile(r'\$[A-Z]{1,2}\$\d+')

def check(path):
    try:
        import openpyxl
    except ImportError:
        print('openpyxl 이 필요합니다 : pip install openpyxl'); return None
    wf = openpyxl.load_workbook(path, data_only=False)
    wv = openpyxl.load_workbook(path, data_only=True)
    hard, zero, dup, temp = [], [], [], []
    for ws in wf.worksheets:
        vs = wv[ws.title]
        names = {}
        for row in ws.iter_rows():
            for c in row:
                f = c.value
                if isinstance(f, str) and f.startswith('='):
                    if HARD.search(f) and not RATE_CELL.search(f):
                        hard.append([ws.title, c.coordinate, f[:70]])
                a = row[0].value
                if isinstance(a, str) and a.strip():
                    key = a.strip()
                    if c.column == 1:
                        names.setdefault(key, []).append(c.row)
            a0 = row[0].value
            if isinstance(a0, str) and '검산' in a0:
                v = vs.cell(row=row[0].row, column=row[0].column + 1).value
                for col in range(2, 12):
                    v2 = vs.cell(row=row[0].row, column=col).value
                    if isinstance(v2, (int, float)):
                        v = v2; break
                if isinstance(v, (int, float)) and abs(v) > 0.5:
                    zero.append([ws.title, row[0].coordinate, a0[:40], v])
        for k, rr in names.items():
            if len(rr) > 1 and len(k) > 3 and not k.startswith(('합계', '소계', '검산', '품목')):
                dup.append([ws.title, k[:40], ', '.join(map(str, rr[:8]))])
        for row in ws.iter_rows():
            for c in row:
                try:
                    fill = c.fill.start_color.rgb if c.fill and c.fill.start_color else None
                except Exception:
                    fill = None
                if fill and str(fill).upper() in ('FFFFFF00', '00FFFF00', 'FFFFE699'):
                    v = vs[c.coordinate].value
                    if isinstance(v, (int, float)) and v:
                        temp.append([ws.title, c.coordinate, v])
    return hard, zero, dup, temp

def run(path=None):
    title('7. 단가장 검진 (경비원)')
    path = path or ask('단가장 xlsx 경로 > ')
    if not os.path.isfile(path):
        print('파일을 찾지 못했습니다.'); return
    r = check(path)
    if r is None:
        return
    hard, zero, dup, temp = r
    def show(name, rows, cols, limit=15):
        print('')
        print('[%s] %s건' % (name, won(len(rows))))
        for x in rows[:limit]:
            print('   ' + ' | '.join(str(i)[:46] for i in x))
        if len(rows) > limit:
            print('   ... 외 %s건' % won(len(rows) - limit))
    show('요율 하드코딩 (입력칸 참조로 바꿔야 함)', hard, 3)
    show('검산이 0 이 아님', zero, 4)
    show('같은 이름이 여러 행 (동명이인)', dup, 3)
    show('노란칸 임시단가 (확정되면 교체)', temp, 3)
    od = outdir('단가장검진')
    for nm, rows, head in (('요율하드코딩', hard, ['시트', '칸', '수식']),
                           ('검산불일치', zero, ['시트', '칸', '라벨', '값']),
                           ('동명이인', dup, ['시트', '품목', '행']),
                           ('임시단가', temp, ['시트', '칸', '값'])):
        write_csv(os.path.join(od, '%s_%s.csv' % (nm, ymd6())), rows, head)
    log('단가장검진', '하드%d 검산%d 중복%d 임시%d' % (len(hard), len(zero), len(dup), len(temp)))
    print('')
    print('저장 폴더 : %s' % od)
    if not (hard or zero):
        print('요율/검산은 이상 없습니다.')

if __name__ == '__main__':
    a = sys.argv[1:]
    run(a[0] if a else None); pause()
