# -*- coding: utf-8 -*-
"""14. 사진대지 - 현장 사진 폴더를 A4 세로 엑셀 사진대지로. 촬영일시는 EXIF 에서 자동.
회사 원틀 3종 중 B(2장/쪽) 기본, 4장/쪽도 가능. 설치완료 검수/수금 근거용."""
import os, datetime
from common import *

def exif_dt(path):
    try:
        from PIL import Image
        im = Image.open(path)
        ex = im._getexif() or {}
        for tag in (36867, 36868, 306):
            v = ex.get(tag)
            if v:
                return str(v)[:19].replace(':', '-', 2)
    except Exception:
        pass
    return datetime.date.fromtimestamp(os.path.getmtime(path)).isoformat()

def prep(path, tmpdir, maxw=900):
    """회전 보정 + 축소. PIL 없으면 원본 그대로."""
    try:
        from PIL import Image, ImageOps
        im = Image.open(path)
        im = ImageOps.exif_transpose(im)
        if im.width > maxw:
            im = im.resize((maxw, int(im.height * maxw / im.width)))
        if im.mode != 'RGB':
            im = im.convert('RGB')
        out = os.path.join(tmpdir, safe_name(os.path.basename(path)) + '.jpg')
        im.save(out, quality=82)
        return out
    except Exception:
        return path

def build(src, site, gongjong, per_page=2):
    try:
        import openpyxl
        from openpyxl.drawing.image import Image as XLImage
        from openpyxl.styles import Font, Alignment, Border, Side
    except ImportError:
        print('openpyxl 이 필요합니다. 명령창에 : pip install openpyxl pillow')
        return None
    pics = sorted([p for p in walk_files(src, {'.jpg', '.jpeg', '.png'})])
    if not pics:
        print('사진이 없습니다 : %s' % src); return None
    od = outdir('사진대지')
    tmp = os.path.join(od, '_tmp'); os.makedirs(tmp, exist_ok=True)
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = '사진대지'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = 'portrait'
    ws.print_area = 'A1:F%d' % (len(pics) * 14 + 6)
    thin = Side(style='thin', color='000000')
    bd = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws['A1'] = '사 진 대 지'
    ws['A1'].font = Font(name='맑은 고딕', size=18, bold=True)
    ws.merge_cells('A1:F1'); ws['A1'].alignment = Alignment(horizontal='center')
    ws['A2'] = '공 사 명'; ws['B2'] = site
    ws['A3'] = '공    종'; ws['B3'] = gongjong
    for c in ('A2', 'A3'):
        ws[c].font = Font(name='맑은 고딕', size=10, bold=True)
    for c in ('B2', 'B3'):
        ws[c].font = Font(name='맑은 고딕', size=10)
    rows_per = 15 if per_page == 2 else 11
    r = 5
    for i, p in enumerate(pics):
        img = prep(p, tmp)
        try:
            xi = XLImage(img); xi.width, xi.height = (430, 300) if per_page == 2 else (300, 210)
            ws.add_image(xi, 'A%d' % r)
        except Exception as e:
            ws['A%d' % r] = '[사진 삽입 실패] %s' % os.path.basename(p)
        lab = r + (rows_per - 2)
        ws['A%d' % lab] = '촬영일'; ws['B%d' % lab] = exif_dt(p)
        ws['C%d' % lab] = '내용';  ws['D%d' % lab] = os.path.splitext(os.path.basename(p))[0][:40]
        for col in 'ABCD':
            ws['%s%d' % (col, lab)].font = Font(name='맑은 고딕', size=9)
            ws['%s%d' % (col, lab)].border = bd
        r += rows_per
        if (i + 1) % per_page == 0:
            ws.row_breaks.append(openpyxl.worksheet.pagebreak.Break(id=r - 1))
    for col, w in zip('ABCDEF', (14, 22, 10, 30, 12, 12)):
        ws.column_dimensions[col].width = w
    out = os.path.join(od, '%s_사진대지_%s.xlsx' % (safe_name(site), ymd6()))
    wb.save(out)
    return out, len(pics)

def run():
    title('14. 사진대지 만들기')
    src = ask('사진이 든 폴더 경로 > ')
    if not os.path.isdir(src):
        print('폴더를 찾지 못했습니다.'); return
    site = ask('공사명(현장) > ', '현장미정')
    gj = ask('공종 (기본: 객실관리설비공사) > ', '객실관리설비공사')
    per = int(ask('한 쪽에 몇 장? (2 또는 4, 기본 2) > ', '2') or 2)
    res = build(src, site, gj, per)
    if res:
        print('사진 %s장 -> %s' % (won(res[1]), res[0]))

if __name__ == '__main__':
    run(); pause()
