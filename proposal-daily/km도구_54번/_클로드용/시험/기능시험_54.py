import sys, os, io, json, copy, glob, importlib
sys.path.insert(0, os.getcwd())
OUT = sys.argv[1]
from pptx import Presentation
from pptx.util import Inches
import deck_core as deck, t54_deck as t54, deck_check as dc
ok = lambda c, m: print(('통과 ' if c else '실패 ') + m)

# B. 옛 Pillow : getlength·getbbox 없음
from PIL import ImageFont
F = type(ImageFont.truetype('/usr/share/fonts/truetype/nanum/NanumGothic.ttf', 10))
saved = {k: getattr(F, k) for k in ('getlength', 'getbbox')}
for k in saved: delattr(F, k)
dc._cache.clear()
try:
    spec = json.load(io.open('제안서_내용/수1_FIP_사용매뉴얼.json', encoding='utf-8'))
    spec = deck.fill(spec, {'{{현장}}': '청담PJ', '{{현장_제목}}': '청담PJ ', '{{날짜}}': 'x'})
    f = os.path.join(OUT, 'oldpillow.pptx'); deck.build(spec, f)
    ok(os.path.exists(f), '옛 Pillow(getlength 없음)에서도 제안서가 만들어짐')
finally:
    for k, v in saved.items(): setattr(F, k, v)
    dc._cache.clear()

# C+D. 목차 쪽번호 (4개 목차 전부, 물량 장 넣은 것 · 안 넣은 것)
rows = [['품목%d' % i, str(10 + i), '블록'] for i in range(12)]
extras0 = t54.qty_slides('청담PJ', rows, 'x.csv')
ok(len(extras0) == 3 and len(extras0[0]['rows']) == 9 and len(extras0[1]['rows']) == 3, '12품목 -> 표 2장(9+3) + 적용범위 1장, 빠진 품목 없음')
def page_titles(f):
    out = {}
    for s in Presentation(f).slides:
        num = None; ttl = None
        for sh in s.shapes:
            if sh.has_text_frame:
                t = sh.text_frame.text.strip()
                if abs(sh.top / 914400 - 5.22) < 0.01 and t.isdigit(): num = int(t)
                if abs(sh.top / 914400 - 0.69) < 0.02 or abs(sh.top / 914400 - 0.45) < 0.02:
                    if len(t) > 1 and ttl is None and sh.text_frame.paragraphs[0].runs and sh.text_frame.paragraphs[0].runs[0].font.size and sh.text_frame.paragraphs[0].runs[0].font.size.pt >= 13: ttl = t
        if num: out[num] = ttl
    return out
allgood = True
for name in ('금4_파트너공동제안_도어락', '수1_FIP_사용매뉴얼', '수2_방재실_운영PC매뉴얼', '수4_준공인계_매뉴얼'):
    for with_qty in (False, True):
        spec = json.load(io.open('제안서_내용/%s.json' % name, encoding='utf-8'))
        d2 = deck.fill(spec, {'{{현장}}': '청담PJ', '{{현장_제목}}': '청담PJ ', '{{날짜}}': 'x'})
        sl = d2['slides']
        if with_qty:
            at = 1
            if sl[at].get('type') == 'toc': at += 1
            ex = t54.qty_slides('청담PJ', rows, 'x.csv')
            for x in sl:
                if x.get('type') == 'toc':
                    x['items'] = [{'text': '그 현장 물량 · 적용 범위', 'goto': ex[0]['title']}] + x['items']
            for k, e in enumerate(ex): sl.insert(at + k, e)
        f = os.path.join(OUT, 'toc_%s_%d.pptx' % (name[:2], with_qty)); t54.build_merged(d2, None, [], f, {})
        pt = page_titles(f)
        toc = next(x for x in sl if x['type'] == 'toc')
        for it in toc['items']:
            got = pt.get(it.get('page'))
            good = bool(got) and (got == it['goto'] or it['goto'] in (got or ''))
            allgood &= good
            if not good: print('   어긋남 %s %s : 「%s」 -> %s쪽 = %s' % (name[:2], with_qty, it['text'], it.get('page'), got))
ok(allgood, '목차 4개 x (물량 장 없음/있음) : 모든 항목의 쪽번호가 그 제목의 장을 가리킴')

# E. 회사 원본 장은 「요청」 장 앞에
src = Presentation(); s1 = src.slides.add_slide(src.slide_layouts[6])
tb = s1.shapes.add_textbox(0, 0, Inches(3), Inches(1)); tb.text = '회사개요 COMPANY'
src.slide_width, src.slide_height = Inches(10), Inches(5.625); src.save(os.path.join(OUT, 'src.pptx'))
spec = json.load(io.open('제안서_내용/목5_에어컨EHP_연동설명.json', encoding='utf-8'))
spec['slides'].append({'type': 'request', 'title': '요청드리는 사항', 'items': [{'text': 'x'}]})
f = os.path.join(OUT, 'order.pptx'); t54.build_merged(spec, os.path.join(OUT, 'src.pptx'), [1], f, {})
texts = [' '.join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame) for s in Presentation(f).slides]
ok('회사개요' in texts[-2] and '요청' in texts[-1], '회사 원본 장이 「요청드리는 사항」 바로 앞에 들어감 (%d장)' % len(texts))

# G. 점유율 숫자가 글자 조각 둘로 나뉜 경우
import errata_apply as ea
p = Presentation(); s = p.slides.add_slide(p.slide_layouts[6])
tf = s.shapes.add_textbox(0, 0, Inches(8), Inches(1)).text_frame
para = tf.paragraphs[0]; r1 = para.add_run(); r1.text = '국내 5성급 호텔 7'; r2 = para.add_run(); r2.text = '0% 이상 점유'
tf2 = s.shapes.add_textbox(0, Inches(2), Inches(8), Inches(1)).text_frame; tf2.text = '5성급 호텔 75% 이상 점유, 170% 증가'
p.save(os.path.join(OUT, 'err.pptx'))
dst, log = ea.fix(os.path.join(OUT, 'err.pptx'))
t = [sh.text_frame.text for sh in Presentation(dst).slides[0].shapes]
ok(any('못 바꿈' in x[2] for x in log), '조각난 숫자는 못 바꿨다고 기록에 남김')
ok('80%' in t[1] and '170%' in t[1], '75%% -> 80%%, 170%% 는 그대로 : %s' % t[1])
dst2, _ = ea.fix(os.path.join(OUT, 'err.pptx'))
ok(dst2 != dst, '두 번째 고치기는 새 이름 : %s' % os.path.basename(dst2))
