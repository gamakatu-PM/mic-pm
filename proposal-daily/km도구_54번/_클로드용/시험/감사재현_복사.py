import sys, os, zipfile, io, warnings
sys.path.insert(0, '/home/user/mic-pm/proposal-daily/km도구_54번/코드/km_tools')
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from pptx import Presentation
from pptx.util import Inches
from PIL import Image
import t54_deck
# make two images
for n,c in (('a.png','red'),('b.png','blue')):
    Image.new('RGB',(20,20),c).save(n)
src=Presentation()
s1=src.slides.add_slide(src.slide_layouts[6])
s2=src.slides.add_slide(src.slide_layouts[6])
p1=s1.shapes.add_picture('a.png',0,0,Inches(1))
p2=s1.shapes.add_picture('b.png',Inches(2),0,Inches(1))
# hyperlink to slide 2
tb=s1.shapes.add_textbox(0,Inches(2),Inches(2),Inches(1)); tb.text='목차'
tb.click_action.target_slide=s2
src.save('src.pptx')
# rewrite slide1 rels so layout is last: renumber rIds
z=zipfile.ZipFile('src.pptx'); files={n:z.read(n) for n in z.namelist()}; z.close()
rels=files['ppt/slides/_rels/slide1.xml.rels'].decode(); xml=files['ppt/slides/slide1.xml'].decode()
print(rels)
m={'rId1':'X9','rId2':'X1','rId3':'X2'}
import re
def ren(s):
    return re.sub(r'"(rId[123])"', lambda mm: '"%s"'%m[mm.group(1)], s).replace('"X','"rId')
files['ppt/slides/_rels/slide1.xml.rels']=ren(rels).encode(); files['ppt/slides/slide1.xml']=ren(xml).encode()
with zipfile.ZipFile('src2.pptx','w',zipfile.ZIP_DEFLATED) as zo:
    for n,d in files.items(): zo.writestr(n,d)
src=Presentation('src2.pptx')
print([ (r.rId, r.reltype.split('/')[-1]) for r in src.slides[0].part.rels.values()])
dst=Presentation(); dst.slide_width=Inches(10); dst.slide_height=Inches(5.625)
d=t54_deck.copy_slide(src.slides[0], dst)
from pptx.enum.shapes import MSO_SHAPE_TYPE
pics=[sh for sh in d.shapes if sh.shape_type==MSO_SHAPE_TYPE.PICTURE]
print('pic embeds', [sh._element.blipFill.blip.rEmbed for sh in pics], [sh.image.blob[:0] or sh.image.filename for sh in pics])
print('colors', [Image.open(io.BytesIO(sh.image.blob)).getpixel((1,1)) for sh in pics])
with warnings.catch_warnings(record=True) as w:
    warnings.simplefilter('always')
    dst.save('out.pptx')
    print('warnings', [str(x.message) for x in w])
names=zipfile.ZipFile('out.pptx').namelist()
import collections
print('dups', [n for n,c in collections.Counter(names).items() if c>1])
print([n for n in names if 'slide' in n and 'Layout' not in n])
