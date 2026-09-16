# -*- coding: utf-8 -*-
"""만들어진 pptx 의 실제 도형 좌표·글자를 읽어 HTML 로 다시 그리고 PNG 로 찍는다.
LibreOffice 가 없는 환경에서 렌더 검수를 대신한다."""
import sys, os, html
from pptx import Presentation
from pptx.util import Emu

PX = 96.0 / 914400.0  # EMU -> px


def esc(t):
    return html.escape(t).replace('\n', '<br>')


def slide_html(slide, w, h):
    out = ['<div class="s" style="width:%dpx;height:%dpx">' % (w * PX, h * PX)]
    for sh in slide.shapes:
        x, y = sh.left * PX, sh.top * PX
        cw, ch = sh.width * PX, sh.height * PX
        if sh.has_table:
            tbl = sh.table
            out.append('<table style="left:%.1fpx;top:%.1fpx;width:%.1fpx">' % (x, y, cw))
            for r in tbl.rows:
                out.append('<tr>')
                for c in r.cells:
                    bg = ''
                    try:
                        bg = 'background:#%s;' % (str(c.fill.fore_color.rgb),)
                    except Exception:
                        pass
                    col = '#333'
                    fs = 8.5
                    bold = ''
                    for p in c.text_frame.paragraphs:
                        for run in p.runs:
                            try:
                                col = '#%s' % (str(run.font.color.rgb),)
                            except Exception:
                                pass
                            if run.font.size:
                                fs = run.font.size.pt
                            if run.font.bold:
                                bold = 'font-weight:700;'
                    out.append('<td style="%scolor:%s;font-size:%.1fpx;%s">%s</td>'
                               % (bg, col, fs * 96 / 72.0, bold, esc(c.text)))
                out.append('</tr>')
            out.append('</table>')
            continue
        fill = ''
        try:
            if sh.fill.type is not None and sh.fill.type == 1:
                fill = 'background:#%s;' % (str(sh.fill.fore_color.rgb),)
        except Exception:
            pass
        border = ''
        try:
            if sh.line.color and sh.line.color.rgb is not None:
                border = 'border:1px solid #%s;' % (str(sh.line.color.rgb),)
        except Exception:
            pass
        radius = 'border-radius:6px;' if 'ROUNDED' in str(sh.shape_type) else ''
        if str(getattr(sh, 'shape_type', '')).find('OVAL') >= 0:
            radius = 'border-radius:50%;'
        inner = ''
        if sh.has_text_frame and sh.text_frame.text.strip():
            anchor = 'flex-start'
            if str(sh.text_frame.vertical_anchor).find('MIDDLE') >= 0:
                anchor = 'center'
            parts = []
            for p in sh.text_frame.paragraphs:
                al = 'left'
                if str(p.alignment).find('CENTER') >= 0:
                    al = 'center'
                elif str(p.alignment).find('RIGHT') >= 0:
                    al = 'right'
                for run in p.runs:
                    col = '#333'
                    try:
                        col = '#%s' % (str(run.font.color.rgb),)
                    except Exception:
                        pass
                    fs = run.font.size.pt if run.font.size else 11
                    bw = '700' if run.font.bold else '400'
                    parts.append('<div style="text-align:%s;color:%s;font-size:%.1fpx;'
                                 'font-weight:%s;line-height:1.25">%s</div>'
                                 % (al, col, fs * 96 / 72.0, bw, esc(run.text) or '&nbsp;'))
            inner = ('<div style="display:flex;flex-direction:column;justify-content:%s;'
                     'height:100%%;overflow:visible">%s</div>' % (anchor, ''.join(parts)))
        out.append('<div class="b" style="left:%.1fpx;top:%.1fpx;width:%.1fpx;height:%.1fpx;%s%s%s">%s</div>'
                   % (x, y, cw, ch, fill, border, radius, inner))
    out.append('</div>')
    return '\n'.join(out)


def main(path, outdir):
    prs = Presentation(path)
    w, h = prs.slide_width, prs.slide_height
    css = ('<style>body{margin:0;background:#888;font-family:"Noto Sans CJK KR","Malgun Gothic",sans-serif}'
           '.s{position:relative;background:#fff;margin:0 0 14px 0;overflow:hidden}'
           '.b{position:absolute;box-sizing:border-box;padding:1px}'
           'table{position:absolute;border-collapse:collapse;table-layout:fixed}'
           'td{border:1px solid #d9d9d9;padding:2px 4px;vertical-align:middle;'
           'word-break:break-all;overflow:hidden}</style>')
    pages = [slide_html(s, w, h) for s in prs.slides]
    doc = '<!doctype html><meta charset="utf-8">' + css + ''.join(pages)
    hp = os.path.join(outdir, 'preview.html')
    open(hp, 'w', encoding='utf-8').write(doc)
    from playwright.sync_api import sync_playwright
    with sync_playwright() as p:
        b = p.chromium.launch(executable_path='/opt/pw-browsers/chromium-1194/chrome-linux/chrome', args=['--no-sandbox'])
        pg = b.new_page(viewport={'width': int(w * PX) + 20, 'height': int(h * PX) + 20})
        pg.goto('file://' + os.path.abspath(hp))
        for i, el in enumerate(pg.query_selector_all('.s'), 1):
            el.screenshot(path=os.path.join(outdir, 'slide%d.png' % i))
        b.close()
    print('slides:', len(pages))


if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else '.')
