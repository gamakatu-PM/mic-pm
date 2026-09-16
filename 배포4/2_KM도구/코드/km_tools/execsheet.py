# -*- coding: utf-8 -*-
"""실행산출 6시트 만들기 - 정답본(주일능 v7) 틀을 복사해 값만 채운다. 사람 손이 안 들어가므로 매번 같은 모양.
28번이 부르고, 클로드도 이 함수로만 만든다.

시트 : 0.자가진단 / 1.입력판 / 2.CB 실행 / 3.견적↔실행 대조 / 4.확인 목록 / 5.변경 이력
"""
import os, re, copy, shutil, collections, glob
from common import *

SHEETS = ['0.자가진단', '1.입력판', '2.CB 실행', '3.견적↔실행 대조', '4.확인 목록', '5.변경 이력']
YEL = 'FFF2CC'; GRY = 'E8EAED'; GRN = 'D6EFD8'

def template():
    """정답본 폴더의 실행산출 정답본 (없으면 None -> 같은 구조로 새로 그림)"""
    d = os.path.join(cfg('template'), '정답본')
    for p in sorted(glob.glob(os.path.join(d, '*.xlsx')), key=os.path.getmtime, reverse=True):
        b = os.path.basename(p)
        if b.startswith('~$'):
            continue
        try:
            import openpyxl
            wb = openpyxl.load_workbook(p, read_only=True)
            ok = all(any(re.sub(r'\s', '', s) == re.sub(r'\s', '', n) for s in wb.sheetnames) for n in SHEETS)
            wb.close()
            if ok:
                return p
        except Exception:
            continue
    return None

def _fill(c):
    from openpyxl.styles import PatternFill
    return PatternFill('solid', fgColor=c)

def build(site, out_path, qty, cb_qty, cb_rows, mult, unknown, meta):
    """
    qty      : [(품명, 규격, 단위, 수량, 실행단가 or None, 비고, 구역)]  구역 = '중앙'|'객실'|'노무'|'도어락'
    cb_qty   : [(타입이름, 대수)]  예 [('CB1 스텐다드-1~4',270), ...]
    cb_rows  : [(모듈명, 도면표기, [1대당 수량 per 타입], 실행단가 or None, 비고)]
    mult     : {'조립비율':0.3,'계약배수':1.5,'견적배수':1.8,'예산배수':2.1}
    unknown  : [(무엇을, 왜, 후보/참고)]  -> 1.입력판 노란칸 + 4.확인 목록
    meta     : {'도면':..., '견적서':..., '단가장':..., '출처':...}
    """
    import openpyxl
    from openpyxl.utils import get_column_letter as L
    tpl = template()
    if tpl:
        shutil.copy(tpl, out_path); wb = openpyxl.load_workbook(out_path)
        # 시트 이름을 표준으로
        for ws in wb.worksheets:
            for n in SHEETS:
                if re.sub(r'\s', '', ws.title) == re.sub(r'\s', '', n) and ws.title != n:
                    ws.title = n
        styles = {n: _snap(wb[n]) for n in SHEETS}
        for n in SHEETS:
            _clear(wb[n])
    else:
        wb = openpyxl.Workbook(); wb.remove(wb.active)
        for n in SHEETS: wb.create_sheet(n)
        styles = {n: None for n in SHEETS}
    hdr_fill = _fill('2A6099')
    from openpyxl.styles import Font, Alignment
    def put(ws, r, c, v, st=None, fill=None, bold=False, white=False, num=None):
        cell = ws.cell(r, c, v)
        if st is not None: cell._style = copy.copy(st)
        elif not tpl:
            cell.font = Font(name='맑은 고딕', size=10, bold=bold, color='FFFFFF' if white else '000000')
        if fill: cell.fill = _fill(fill)
        if num: cell.number_format = num
        return cell
    def hdr(ws, r, labels, st):
        for i, h in enumerate(labels, 1):
            put(ws, r, i, h, st, fill='2A6099' if not st else None, bold=True, white=True)

    # ---------- 1.입력판 ----------
    ws = wb['1.입력판']; S = styles['1.입력판'] or {}
    sI, sB, sC, sH = S.get('item'), S.get('itemB'), S.get('itemC'), S.get('head')
    ws['A1'] = '%s · 객실관리시스템 — 실행 산출 입력판' % site
    ws['A2'] = '수량 출처: %s / 견적서: %s / 단가: %s' % (meta.get('도면', '-'), meta.get('견적서', '미발행'), meta.get('단가장', '-'))
    ws['A3'] = '★ 노란 칸만 고치십시오. 나머지는 전부 자동 계산됩니다. 모르시면 비워 두십시오 — 그 줄은 0으로 잡히고 자가진단이 알립니다.'
    r = 5
    put(ws, r, 1, '■ 요율 입력 (전부 실행 기준)', sH, fill=None if sH else '2A6099', bold=True, white=True); ws.merge_cells('A%d:C%d' % (r, r)); r += 1
    R = {}
    for key, lab, note in (('조립비율', '조립비율 (CB 내부 자재비 × 이 비율, 외함 제외)', '표준 0.30'),
                           ('계약배수', '계약 실행가 배수', '표준 1.5'),
                           ('견적배수', '견적가 배수', '표준 1.8'),
                           ('예산배수', '설계 예산가 배수', '표준 2.1')):
        put(ws, r, 1, lab, sI); put(ws, r, 2, mult.get(key), sB, fill=YEL, num='0.00'); put(ws, r, 3, note, sC); R[key] = r; r += 1
    r += 1
    put(ws, r, 1, '■ 수량 (견적서/수량표에서 읽음)', sH, fill=None if sH else '2A6099', bold=True, white=True); ws.merge_cells('A%d:C%d' % (r, r)); r += 1
    R['cb'] = []
    for t in cb_qty:
        name, n = t[0], t[1]; desc = t[2] if len(t) > 2 else ''
        put(ws, r, 1, 'CONTROL BOX %s 대수%s' % (name, (' (%s)' % desc) if desc else ''), sI); put(ws, r, 2, n, sB, num='#,##0'); put(ws, r, 3, '수량표', sC); R['cb'].append(r); r += 1
    put(ws, r, 1, '총 실 수', sI); put(ws, r, 2, '=' + '+'.join('B%d' % x for x in R['cb']) if R['cb'] else 0, sB, num='#,##0')
    put(ws, r, 3, 'CB 대수 합. 기구물 수량과 일치해야 함', sC); R['rooms'] = r; r += 2
    put(ws, r, 1, '■ 미확인 실행단가 — 숫자를 넣으면 전 시트가 자동으로 바뀝니다', sH, fill=None if sH else '2A6099', bold=True, white=True); ws.merge_cells('A%d:C%d' % (r, r)); r += 1
    R['unk'] = {}
    for what, why, cand in unknown:
        put(ws, r, 1, what, sI); put(ws, r, 2, None, sB, fill=YEL, num='#,##0'); put(ws, r, 3, '%s %s' % (why, ('— 후보: ' + cand) if cand else ''), sC); R['unk'][what] = r; r += 1
    R['unk_first'] = min(R['unk'].values()) if R['unk'] else r; R['unk_last'] = max(R['unk'].values()) if R['unk'] else r
    for col, w in (('A', 46), ('B', 14), ('C', 95)): ws.column_dimensions[col].width = w

    # ---------- 2.CB 실행 ----------
    ws = wb['2.CB 실행']; S = styles['2.CB 실행'] or {}
    nT = max(1, len(cb_qty))
    ws['A1'] = 'CONTROL BOX 실행 산출 — %s (%s)' % (site, meta.get('CB출처', 'CB구성.csv'))
    labels = ['품목명 (형번)', '견적서 표기 (원문)'] + ['%s 수량(1대당)' % t[0] for t in cb_qty] + ['실행단가'] + ['%s 금액' % t[0] for t in cb_qty] + ['비고']
    hdr(ws, 3, labels, S.get('head'))
    cQ0 = 3; cP = 3 + nT; cA0 = cP + 1; cN = cA0 + nT
    r = 4; first = r
    for mod, txt, per, cost, memo in cb_rows:
        put(ws, r, 1, mod, S.get('item')); put(ws, r, 2, txt, S.get('itemB'))
        for i in range(nT): put(ws, r, cQ0 + i, per[i] if i < len(per) else per[0], S.get('num'), num='#,##0')
        if cost is None:
            key = _key(R, mod)
            put(ws, r, cP, "='1.입력판'!$B$%d" % R['unk'][key] if key else None, S.get('price'), fill=YEL, num='#,##0')
        else:
            put(ws, r, cP, cost, S.get('price'), num='#,##0')
        for i in range(nT): put(ws, r, cA0 + i, '=%s%d*$%s%d' % (L(cQ0 + i), r, L(cP), r), S.get('amt'), num='#,##0')
        put(ws, r, cN, memo, S.get('memo'), fill=YEL if ('[확인]' in memo) else None)
        r += 1
    last = r - 1
    def sub(label, fn, fill=None, memo=''):
        nonlocal r
        put(ws, r, 1, label, S.get('item'), fill=fill)
        for i in range(nT): put(ws, r, cA0 + i, fn(L(cA0 + i)), S.get('amt'), fill=fill, num='#,##0')
        put(ws, r, cN, memo, S.get('memo'), fill=fill)
        rr = r; r += 1; return rr
    R_MAT = sub('▣ CB 내부 자재비 소계', lambda c: '=SUM(%s%d:%s%d)' % (c, first, c, last), GRY)
    R_ASM = sub('▣ 조립비 (자재 × 조립비율)', lambda c: "=ROUND(%s%d*'1.입력판'!$B$%d,0)" % (c, R_MAT, R['조립비율']))
    ws.cell(R_ASM, 2, "=CONCATENATE(\"자재비 × \",TEXT('1.입력판'!$B$%d,\"0.00\"))" % R['조립비율'])
    R_BODY = sub('■ CB 본체 실행가 (외함 제외, 1대)', lambda c: '=%s%d+%s%d' % (c, R_MAT, c, R_ASM), GRN)
    enc_key = next((k for k in R['unk'] if '외함' in k), None)
    R_ENC = sub('■ CB 외함 세트 실행가 (1대)', lambda c: ("='1.입력판'!$B$%d" % R['unk'][enc_key]) if enc_key else 0, None,
                '[확인] 1.입력판 노란칸' if enc_key else meta.get('외함비고', ''))
    R_A = sub('■【A】CB 1대 실행가 (외함 포함)', lambda c: '=%s%d+%s%d' % (c, R_BODY, c, R_ENC), GRN)
    r += 1
    put(ws, r, 1, '◆ 대수', S.get('item'))
    for i in range(nT): put(ws, r, cA0 + i, "='1.입력판'!$B$%d" % R['cb'][i], S.get('amt'), num='#,##0')
    R_N = r; r += 1
    R_SA = sub('▣【A】실행 합계 (외함 포함 × 대수)', lambda c: '=%s%d*%s%d' % (c, R_A, c, R_N)); ws.cell(R_SA, cN, '=' + '+'.join('%s%d' % (L(cA0 + i), R_SA) for i in range(nT)))
    R_SB = sub('▣【B】실행 합계 (본체만 × 대수 / 외함 별도 라인)', lambda c: '=%s%d*%s%d' % (c, R_BODY, c, R_N)); ws.cell(R_SB, cN, '=' + '+'.join('%s%d' % (L(cA0 + i), R_SB) for i in range(nT)))
    r += 1
    R_Q = sub('▷ 참고 — 견적서 CONTROL BOX 단가', lambda c: "=ROUND(%s%d*'1.입력판'!$B$%d,0)" % (c, R_A, R['견적배수']), None,
              '견적서 발행 전: 【A】× 견적배수 잠정' if not meta.get('견적서') else '견적서 발행단가로 교체하십시오')
    sub('▷ 견적단가 ÷ 【A】실행가 = 실제 적용된 배수', lambda c: '=IF(%s%d=0,"",%s%d/%s%d)' % (c, R_A, c, R_Q, c, R_A))
    for i, w in enumerate([30, 46] + [13] * nT + [13] + [14] * nT + [70], 1): ws.column_dimensions[L(i)].width = w

    # ---------- 3.견적↔실행 대조 ----------
    ws = wb['3.견적↔실행 대조']; S = styles['3.견적↔실행 대조'] or {}
    ws['A1'] = '견적서 전 라인 ↔ 실행 대조 · 이윤 산출 — %s%s' % (site, '' if meta.get('견적서') else ' (견적서 미발행: 견적단가 = 실행 × 견적배수 잠정)')
    hdr(ws, 3, ['순번', '품명/기능', '규격', '단위', '수량', '견적단가', '견적금액', '실행단가', '실행금액', '이윤', '이윤율', '비고'], S.get('head'))
    r = 4; n = 0; lines = []
    def sec(t):
        nonlocal r
        put(ws, r, 1, t, S.get('sec'), fill=GRY)
        for c in range(2, 13): ws.cell(r, c).fill = _fill(GRY)
        r += 1
    def line(name, spec, unit, q, cost, memo, ref=None):
        nonlocal r, n
        n += 1
        vals = [n, name, spec, unit, q, "=ROUND(H%d*'1.입력판'!$B$%d,0)" % (r, R['견적배수']), '=E%d*F%d' % (r, r),
                ref if ref else cost, '=E%d*H%d' % (r, r), '=G%d-I%d' % (r, r), '=IF(G%d=0,"",(G%d-I%d)/G%d)' % (r, r, r, r), memo]
        for c, v in enumerate(vals, 1):
            put(ws, r, c, v, (S.get('cols') or [None] * 12)[c - 1], num='#,##0' if c in (5, 6, 7, 8, 9, 10) else ('0.0%' if c == 11 else None))
        if ref and "입력판" in str(ref): ws.cell(r, 8).fill = _fill(YEL); ws.cell(r, 12).fill = _fill(YEL)
        lines.append(r); r += 1
    groups = collections.OrderedDict((('중앙', '1. 중앙 시스템'), ('객실', '2. 객실 (CB · 기구물 · 노무)'), ('도어락', '3. RF DOOR LOCK SYSTEM (직접단가)')))
    for gk, gt in groups.items():
        rows = [x for x in qty if x[6] == gk] + ([] if gk != '객실' else [])
        if gk == '객실':
            sec(gt)
            for i, t in enumerate(cb_qty):
                line('CONTROL BOX %s%s' % (t[0], (' (%s)' % t[2]) if len(t) > 2 and t[2] else ''), meta.get('CB형번', 'CB-30BCB_4F 계열'), 'EA', "='1.입력판'!$B$%d" % R['cb'][i], None,
                     '【A】외함 포함 실행가. 2.CB 실행 참조', ref="='2.CB 실행'!$%s$%d" % (L(cA0 + i), R_A))
            for name, spec, unit, q, cost, memo, g in rows + [x for x in qty if x[6] == '노무']:
                key = _key(R, name) if cost is None else None
                line(name, spec, unit, q, cost, memo, ref=("='1.입력판'!$B$%d" % R['unk'][key]) if key else None)
                if cost is None and not key:
                    ws.cell(r - 1, 8).fill = _fill(YEL); ws.cell(r - 1, 12).fill = _fill(YEL)
        elif rows:
            sec(gt)
            for name, spec, unit, q, cost, memo, g in rows:
                key = _key(R, name) if cost is None else None
                line(name, spec, unit, q, cost, memo, ref=("='1.입력판'!$B$%d" % R['unk'][key]) if key else None)
                if cost is None and not key:
                    ws.cell(r - 1, 8).fill = _fill(YEL); ws.cell(r - 1, 12).fill = _fill(YEL)
    r += 1; R_TOT = r
    tot = [None, 'TOTAL (견적서 발행본 대조)', None, None, None, None, '=' + '+'.join('G%d' % x for x in lines), None,
           '=' + '+'.join('I%d' % x for x in lines), '=G%d-I%d' % (r, r), '=IF(G%d=0,"",(G%d-I%d)/G%d)' % (r, r, r, r), None]
    for c, v in enumerate(tot, 1): put(ws, r, c, v, (S.get('tot') or [None] * 12)[c - 1], fill=GRN, num='#,##0' if c in (7, 9, 10) else ('0.0%' if c == 11 else None))
    r += 1; R_NOGO = r
    put(ws, r, 2, '견적서 표기 총액 (NOGO)', (S.get('cols') or [None] * 12)[1]); put(ws, r, 7, meta.get('견적총액'), (S.get('cols') or [None] * 12)[6], fill=YEL if not meta.get('견적총액') else None, num='#,##0')
    put(ws, r, 12, '견적서 발행 후 총액을 넣으십시오' if not meta.get('견적총액') else '', (S.get('cols') or [None] * 12)[11])
    r += 1; R_CHK = r
    put(ws, r, 2, '검산 — 위 TOTAL 과의 차이', (S.get('cols') or [None] * 12)[1]); put(ws, r, 7, '=IF(G%d="","견적서 미발행",G%d-G%d)' % (R_NOGO, R_TOT, R_NOGO), (S.get('cols') or [None] * 12)[6], num='#,##0')
    put(ws, r, 12, '0 이면 전 라인 누락 없음 ✔', (S.get('cols') or [None] * 12)[11])
    for i, w in enumerate([6, 38, 24, 6, 9, 12, 14, 12, 14, 12, 9, 70], 1): ws.column_dimensions[L(i)].width = w

    # ---------- 4.확인 목록 ----------
    ws = wb['4.확인 목록']; S = styles['4.확인 목록'] or {}
    ws['A1'] = '확인해 주실 것 — 숫자만 알려주시면 1.입력판에 넣고 전 시트가 바뀝니다'
    hdr(ws, 3, ['#', '무엇을', '왜 못 정했나', '차장님 답', '금액 영향'], S.get('head'))
    r = 4
    for i, (what, why, cand) in enumerate(unknown, 1):
        for c, v in enumerate([i, what, why, None, cand], 1): put(ws, r, c, v, (S.get('cols') or [None] * 5)[c - 1], fill=YEL if c == 4 else None)
        r += 1
    for i, w in enumerate([5, 40, 58, 16, 44], 1): ws.column_dimensions[L(i)].width = w

    # ---------- 5.변경 이력 ----------
    ws = wb['5.변경 이력']; S = styles['5.변경 이력'] or {}
    hdr(ws, 1, ['버전', '일자', '등급', '무엇을', '근거', '이전 값'], S.get('head'))
    r = 2
    for row in [('v1', today().isoformat(), 'A', '%s 실행 산출 신규 작성 (28번 코드 생성, 정답본 틀 %s)' % (site, os.path.basename(tpl) if tpl else '내장'),
                 '수량 %s / 단가 %s / CB 구성 %s' % (meta.get('도면', '-'), meta.get('단가장', '-'), meta.get('CB출처', '-')), '—'),
                ('v1', today().isoformat(), 'A', '미확인 %d건은 1.입력판 노란칸으로 남김 (추측 금지)' % len(unknown), 'km-quote-builder 「셀을 못 채울 때의 철칙」', '—')]:
        for c, v in enumerate(row, 1): put(ws, r, c, v, (S.get('cols') or [None] * 6)[c - 1])
        r += 1
    for i, w in enumerate([8, 12, 6, 58, 70, 10], 1): ws.column_dimensions[L(i)].width = w

    # ---------- 0.자가진단 ----------
    ws = wb['0.자가진단']; S = styles['0.자가진단'] or {}
    ws['A1'] = '0. 자가진단 — 이 파일이 스스로 검사합니다 (%s)' % site
    ws['A2'] = '파일을 열면 이 시트부터 보십시오. 빨간 판정이 있으면 1.입력판 노란칸을 채우십시오.'
    hdr(ws, 4, ['검사 항목', '현재값', '기준', '판정', '뜻'], S.get('head'))
    rows0 = [('견적 총액 검산 (라인 합계 − NOGO 표기)', "='3.견적↔실행 대조'!$G$%d" % R_CHK, '0', '=IF(B5="견적서 미발행","— 견적서 미발행",IF(N(B5)=0,"✔ 정상","⚠ 조치 필요"))', '견적서가 나오면 0이어야 전 라인 누락 없음'),
             ('실 수 검산 (CB1+CB2+… − 기구물 수량)', "='1.입력판'!$B$%d-%d" % (R['rooms'], meta.get('객실수', 0)), '0', '=IF(N(B6)=0,"✔ 정상","⚠ 조치 필요")', '0이면 CB 대수와 실 수 일치'),
             ('단가 미입력 노란칸 수', "=COUNTBLANK('1.입력판'!B%d:B%d)" % (R['unk_first'], R['unk_last']), '0', '=IF(N(B7)=0,"✔ 정상","⚠ 조치 필요")', '남아 있으면 그만큼 실행금액이 0으로 잡혀 이윤이 과대평가됩니다'),
             ('현재 산출된 실행 총액', "='3.견적↔실행 대조'!$I$%d" % R_TOT, '—', '—', '미입력 칸이 0인 상태의 값'),
             ('현재 이윤율', "=IF('3.견적↔실행 대조'!$G$%d=0,\"\",'3.견적↔실행 대조'!$K$%d)" % (R_TOT, R_TOT), '—', '—', '견적서 미발행이면 (배수−1)/배수 로 고정. 발행 후 F열을 발행단가로 바꾸면 진짜 이윤율')]
    r = 5
    for row in rows0:
        for c, v in enumerate(row, 1): put(ws, r, c, v, (S.get('cols') or [None] * 5)[c - 1], num='#,##0' if c == 2 else None)
        r += 1
    r += 1
    put(ws, r, 1, '■ 이번 산출에서 확신 없이 넣은 것 / 규칙을 어긴 것', S.get('head2') or S.get('head'), fill=None if S else '2A6099', bold=True, white=True); r += 1
    for t in meta.get('자가신고', []):
        put(ws, r, 1, '· ' + t, S.get('note')); r += 1
    for i, w in enumerate([44, 18, 10, 14, 80], 1): ws.column_dimensions[L(i)].width = w
    wb.active = 0
    wb.save(out_path)
    return out_path

def _n(s): return re.sub(r'[^0-9A-Za-z가-힣]', '', str(s or '')).upper()

def _key(R, name):
    """1.입력판 노란칸 중 이 품목에 해당하는 것 : 똑같은 이름 > 앞 6글자"""
    n = _n(name)
    for k in R['unk']:
        if _n(k) == n: return k
    for k in R['unk']:
        if n[:6] and (n[:6] in _n(k) or _n(k)[:6] in n): return k
    return None

def _snap(ws):
    """정답본 시트에서 서식 견본을 뜬다 (머리글/품목/숫자/소계 등)"""
    S = {}
    try:
        t = ws.title
        if t == '1.입력판':
            S['head'] = copy.copy(ws['A5']._style); S['item'] = copy.copy(ws['A6']._style); S['itemB'] = copy.copy(ws['B6']._style); S['itemC'] = copy.copy(ws['C6']._style)
        elif t == '2.CB 실행':
            S['head'] = copy.copy(ws['A3']._style); S['item'] = copy.copy(ws['A4']._style); S['itemB'] = copy.copy(ws['B4']._style)
            S['num'] = copy.copy(ws['C4']._style); S['price'] = copy.copy(ws['E4']._style); S['amt'] = copy.copy(ws['F4']._style); S['memo'] = copy.copy(ws['H4']._style)
        elif t == '3.견적↔실행 대조':
            S['head'] = copy.copy(ws['A3']._style); S['sec'] = copy.copy(ws['A4']._style)
            S['cols'] = [copy.copy(ws.cell(5, c)._style) for c in range(1, 13)]; S['tot'] = [copy.copy(ws.cell(33, c)._style) for c in range(1, 13)]
        elif t == '4.확인 목록':
            S['head'] = copy.copy(ws['A3']._style); S['cols'] = [copy.copy(ws.cell(4, c)._style) for c in range(1, 6)]
        elif t == '5.변경 이력':
            S['head'] = copy.copy(ws['A1']._style); S['cols'] = [copy.copy(ws.cell(2, c)._style) for c in range(1, 7)]
        elif t == '0.자가진단':
            S['head'] = copy.copy(ws['A4']._style); S['cols'] = [copy.copy(ws.cell(5, c)._style) for c in range(1, 6)]
            S['head2'] = copy.copy(ws['A11']._style); S['note'] = copy.copy(ws['A12']._style)
    except Exception:
        pass
    return S

def _clear(ws):
    from openpyxl.styles import PatternFill
    for m in list(ws.merged_cells.ranges): ws.unmerge_cells(str(m))
    for row in ws.iter_rows():
        for c in row:
            c.value = None; c.fill = PatternFill(fill_type=None)
