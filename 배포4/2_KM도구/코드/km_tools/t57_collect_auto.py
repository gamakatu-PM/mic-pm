# -*- coding: utf-8 -*-
"""
57. 수금대장이 스스로 채워짐  (km_tools / t57_collect_auto)

최우선 5번 「수금」 A안 (2026-09-28, 차장님이 설계 질문 8개에 답하신 대로). 클로드(AI)를 쓰지 않는다 → 사용량 0.

    python t57_collect_auto.py [<meta폴더>] [--today 260928] [--quiet]
    (메타 폴더를 안 주면 회의록 쪽은 건너뛰고 확정 대장·견적서만 본다)

차장님 답 (2026-09-27)
    1 납품 사실  : 회의록에 나온다 + 차장님 한 줄 + 계산서 발행이 곧 납품  (세 통로 다)
    2 금액       : 계약(확정 대장 「계약금액」)
    3 계산서     : 경리(회사) → 도구는 경리에게 보낼 메일 본문만
    4 대장의 축  : 납품 3회 그대로 (외함 · 속판 · 기구물)
    5 결제조건   : 현장마다 다르다(계약서) → 확정 대장 「결제조건」 (없으면 빈칸, 독촉은 안 만든다)
    6 3회 나눔   : 견적서 구분별 소계 비율 (비율은 수금비율.csv 입력칸. 자동값은 「자동」, 차장님이 고치면 「손」)
    7 경리 메일  : 본문만 (보내는 것은 차장님)
    8 독촉       : 담당자에게 보낼 문안 (전화 첫마디 + 메일)

들어오는 것 (전부 이미 있는 것. 새 대장 없음)
    · 회의록 meta.json 「일정」「할 일」 의 납품 줄  →  납품예정일       (현장명은 meta.site 그대로. 복합회의 등은 「현장 확인」)
    · 확정 대장(43번·받은답 「확정,현장,항목,값」·앱 회신) :
        항목 「납품 외함」「계산서 속판」「입금 기구물」 + 값=날짜  → 납품일 / 계산서발행일 / 입금일
        항목 「계약금액」 → 금액의 뿌리   항목 「결제조건」 → 결제조건일수
    · 28번 견적서 xlsx (단가붙이기\\*\\{현장}_*_견적서_v*.xlsx) 내역서 → 외함/속판/기구물 소계 비율

나가는 것
    · _도구결과\\수금\\수금대장.csv  (9번·26번·아침 한 장이 그대로 읽는 7칸 + 납품예정일 + 근거)   ★ 빈칸만 채운다. 차장님이 적은 값은 절대 안 바꾼다
    · _도구결과\\수금\\수금비율.csv   (현장 | 외함 | 속판 | 기구물 | 근거 | 상태)   상태=손 이면 다시 계산하지 않는다
    · _도구결과\\수금\\수금대장_이력.csv (날짜 | 현장 | 구분 | 칸 | 이전 | 이후 | 근거)  B등급 이력
    · _도구결과\\수금\\{오늘}\\계산서요청_경리_{현장}_{구분}.txt   납품했는데 계산서가 없는 줄마다 (본문만)
    · _도구결과\\수금\\{오늘}\\독촉_{현장}_{구분}.txt               계산서 뒤 결제조건이 지났는데 입금이 없는 줄마다 (전화 첫마디 + 메일)
    · _도구결과\\수금\\{오늘}\\수금자동_{오늘}.json                 자가진단 (몇 줄 채웠나·왜 못 채웠나)

지키는 것
    · 금액·비율을 도구가 정하지 않는다 — 계약금액 × 비율(입력칸) 만. 둘 중 하나가 없으면 빈칸으로 두고 이유를 근거에 적는다
    · 현장명은 meta.site 그대로. 「이건 사실 ○○ 현장 것」 판단 없음
    · 「N일 지남」 을 메일 본문에 쓰지 않는다 (아침 한 장·9번이 D-day 를 보여 준다)
    · 메일을 보내지 않는다. 본문 파일까지
"""
from __future__ import print_function
import os, sys, io, re, csv, json, glob, datetime

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from common import *
import facts as FX

VERSION = 'v2.1 2026-09-28'   # v2.1 : 번호 56→57 (다른 창이 같은 날 t56_makequeue 를 먼저 냄, km-00 41번) · v2 : 독립 감사 9건 반영 — 현장은 정확히 같을 때만(유일한 앞머리 일치만 허용) · 재저장 때 줄 안 지움 · 「소」「합」 품목 누락 · 결제조건/계약금액은 숫자만 · 특수문자 현장명 · --today 엉터리 · 폴더 날짜
# v1 2026-09-28
TOOL = '수금'
KINDS = ('외함', '속판', '기구물')
HEAD = ['현장', '구분(외함/속판/기구물)', '납품일', '금액', '계산서발행일', '입금일', '결제조건일수', '납품예정일', '근거']
RATIO_HEAD = ['현장', '외함', '속판', '기구물', '근거', '상태']
HIST_HEAD = ['날짜', '현장', '구분', '칸', '이전', '이후', '근거']
NOT_A_SITE = ('복합회의', '확인필요', '확인 필요', '현장미정', '미정', '수금채크', '수금체크', '삭제요망')

KIND_WORDS = {'외함': ('외함',), '속판': ('속판', 'CB', 'C/B', '컨트롤박스', '컨트롤 박스', '제어분전', '제어 분전'),
              '기구물': ('기구물', '기구', '챠임', '차임', '키센서', '온도조절', '스위치', 'BSP', '인디케이터')}
DELIV_WORDS = ('납품', '출고', '입고', '반입')
FACT_RE = re.compile(r'^(납품|납품일|계산서|계산서발행|계산서발행일|입금|입금일)\s*[-_ ]?\s*(외함|속판|기구물)$')


# ── 작은 도우미 ──────────────────────────────────────────────
def _n(s):
    # 빈칸·문장부호·cp949 밖 글자(😀 등, write_csv 가 ? 로 바꿈)를 다 떼고 글자·숫자만 남긴다
    return re.sub(r'[^0-9A-Za-z가-힣]', '', str(s or ''))


def same_site(a, b):
    """정확히 같은 현장만 True. (v2 : 「조선호텔」 ⊂ 「조선호텔 리뉴얼」 같은 부분일치는 오귀속 사고라 뺌 — CLAUDE.md 3-1)"""
    a, b = _n(a), _n(b)
    return bool(a) and a == b


def pick_site(site, names):
    """names 중 site 와 맞는 하나. 정확 일치 → 그것. 없으면 앞머리 일치(동구로초 ⊂ 동구로초등학교)가 **딱 하나**일 때만 그것.
    둘 이상이면 None (어느 것인지 도구가 정하지 않는다)."""
    n = _n(site)
    if not n:
        return None
    exact = [x for x in names if _n(x) == n]
    if exact:
        return exact[0]
    pre = [x for x in names if len(n) >= 3 and len(_n(x)) >= 3 and (_n(x).startswith(n) or n.startswith(_n(x)))]
    return pre[0] if len(pre) == 1 else None


def is_site(name):
    n = re.sub(r'\s+', '', name or '')
    if not n:
        return False
    return not any(x.replace(' ', '') in n for x in NOT_A_SITE)


def d6(s):
    """260929 / 2026-09-29 / 2026.9.29 → date, 아니면 None"""
    s = str(s or '').strip()
    m = re.match(r'^(\d{2}|\d{4})\D+(\d{1,2})\D+(\d{1,2})$', s)
    if m:
        y = int(m.group(1)); y = y + 2000 if y < 100 else y
        try:
            return datetime.date(y, int(m.group(2)), int(m.group(3)))
        except Exception:
            return None
    d = re.sub(r'\D', '', s)
    if len(d) == 8:
        d = d[2:]
    if len(d) == 6:
        try:
            return datetime.date(2000 + int(d[:2]), int(d[2:4]), int(d[4:6]))
        except Exception:
            return None
    return None


def iso(dt):
    return dt.isoformat() if dt else ''


def money(s):
    """「90,000,000원」 「90000000」 만 숫자로. 「9천만원」「1.5억」「30일/60일」「선수금 30%」 처럼 숫자를 이어 붙이면 뜻이 바뀌는 글은 None."""
    t = re.sub(r'[\s,]', '', str(s or ''))
    t = re.sub(r'(원|일|VAT별도)$', '', t)
    return int(t) if re.match(r'^\d+$', t) else None


def kind_of(text):
    t = str(text or '')
    for k, words in KIND_WORDS.items():
        if any(w.lower() in t.lower() for w in words):
            return k
    return ''


# ── 1. 회의록 → 납품 예정 ───────────────────────────────────
SCHED_RE = re.compile(r'^\s*[-·•]?\s*([^|]+?)\s*\|\s*(\d{6})\s*\|\s*(.*?)\s*$')
TODO_RE = re.compile(r'^\s*[-·•]?\s*(\d{6})\s*\|\s*(.+?)\s*(?:\|\s*(.*?))?\s*$')


def _lines(v):
    if v is None:
        return []
    if isinstance(v, str):
        return [x for x in v.split('\n') if x.strip()]
    out = []
    for it in v:
        if isinstance(it, (list, tuple)):
            out.append(' | '.join(str(c) for c in it if str(c or '').strip()))
        elif isinstance(it, dict):
            out.extend(_lines(list(it.values())))
        elif str(it or '').strip():
            out.append(str(it))
    return out


def scan_meta(meta_dir):
    """[(현장, 구분, 예정일 date, 근거)] , [현장 확인 목록]"""
    out, odd = [], []
    if not meta_dir or not os.path.isdir(meta_dir):
        return out, odd
    for root, _d, files in os.walk(meta_dir):
        for fn in sorted(files):
            if not fn.endswith('meta.json') or fn.startswith('_삭제요망'):
                continue
            try:
                j = json.load(io.open(os.path.join(root, fn), encoding='utf-8'))
            except Exception:
                continue
            m = j.get('meta') or {}
            s1 = (j.get('sec') or {}).get('1') or {}
            s2 = (j.get('sec') or {}).get('2') or {}
            site = (m.get('site') or s1.get('현장') or '').strip()
            cands = []
            for ln in _lines(s1.get('일정')):
                mm = SCHED_RE.match(ln)
                if mm and (any(w in mm.group(1) for w in DELIV_WORDS) or any(w in mm.group(3) for w in DELIV_WORDS)):
                    cands.append((mm.group(2), mm.group(1) + ' ' + mm.group(3)))
            for ln in _lines(s2.get('할 일')):
                mm = TODO_RE.match(ln)
                if mm and any(w in mm.group(2) for w in DELIV_WORDS):
                    cands.append((mm.group(1), mm.group(2)))
            for ymd, text in cands:
                dt = d6(ymd)
                if not dt:
                    continue
                k = kind_of(text)
                if not is_site(site):
                    odd.append((site, k, iso(dt), text, fn))
                    continue
                out.append((site, k, dt, '회의록 %s %s「%s」' % (fn[:22], ymd, text[:40])))
    return out, odd


# ── 2. 확정 대장 → 납품일·계산서·입금·계약금액·결제조건 ──────
def scan_facts():
    """{(현장,구분): {'납품일':date,'계산서발행일':date,'입금일':date, 근거…}} , {현장: 계약금액}, {현장: 결제조건}"""
    got, contract, term, bad = {}, {}, {}, []
    try:
        rows = FX.load(active_only=True)
    except Exception:
        rows = []
    for r in rows:
        item = re.sub(r'\s+', '', r.get('항목') or '')
        site = (r.get('현장') or '').strip()
        val = (r.get('값') or '').strip()
        if not site or not is_site(site):
            continue
        if item == '계약금액':
            v = money(val)
            if v is None:
                bad.append('%s 계약금액 「%s」 는 숫자만 받습니다 (예 90000000 · 90,000,000원)' % (site, val))
            elif site not in contract:
                contract[site] = (v, '확정 대장 %s' % (r.get('일자') or ''))
            continue
        if item in ('결제조건', '결제조건일수'):
            v = money(val)
            if v is None:
                bad.append('%s 결제조건 「%s」 는 일수 숫자만 받습니다 (예 30)' % (site, val))
            elif site not in term:
                term[site] = (v, '확정 대장 %s' % (r.get('일자') or ''))
            continue
        mm = FACT_RE.match(item)
        if not mm:
            continue
        what = {'납품': '납품일', '납품일': '납품일', '계산서': '계산서발행일', '계산서발행': '계산서발행일',
                '계산서발행일': '계산서발행일', '입금': '입금일', '입금일': '입금일'}[mm.group(1)]
        dt = d6(val)
        if not dt:
            continue
        key = (site, mm.group(2))
        got.setdefault(key, {})
        if what not in got[key]:              # 최신이 앞이라 처음 것만
            got[key][what] = (dt, '확정 대장 %s 「%s」' % (r.get('일자') or '', r.get('항목') or ''))
    return got, contract, term, bad


# ── 3. 견적서 → 3회 비율 ────────────────────────────────────
def quote_split(path):
    """내역서에서 외함/속판/기구물 금액 합. (dict, 근거) — 못 읽으면 (None, 이유)"""
    try:
        import openpyxl
        wb = openpyxl.load_workbook(path, data_only=True)
    except Exception as e:
        return None, '견적서 못 엶 %s' % e
    ws = next((wb[s] for s in wb.sheetnames if '내역서' in s), None)
    if ws is None:
        return None, '내역서 시트 없음'
    tot = {'외함': 0, '속판': 0, '기구물': 0}
    sec = ''
    for r in range(1, (ws.max_row or 1) + 1):
        a = str(ws.cell(r, 1).value or '').strip()
        if re.match(r'^\d+\.\s*중앙', a):
            sec = '중앙'; continue
        if re.match(r'^\d+\.\s*객실', a):
            sec = '객실'; continue
        if not a or re.match(r'^(소\s*계|합\s*계)$', a) or a.startswith('*') or a.startswith('-'):
            continue
        v = ws.cell(r, 11).value
        if not isinstance(v, (int, float)):
            q, p1, p2 = ws.cell(r, 4).value, ws.cell(r, 5).value, ws.cell(r, 7).value
            p = p1 if isinstance(p1, (int, float)) else (p2 if isinstance(p2, (int, float)) else None)
            v = q * p if isinstance(q, (int, float)) and p is not None else None
        if not isinstance(v, (int, float)):
            continue
        if '외함' in a:
            tot['외함'] += v
        elif sec == '중앙' or a.upper().startswith('CONTROL BOX'):
            tot['속판'] += v
        elif sec == '객실':
            tot['기구물'] += v
    s = sum(tot.values())
    if s <= 0:
        return None, '내역서 금액이 0 (수식 값이 파일에 없으면 28번을 다시 돌리십시오)'
    return tot, '견적서 %s' % os.path.basename(path)


def find_quote(site):
    base = os.path.join(cfg('out'), '단가붙이기')
    cands = []
    names = {}
    for p in glob.glob(os.path.join(base, '*', '*_견적서_v*.xlsx')):
        m = re.match(r'^(.*?)(?:_\d{6})?_견적서_v', os.path.basename(p))
        if m:
            names.setdefault(m.group(1), []).append(p)
    hit = pick_site(site, list(names.keys()))
    if not hit:
        return None
    return sorted(names[hit], key=os.path.getmtime)[-1]


def ratio_path():
    p = os.path.join(cfg('out'), '수금')
    os.makedirs(p, exist_ok=True)
    return os.path.join(p, '수금비율.csv')


def _read_csv(p):
    for enc in ('cp949', 'utf-8-sig', 'utf-8'):
        try:
            with io.open(p, 'r', encoding=enc, newline='') as fp:
                return [r for r in csv.reader(fp)]
        except Exception:
            continue
    return []


def load_ratio():
    p = ratio_path()
    out = {}
    if os.path.exists(p):
        for r in _read_csv(p)[1:]:
            r = (r + [''] * 6)[:6]
            if r[0].strip():
                out[r[0].strip()] = {'외함': r[1], '속판': r[2], '기구물': r[3], '근거': r[4], '상태': r[5].strip() or '자동'}
    return out


def ensure_ratio(site, ratios):
    """현장의 비율 줄. 상태=손 이면 그대로. 없거나 자동이면 견적서에서 계산. 못 하면 빈 줄(노란)."""
    row = ratios.get(site)
    if row and row['상태'] == '손':
        return row, False
    q = find_quote(site)
    if q:
        tot, why = quote_split(q)
        if tot:
            s = float(sum(tot.values()))
            new = {'외함': '%.4f' % (tot['외함'] / s), '속판': '%.4f' % (tot['속판'] / s), '기구물': '%.4f' % (tot['기구물'] / s), '근거': why, '상태': '자동'}
            changed = (row or {}).get('근거') != why or not row
            ratios[site] = new
            return new, changed
        why2 = why
    else:
        why2 = '견적서 없음 (단가붙이기 폴더)'
    if not row:
        ratios[site] = {'외함': '', '속판': '', '기구물': '', '근거': why2, '상태': '자동'}
        return ratios[site], True
    return row, False


def save_ratio(ratios):
    rows = [[s, v['외함'], v['속판'], v['기구물'], v['근거'], v['상태']] for s, v in sorted(ratios.items())]
    write_csv(ratio_path(), rows, RATIO_HEAD)


# ── 4. 수금대장 읽기·채우기 ─────────────────────────────────
def book_path():
    p = os.path.join(cfg('out'), '수금')
    os.makedirs(p, exist_ok=True)
    return os.path.join(p, '수금대장.csv')


def load_book():
    """대장의 모든 줄(「예)」 줄·현장 빈칸 줄도) 그대로. 채우는 대상은 work() 가 고른다. 재저장 때 한 줄도 안 지운다 (v2)."""
    p = book_path()
    rows = []
    if os.path.exists(p):
        for r in _read_csv(p)[1:]:
            r = (r + [''] * 9)[:9]
            rows.append([c.strip() for c in r])
    return rows


def workable(r):
    return bool(r[0].strip()) and not r[0].startswith('예)')


_T0 = [None]


def hist_add(hist, site, kind, col, before, after, why):
    hist.append([iso(_T0[0] or datetime.date.today()), site, kind, col, before, after, why])


def run(meta_dir=None, today_ymd=None, quiet=False):
    t0 = (d6(today_ymd) if today_ymd else None) or today()
    if today_ymd and not d6(today_ymd):
        print('※ --today %s 를 날짜로 못 읽어 오늘(%s)로 합니다' % (today_ymd, t0))
    _T0[0] = t0
    if not quiet:
        title('57. 수금대장이 스스로 채워짐   (회의록·확정 대장·견적서 → 수금대장. 빈칸만. 토큰 0)')
    book = load_book()
    ratios = load_ratio()
    hist, notes = [], []
    def row_for(site, kind, make=True):
        names = sorted(set(r[0] for r in book if workable(r) and r[1].strip() == kind))
        hit = pick_site(site, names)
        if hit:
            return next(r for r in book if r[0] == hit and r[1].strip() == kind)
        if not make:
            return None
        r = [site, kind, '', '', '', '', '', '', '']
        book.append(r)
        hist_add(hist, site, kind, '줄', '', '새 줄', '자동')
        return r

    def fill(r, col, val, why):
        i = HEAD.index(col)
        if r[i].strip():
            return False
        r[i] = val
        hist_add(hist, r[0], r[1], col, '', val, why)
        return True

    # 1) 회의록 → 납품예정일
    plans, odd = scan_meta(meta_dir)
    n_plan = 0
    for site, kind, dt, why in plans:
        kinds = [kind] if kind else []
        if not kinds:
            notes.append('%s %s 납품 줄에 구분(외함/속판/기구물)이 없어 예정일을 못 넣음 — %s' % (site, iso(dt), why))
            continue
        r = row_for(site, kind)
        if fill(r, '납품예정일', iso(dt), why):
            n_plan += 1
    for site, k, dt, text, fn in odd:
        notes.append('현장 확인 : 「%s」 회의록의 납품 줄(%s %s) — 현장이 아닌 이름이라 넣지 않음 (%s)' % (site or '빈칸', dt, text[:30], fn[:22]))

    # 2) 확정 대장 → 납품일·계산서·입금·결제조건
    got, contract, term, bad = scan_facts()
    notes.extend(bad)
    n_fact = 0
    for (site, kind), d in got.items():
        r = row_for(site, kind)
        for col in ('납품일', '계산서발행일', '입금일'):
            if col in d and fill(r, col, iso(d[col][0]), d[col][1]):
                n_fact += 1
        if '계산서발행일' in d and not r[HEAD.index('납품일')].strip():   # 답 1-다 : 계산서 발행이 곧 납품
            if fill(r, '납품일', iso(d['계산서발행일'][0]), '계산서 발행일 = 납품일 (차장님 답 1-다)'):
                n_fact += 1
    for r in book:
        if not workable(r):
            continue
        hit = pick_site(r[0], list(term.keys()))
        if hit and fill(r, '결제조건일수', str(term[hit][0]), term[hit][1]):
            n_fact += 1

    # 3) 금액 = 계약금액 × 비율 (둘 다 있을 때만)
    n_amt = 0
    ratio_changed = False
    for r in book:
        site, kind = r[0], r[1]
        if not workable(r) or r[HEAD.index('금액')].strip() or kind not in KINDS:
            continue
        hit = pick_site(site, list(contract.keys()))
        c = contract[hit] if hit else None
        if not c:
            notes.append('%s %s 금액 빈칸 — 확정 대장에 「계약금액」 이 없음 (43번 : 확정,%s,계약금액,금액)' % (site, kind, site))
            continue
        rr, ch = ensure_ratio(site, ratios)
        ratio_changed = ratio_changed or ch
        try:
            f = float(rr[kind])
        except Exception:
            notes.append('%s %s 금액 빈칸 — 수금비율.csv 의 %s 비율이 비어 있음 (%s)' % (site, kind, kind, rr.get('근거') or ''))
            continue
        amt = int(round(c[0] * f))
        if fill(r, '금액', str(amt), '%s × %s비율 %s (%s, %s)' % (c[1], kind, rr[kind], rr['상태'], rr['근거'][:30])):
            n_amt += 1

    # 4) 근거 칸 갱신 + 저장
    for r in book:
        if not workable(r):
            continue
        why = [h[6] for h in hist if h[1] == r[0] and h[2] == r[1] and h[3] != '줄']
        if why:
            r[8] = (r[8] + ' / ' if r[8] else '') + ' / '.join(w[:60] for w in why[-3:])
    changed = bool(hist)
    if changed or not os.path.exists(book_path()):
        write_csv(book_path(), book, HEAD)
    if ratio_changed or (ratios and not os.path.exists(ratio_path())):
        save_ratio(ratios)
    if hist:
        hp = os.path.join(cfg('out'), '수금', '수금대장_이력.csv')
        old = _read_csv(hp)[1:] if os.path.exists(hp) else []
        write_csv(hp, old + hist, HIST_HEAD)

    # 5) 경리 메일 본문 · 독촉 문안
    od = os.path.join(cfg('out'), TOOL, ymd6(t0))
    os.makedirs(od, exist_ok=True)
    mails, dun = [], []
    for r in book:
        if not workable(r):
            continue
        site, kind, dlv, amt, bill, paid, tm = r[:7]
        if dlv and not bill:
            p = os.path.join(od, '계산서요청_경리_%s_%s.txt' % (safe_name(site, 20), kind))
            io.open(p, 'w', encoding='utf-8').write(mail_body(site, kind, dlv, amt))
            mails.append(p)
        elif bill and not paid:
            bd, tn = d6(bill), money(tm)
            if bd and tn is not None and (t0 - bd).days > tn:
                p = os.path.join(od, '독촉_%s_%s.txt' % (safe_name(site, 20), kind))
                io.open(p, 'w', encoding='utf-8').write(dun_body(site, kind, bill, amt, tn))
                dun.append(p)
            elif bd and tn is None:
                notes.append('%s %s 입금 대기인데 결제조건이 없거나 숫자가 아니라(「%s」) 독촉 판단 못 함 (43번 : 확정,%s,결제조건,30)' % (site, kind, tm, site))

    rep = {'version': VERSION, 'today': iso(t0), '줄': len(book), '납품예정 채움': n_plan, '확정대장 채움': n_fact, '금액 채움': n_amt,
           '경리메일': [os.path.basename(x) for x in mails], '독촉': [os.path.basename(x) for x in dun], '못 채운 이유': notes,
           '이력': len(hist), '대장': book_path(), '비율': ratio_path()}
    io.open(os.path.join(od, '수금자동_%s.json' % ymd6(t0)), 'w', encoding='utf-8').write(json.dumps(rep, ensure_ascii=False, indent=1))
    if not quiet:
        print('줄 %d · 납품예정 %d · 확정대장 %d · 금액 %d 채움 / 경리 메일 %d · 독촉 %d' % (len(book), n_plan, n_fact, n_amt, len(mails), len(dun)))
        for n in notes:
            print('  ※ ' + n)
        print('대장 : %s' % book_path())
    try:
        log(TOOL, '자동 %d줄 채움' % len(hist))
    except Exception:
        pass
    return rep


def mail_body(site, kind, dlv, amt):
    return '\n'.join([
        '제목: [세금계산서 발행 요청] %s %s 납품분 (%s)' % (site, kind, dlv),
        '',
        '경리팀 담당자님께,',
        '',
        '%s 현장 %s 납품이 %s 에 완료되어 세금계산서 발행을 요청드립니다.' % (site, kind, dlv),
        '',
        '  · 현장 : %s' % site,
        '  · 구분 : %s (납품 3회 중)' % kind,
        '  · 납품일 : %s' % dlv,
        '  · 금액 : %s' % ((won(int(amt)) + '원 (VAT 별도)') if str(amt).strip().isdigit() else '[   ]원 (VAT 별도)'),
        '  · 받는 곳 : [   ]  (사업자·담당자)',
        '  · 첨부 : 납품확인서 [   ] / 사진대지 [   ]',
        '',
        '발행되면 발행일을 알려 주십시오. (수금대장에 적고 결제조건에 맞춰 입금을 확인하겠습니다)',
        '',
        '한국마이크로닉(주) 배성윤 드림',
        '',
        '※ 이 본문은 수금대장에서 만든 초안입니다. 금액·받는 곳을 확인하신 뒤 보내십시오. 발행일은 43번에 「확정,%s,계산서 %s,YYMMDD」.' % (site, kind),
    ])


def dun_body(site, kind, bill, amt, term):
    a = (won(int(amt)) + '원') if str(amt).strip().isdigit() else '[   ]원'
    return '\n'.join([
        '[전화 첫마디]',
        '안녕하세요, 한국마이크로닉 배성윤입니다. %s 현장 %s 분 세금계산서(%s 발행) 입금 건으로 연락드렸습니다. 결제 예정일을 여쭤봐도 될까요?' % (site, kind, bill),
        '',
        '[메일]',
        '제목: [입금 확인 요청] %s %s 납품분 세금계산서 (%s 발행)' % (site, kind, bill),
        '',
        '[담당자] 님께,',
        '',
        '%s 현장 %s 납품분 세금계산서를 %s 에 발행해 드렸습니다 (%s, 결제조건 %d일).' % (site, kind, bill, a, term),
        '아직 입금이 확인되지 않아 결제 예정일을 여쭙습니다. 처리 중이시면 예정일만 알려 주시면 됩니다.',
        '',
        '  · 현장 : %s' % site,
        '  · 구분 : %s' % kind,
        '  · 계산서 발행일 : %s' % bill,
        '  · 금액 : %s (VAT 별도)' % a,
        '',
        '감사합니다.',
        '한국마이크로닉(주) 배성윤 드림',
        '',
        '※ 수금대장에서 만든 초안입니다. 담당자 이름은 현장대장·전화번호사전에서 확인해 넣으십시오. 입금되면 43번에 「확정,%s,입금 %s,YYMMDD」.' % (site, kind),
    ])


def main(argv):
    meta = None
    opt = {}
    i = 0
    while i < len(argv):
        if argv[i].startswith('--'):
            if argv[i] == '--quiet':
                opt['quiet'] = True; i += 1
            elif i + 1 < len(argv):
                opt[argv[i][2:]] = argv[i + 1]; i += 2
            else:
                i += 1
        else:
            meta = argv[i]; i += 1
    run(meta, opt.get('today'), opt.get('quiet', False))
    return 0


if __name__ == '__main__':
    sys.exit(main(sys.argv[1:]))
