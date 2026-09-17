# -*- coding: utf-8 -*-
"""44. 회의록 확인 - 저장만 해 둔 PLAUD 회의록을 「읽었다」고 표시한다. 토큰 0.
프로님 : "플라우드 회의록은 저장만 하고 (넘어가네). 내가 확인을 못 하니 미확인. 그것을 매일 표현해 달라."

  · 표시 안 된 회의는 42 아침 한 장 **맨 위 「미확인 회의록」** 에 매일 뜬다 (오래될수록 빨강).
  · 44번을 누르면 미확인 목록이 번호와 함께 나온다. 읽은 번호를 「1,3,5」 처럼 넣으면 그 줄이 사라진다.
  · 「a」 = 전부 확인 / 「o3」 = 3번 회의록 워드 열기 / 엔터 = 나가기
  · 받은답 csv 한 줄 `회의확인,<회의폴더>,확인,메모` 로도 된다 (30번·36번이 반영).
확인 대장 : _도구결과\\_대장\\회의확인.csv  (확인일·회의폴더·현장·메모·누가)
"""
import os, sys, datetime
import common
from common import *
import facts

TOOL = '회의록확인'

def unchecked(days=None):
    """미확인 회의 [dict] 오래된 것 먼저"""
    import t42_morning as MO
    done = facts.checked()
    out = []
    for m in MO.meetings():
        if facts._norm(os.path.basename(m['folder'])) in done:
            continue
        gap = (today() - m['day']).days if m['day'] else None
        if days is not None and gap is not None and gap > days:
            continue
        m['gap'] = gap
        out.append(m)
    out.sort(key=lambda m: m['day'] or datetime.date(2000, 1, 1))
    return out

def line(i, m):
    return '%2d. %-10s %-12s %-16s %s%s' % (
        i, (m['day'].isoformat()[5:] if m['day'] else '날짜?'), (m['site'] or '?')[:12], (m['who'] or '')[:16],
        (m['agenda'] or (m['decisions'][0] if m['decisions'] else ''))[:40],
        ('  [%d일 지남]' % m['gap']) if m['gap'] and m['gap'] > 2 else '')

def run():
    title('44. 회의록 확인   (저장만 해 둔 회의록을 읽었다고 표시. 미확인은 아침 한 장 맨 위에 매일 뜬다)')
    xs = unchecked()
    if not xs:
        print('미확인 회의록이 없습니다. 전부 확인하셨습니다.')
        return True
    xlsx, csvp = build_xlsx(xs, quiet=True)
    print('미확인 회의록 %d건 (오래된 것부터)' % len(xs))
    print('요약 엑셀 : %s' % (xlsx or csvp))
    print('-' * 88)
    for i, m in enumerate(xs, 1):
        print(line(i, m))
        for t in m['todos'][:2]:
            print('     할 일 : %s' % t[:70])
        for d in m['decisions'][:1]:
            print('     결정   : %s' % d[:70])
    print('-' * 88)
    print('읽은 번호를 넣으십시오.  예) 1,3,5   /   a = 전부 확인   /   o3 = 3번 워드 열기   /   엔터 = 나가기')
    while True:
        s = ask('> ', '').strip()
        if not s:
            return True
        if s.lower().startswith('o'):
            try:
                m = xs[int(s[1:]) - 1]
                open_file(m['docx']) or open_folder(m['folder'])
                print('열었습니다 : %s' % os.path.basename(m['docx']))
            except Exception:
                print('번호를 다시 넣으십시오.')
            continue
        if s.lower() == 'a':
            pick = list(range(1, len(xs) + 1))
        else:
            pick = []
            for t in s.replace(' ', '').split(','):
                if t.isdigit() and 1 <= int(t) <= len(xs):
                    pick.append(int(t))
        if not pick:
            print('번호를 다시 넣으십시오.')
            continue
        n = 0
        for i in pick:
            m = xs[i - 1]
            if facts.check_meeting(os.path.basename(m['folder']), m['site'], (m['agenda'] or '')[:40]):
                n += 1
        print('%d건 확인 처리했습니다. (대장 : %s)' % (n, facts.chk_path()))
        left = unchecked()
        print('남은 미확인 %d건' % len(left))
        build_xlsx(left, quiet=True)
        try:
            if not common.AUTO:
                import t42_morning as MO
                MO.build(quiet=True)
                print('아침 한 장을 다시 만들었습니다.')
        except Exception:
            pass
        return True

# ---------------- 요약 엑셀 (내용을 빠뜨리지 않는다) ----------------

SUM_HEAD = ['번호', '회의일', '며칠 지남', '현장', '협의자', '안건', '결정사항', '조치사항',
            '할 일', '수량·규격 변경', '★대외 언급 금지', '리스크', '타부서 전달', '회의록 파일']

def summary_rows(xs=None):
    xs = xs if xs is not None else unchecked()
    rows = []
    J = lambda v: '\n'.join(v) if v else ''
    for i, m in enumerate(xs, 1):
        rows.append([i, (m['day'].isoformat() if m['day'] else ''), (m.get('gap') if m.get('gap') is not None else ''),
                     m['site'], m['who'], m['agenda'], J(m['decisions']), J(m['actions']), J(m['todos']),
                     J(m['changes']), J(m['secret']), J(m['risk']), J(m['dept']), m['docx']])
    return rows

def build_xlsx(xs=None, quiet=True):
    """미확인 회의록 요약 엑셀. 요약 1장 + 할일·변경·대외금지 따로 + 전체내용 1장"""
    xs = xs if xs is not None else unchecked()
    od = outdir(TOOL)
    rows = summary_rows(xs)
    csvp = write_csv(os.path.join(od, '미확인회의록_요약_%s.csv' % ymd6()),
                     [[str(c).replace('\n', ' / ') for c in r] for r in rows], SUM_HEAD)
    xlsx = None
    try:
        ensure_pkg('openpyxl', 'openpyxl')
        import openpyxl
        from openpyxl.styles import Font, Alignment, PatternFill
        from openpyxl.utils import get_column_letter
        wb = openpyxl.Workbook()
        HEADF = Font(bold=True, color='FFFFFF'); HEADB = PatternFill('solid', fgColor='2A6099')
        RED = PatternFill('solid', fgColor='FADBD8'); WRAP = Alignment(wrap_text=True, vertical='top')
        ws = wb.active; ws.title = '1.요약'
        ws.append(SUM_HEAD)
        for c in ws[1]:
            c.font = HEADF; c.fill = HEADB; c.alignment = WRAP
        for r in rows:
            ws.append(r)
            if isinstance(r[2], int) and r[2] >= 3:
                for c in ws[ws.max_row]:
                    c.fill = RED
        for i, w in enumerate([5, 11, 8, 14, 18, 30, 40, 34, 40, 34, 28, 30, 26, 46], 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = WRAP
        ws.freeze_panes = 'A2'
        for name, key, head in (('2.할 일', 'todos', ['회의일', '현장', '협의자', '할 일']),
                                ('3.수량·규격 변경', 'changes', ['회의일', '현장', '협의자', '변경 내용']),
                                ('4.대외 언급 금지', 'secret', ['회의일', '현장', '협의자', '★대외 언급 금지']),
                                ('5.리스크·타부서', 'risk', ['회의일', '현장', '협의자', '리스크']),
                                ('6.타부서 전달', 'dept', ['회의일', '현장', '협의자', '타부서 전달 (작업의뢰서)'])):
            w2 = wb.create_sheet(name)
            w2.append(head)
            for c in w2[1]:
                c.font = HEADF; c.fill = HEADB
            for m in xs:
                for x in m.get(key) or []:
                    w2.append([(m['day'].isoformat() if m['day'] else ''), m['site'], m['who'], x])
            for i, wd in enumerate([11, 14, 18, 90], 1):
                w2.column_dimensions[get_column_letter(i)].width = wd
            for row in w2.iter_rows(min_row=2):
                for c in row:
                    c.alignment = WRAP
            w2.freeze_panes = 'A2'
        w3 = wb.create_sheet('7.전체 내용')
        w3.append(['회의', '절', '내용'])
        for c in w3[1]:
            c.font = HEADF; c.fill = HEADB
        for m in xs:
            tag = '%s %s %s' % ((m['day'].isoformat() if m['day'] else ''), m['site'], m['who'])
            w3.append([tag, '안건', m['agenda']])
            for nm, key in (('결정사항', 'decisions'), ('조치사항', 'actions'), ('할 일', 'todos'),
                            ('수량·규격 변경', 'changes'), ('★대외 언급 금지', 'secret'), ('리스크', 'risk'), ('타부서 전달', 'dept')):
                for x in m.get(key) or []:
                    w3.append([tag, nm, x])
            w3.append([tag, '회의록 파일', m['docx']])
        for i, wd in enumerate([34, 16, 100], 1):
            w3.column_dimensions[get_column_letter(i)].width = wd
        for row in w3.iter_rows(min_row=2):
            for c in row:
                c.alignment = WRAP
        w3.freeze_panes = 'A2'
        xlsx = os.path.join(od, '미확인회의록_요약_%s.xlsx' % ymd6())
        wb.save(xlsx)
        try:
            finish_xlsx(xlsx)
        except Exception:
            pass
    except Exception as e:
        if not quiet:
            print('[엑셀 실패] %s (csv 는 나왔습니다)' % e)
    if not quiet:
        print('요약 엑셀 : %s' % (xlsx or '실패'))
        print('요약 csv  : %s' % csvp)
    return xlsx, csvp

if __name__ == '__main__':
    run(); pause()
