# -*- coding: utf-8 -*-
"""46. 앞으로 해야 될 것 - 현장별로 무엇을 언제 해야 하는지 + 클로드가 「제가 만들까요?」 물어야 할 것. 토큰 0.
프로님 (2026-09-17) : "매일 아침 이걸로 보내. 서류 만들 게 있으면 나한테 만들라고 하지 말고 네가 만들까요 하고 물어봐.
                        내가 생각 못 해서 얘기를 못 할 수도 있는데 네가 챙겨서 이렇게 할까요 하고 물어봐야지. 체계를 바꾸라고."

  · 판단(무엇을 해야 하는지)은 클로드가 회의록을 읽고 받은답 `앞으로,현장,때,등급,할일,누가,왜,만들기` 로 넣는다.
  · 이 도구는 대장(_도구결과\\_대장\\앞으로할것.csv)을 읽어 보여 주고, 끝난 것을 번호로 완료 표시한다.
  · 42 아침 한 장 1층(오늘·지남)·2층(현장 카드)·클로드용 md, 35 아침 메일 제목에 매일 들어간다.
  · 「만들기」 칸이 찬 줄은 「제가 만들까요?」 로 뜬다 - 프로님께 만들라고 시키지 않는다.
  · 46번 : 목록 -> 「1,3」 완료 / 「a현장명」 그 현장만 / 엔터 = 나가기
  · 받은답 `앞으로완료,현장,할일 일부` 로도 완료된다.
"""
import os, sys, io
import common
from common import *
import facts

TOOL = '앞으로할것'
SUM_HEAD = ['번호', '때', 'D', '등급', '현장', '할 일', '누가', '왜', '제가 만들까요?', '상태', '적은 날']

def waiting(site=None):
    """답을 못 받은 제안 (확인 전) — 이것은 진행하지 않는다. 3일 넘으면 「묵은 제안」"""
    import datetime
    out = []
    for d in rows(site):
        if d['등급'] == '확정':
            continue
        gap = None
        try:
            y, m, dd = [int(x) for x in d['일자'].split('-')]
            gap = (today() - datetime.date(y, m, dd)).days
        except Exception:
            pass
        d['대기일'] = gap
        d['묵음'] = (gap is not None and gap >= 3)
        out.append(d)
    return out

def rows(site=None):
    out = []
    for d in facts.plan_load(site):
        dd = facts.plan_due(d)
        d['dd'] = dd
        d['level'] = 'red' if (dd is not None and dd <= 0) else ('yel' if (dd is not None and dd <= 3) else ('blu' if dd is not None else 'gry'))
        out.append(d)
    return out

def tag(d):
    dd = d.get('dd')
    if dd is None:
        return '미정' if '미정' in (d['때'] or '') else d['때']
    return '오늘' if dd == 0 else ('D%+d 지남' % dd if dd < 0 else 'D+%d' % dd)

def build_xlsx(xs=None, quiet=True):
    """요약 엑셀 (아침 메일 첨부용) : 1.앞으로 할 것 / 2.제가 만들까요 / 3.현장별"""
    xs = rows() if xs is None else xs
    od = outdir(TOOL)
    cp = write_csv(os.path.join(od, '앞으로할것_%s.csv' % ymd6()),
                   [[i, d['때'], tag(d), d['등급'], d['현장'], d['할일'], d['누가'], d['왜'], d['만들기'], d['상태'], d['일자']] for i, d in enumerate(xs, 1)], SUM_HEAD)
    xp = None
    try:
        import openpyxl
        from openpyxl.styles import Font, Alignment, PatternFill
        wb = openpyxl.Workbook()
        def sheet(name, head, body, widths, wrap=()):
            ws = wb.create_sheet(name)
            ws.append(head)
            for c in ws[1]:
                c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='2A6099')
            for r in body:
                ws.append(r)
            for i, w in enumerate(widths, 1):
                ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
            for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
                for c in row:
                    c.alignment = Alignment(vertical='top', wrap_text=(c.column in wrap))
            ws.freeze_panes = 'A2'
        body = [[i, d['때'], tag(d), d['등급'], d['현장'], d['할일'], d['누가'], d['왜'], d['만들기'], d['상태'], d['일자']] for i, d in enumerate(xs, 1)]
        sheet('1.앞으로 할 것', SUM_HEAD, body, [5, 12, 10, 7, 18, 52, 20, 56, 26, 6, 11], wrap=(6, 8))
        offers = [d for d in xs if d['만들기']]
        sheet('2.제가 만들까요', ['현장', '때', '만들 서류', '관련 할 일', '왜'],
              [[d['현장'], d['때'], d['만들기'], d['할일'], d['왜']] for d in offers], [18, 12, 30, 50, 56], wrap=(4, 5))
        by = {}
        for d in xs:
            by.setdefault(d['현장'], []).append(d)
        sheet('3.현장별', ['현장', '건수', '오늘·지남', '가장 급한 것'],
              [[s, len(v), len([d for d in v if d['level'] == 'red']), (v[0]['때'] + ' ' + v[0]['할일']) if v else ''] for s, v in by.items()],
              [18, 6, 10, 70], wrap=(4,))
        wb.remove(wb['Sheet'])
        xp = os.path.join(od, '앞으로할것_%s.xlsx' % ymd6())
        wb.save(xp)
        try:
            finish_xlsx(xp)
        except Exception:
            pass
    except Exception as e:
        if not quiet:
            print('[엑셀 생략] %s' % e)
    return xp, cp

def line(i, d):
    return '%2d. %-10s %-8s %-4s %-12s %s%s' % (i, (d['때'] or '')[:10], tag(d)[:8], d['등급'], (d['현장'] or '')[:12], d['할일'][:44],
                                              ('  [제가 만들까요? %s]' % d['만들기']) if d['만들기'] else '')

def run():
    title('46. 앞으로 해야 될 것   (현장별 · 때 순 · 「제가 만들까요?」 · 아침 한 장과 아침 메일에 매일)')
    site = None
    while True:
        xs = rows(site)
        if not xs:
            print('앞으로 해야 될 것이 비어 있습니다%s. 클로드가 회의록을 읽고 받은답 `앞으로,...` 로 넣습니다.' % ((' (' + site + ')') if site else ''))
            print('파일 : %s' % facts.plan_path())
            return True
        xp, cp = build_xlsx(xs, quiet=True)
        print('앞으로 해야 될 것 %d건%s  (오늘·지남 %d · 제가 만들까요 %d)' % (
            len(xs), (' · ' + site) if site else '', len([d for d in xs if d['level'] == 'red']), len([d for d in xs if d['만들기']])))
        print('요약 엑셀 : %s' % (xp or cp))
        print('-' * 100)
        for i, d in enumerate(xs, 1):
            print(line(i, d))
        print('-' * 100)
        print('끝난 번호를 넣으십시오.  예) 1,3   /   a현장명 = 그 현장만   /   엔터 = 나가기')
        s = ask('> ', '').strip()
        if not s:
            return True
        if s.startswith('a') and len(s) > 1:
            site = s[1:].strip(); continue
        n = 0
        for tok in s.replace(' ', '').split(','):
            if tok.isdigit() and 1 <= int(tok) <= len(xs):
                d = xs[int(tok) - 1]
                n += facts.plan_done(d['현장'], d['할일'])
        print('완료 %d건 표시했습니다. 다음 아침 한 장부터 빠집니다.' % n)
        try:
            if not common.AUTO and ask('아침 한 장을 지금 다시 만들까요? (y=엔터 / n) > ', 'y').lower().startswith('y'):
                import t42_morning as MO
                MO.run(quiet=True)
        except Exception as e:
            print('[42 실패] %s' % e)
        return True

if __name__ == '__main__':
    run(); pause()
