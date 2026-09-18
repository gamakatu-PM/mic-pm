# -*- coding: utf-8 -*-
"""47. 견적 보낸 곳 · 찾아갈 곳 - 견적을 보내 놓고 안 가보면 그대로 식는다. 토큰 0.
프로님 (2026-09-18) : "견적을 보낸 곳을 찾아가 봐야 된다고, 방문해야 할 것을 메일에 적게 해."

  · 견적을 보내면 대장에 한 줄 (받은답 `견적발송,현장,받는곳,담당자,금액,메모` / 47번 / csv 직접)
  · 보낸 뒤 3일이 지나도록 안 가보셨으면 → 42 아침 한 장 1층·35 아침 메일에 **매일** 빨갛게 뜬다
  · 다녀오시면 47번에서 번호만 넣으시면 됩니다 (받은답 `방문,현장,메모` 도 같은 효과)
  · 다녀온 뒤 14일이 지나면 「다시 가보실 때」 로 다시 뜬다
  · 수주/탈락이 정해지면 상태칸에 「수주」·「탈락」 을 넣으시면 목록에서 빠진다
대장 : _도구결과\\_대장\\견적발송.csv  (보낸날·현장·받는곳·담당자·금액·마지막방문·방문횟수·상태·메모)
"""
import os, sys, io
import common
from common import *
import facts

TOOL = '찾아갈곳'
SUM_HEAD = ['번호', '찾아갈 때', '보낸날', '지난날', '현장', '받는곳', '담당자', '금액', '마지막방문', '방문횟수', '왜', '메모']

def rows():
    return facts.quote_load()

def tag(d):
    if d['need']:
        return '가보실 때' if d['마지막방문'] else '★ 아직 안 감'
    return '다녀옴'

def build_xlsx(xs=None, quiet=True):
    xs = rows() if xs is None else xs
    od = outdir(TOOL)
    body = [[i, tag(d), d['보낸날'], (d['gap'] if d['gap'] is not None else ''), d['현장'], d['받는곳'], d['담당자'],
             d['금액'], d['마지막방문'], d['방문횟수'], d['why'], d['메모']] for i, d in enumerate(xs, 1)]
    cp = write_csv(os.path.join(od, '찾아갈곳_%s.csv' % ymd6()), body, SUM_HEAD)
    xp = None
    try:
        import openpyxl
        from openpyxl.styles import Font, Alignment, PatternFill
        wb = openpyxl.Workbook()
        ws = wb.active; ws.title = '1.찾아갈 곳'
        ws.append(SUM_HEAD)
        for c in ws[1]:
            c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='2A6099')
        for r in body:
            ws.append(r)
        for i, w in enumerate([5, 12, 11, 8, 18, 20, 14, 14, 11, 8, 34, 30], 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            for c in row:
                c.alignment = Alignment(vertical='top', wrap_text=(c.column in (11, 12)))
            if str(row[1].value or '').startswith('★'):
                for c in row:
                    c.fill = PatternFill('solid', fgColor='FDE8E6')
        ws.freeze_panes = 'A2'
        xp = os.path.join(od, '찾아갈곳_%s.xlsx' % ymd6())
        wb.save(xp)
        try:
            finish_xlsx(xp)
        except Exception:
            pass
    except Exception as e:
        if not quiet:
            print('[엑셀 생략] %s' % e)
    return xp, cp

def line(i, d):
    return '%2d. %-11s %-11s %-16s %-18s %-12s %s' % (
        i, tag(d), d['보낸날'], (d['현장'] or '')[:16], (d['받는곳'] or '')[:18], (d['담당자'] or '')[:12], d['why'])

def run():
    title('47. 견적 보낸 곳 · 찾아갈 곳   (보내 놓고 안 가면 식습니다. 아침 한 장·아침 메일에 매일)')
    while True:
        xs = rows()
        if not xs:
            print('견적 보낸 곳 대장이 비어 있습니다.')
            print('파일 : %s' % facts.quote_path())
            print('넣는 길 : 이 파일에 직접 / 받은답 `견적발송,현장,받는곳,담당자,금액,메모` / 아래에서 바로 넣기')
            if common.AUTO:
                return True
            if ask('지금 한 줄 넣을까요? (y=엔터 / n) > ', 'y').lower().startswith('y'):
                s = ask('현장 > ', '').strip()
                if not s:
                    return True
                facts.quote_add(s, ask('받는곳 (회사) > ', '').strip(), ask('담당자 > ', '').strip(),
                                ask('금액 (모르면 엔터) > ', '').strip(), ask('메모 > ', '').strip())
                continue
            return True
        need = [d for d in xs if d['need']]
        xp, cp = build_xlsx(xs, quiet=True)
        print('견적 보낸 곳 %d곳 · 찾아가실 곳 %d곳' % (len(xs), len(need)))
        print('목록 : %s' % (xp or cp))
        print('-' * 104)
        for i, d in enumerate(xs, 1):
            print(line(i, d))
        print('-' * 104)
        print('다녀오신 번호를 넣으십시오.  예) 1,3   /   n = 새로 한 줄 넣기   /   엔터 = 나가기')
        s = ask('> ', '').strip()
        if not s:
            return True
        if s.lower() == 'n':
            site = ask('현장 > ', '').strip()
            if site:
                facts.quote_add(site, ask('받는곳 (회사) > ', '').strip(), ask('담당자 > ', '').strip(),
                                ask('금액 (모르면 엔터) > ', '').strip(), ask('메모 > ', '').strip())
            continue
        n = 0
        for tok in s.replace(' ', '').split(','):
            if tok.isdigit() and 1 <= int(tok) <= len(xs):
                d = xs[int(tok) - 1]
                n += facts.quote_visit(d['현장'], to=d['받는곳'])
        print('%d곳 다녀오신 것으로 표시했습니다. 14일 뒤에 다시 뜹니다.' % n)
        try:
            if not common.AUTO and ask('아침 한 장을 지금 다시 만들까요? (y=엔터 / n) > ', 'y').lower().startswith('y'):
                import t42_morning as MO
                MO.run(quiet=True)
        except Exception as e:
            print('[42 실패] %s' % e)
        return True

if __name__ == '__main__':
    run(); pause()
