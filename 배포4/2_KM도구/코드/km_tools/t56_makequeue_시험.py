# -*- coding: utf-8 -*-
"""56. t56_makequeue 자가시험 — 「만들 차례」 탭 값 → 실제 파일·시트 준비물."""
from __future__ import print_function
import os, sys, io, json, shutil, tempfile

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import t56_makequeue as T

OK = 0
NG = []


def chk(name, cond, why=''):
    global OK
    if cond:
        OK += 1
        print('  OK   ' + name)
    else:
        NG.append(name)
        print('  NG   ' + name + '  ' + str(why))


HEAD = ['날짜', '현장', '할일', '완료', '고르기', '메모', '현장(걸러보기)', '원본탭', '처리일', '생성결과']

PAYLOAD = {'ok': True, 'valueRanges': [{'range': "'만들 차례'!A1:J10", 'values': [
    HEAD,
    ['2026-09-27', '수유초등학교', '- 화재 감지기 재설치', '', '작업의뢰서', '', '수유초등학교', '답요청', '2026-09-28', ''],
    ['2026-09-28', '양양 쏠비치', '- 견적 재확인 메일 보내기', '', '메일', '', '양양 쏠비치', '오늘 할일', '2026-09-28', ''],
    ['2026-09-28', '단양디캠프', '- 이번 주 진행 상황 정리', '', '보고서', '', '단양디캠프', '오늘 할일', '2026-09-28', ''],
    ['2026-09-28', '앵커 호텔', '- 제안 범위 넓혀서 다시', '', '제안서', '', '앵커 호텔', '오늘 할일', '2026-09-28', ''],
    ['2026-10-05', '조선호텔', '- 화재감지기 재설치', '', '캘린더', '', '조선호텔', '오늘 할일', '2026-09-28', ''],
    ['2026-09-27', '연합기숙사', '- 이미 만든 것', '', '메일', '', '연합기숙사', '답요청', '2026-09-28', '만듦(시트) · 메일 탭'],  # 이미 만듦(J열 있음) → 건너뜀
]}]}

print('--- parse : 대상 골라내기')
rows = T.rows_of(PAYLOAD)
items = T.parse(rows)
chk('5건만 골라냄 (이미 만든 줄 제외)', len(items) == 5, len(items))
chk('행 번호 = 인덱스+2 (헤더 다음이 2행)', items[0]['row'] == 2, items[0])
chk('할일 글 앞 「- 」 는 뗀다', items[0]['task'] == '화재 감지기 재설치', items[0]['task'])
chk('제안서도 목록엔 들어감(run 에서 건너뜀)', any(i['pick'] == '제안서' for i in items))
chk('캘린더도 목록엔 들어감(run 에서 이 스크립트가 만들지 않고 캘린더항목 으로만 넘김)', any(i['pick'] == '캘린더' for i in items))

print('--- run : 실제 파일·시트준비물 생성')
tmp = tempfile.mkdtemp(prefix='km_t56_')
try:
    rep = T.run(PAYLOAD, out_root=tmp, today='260928')
    chk('결과에 tab 키 있음 (그대로 markqueue 에 먹일 수 있게)', rep.get('tab') == '만들 차례', rep.get('tab'))
    chk('제안서는 대기 1건으로만 세고 안 만듦', rep['제안서대기'] == 1, rep['제안서대기'])
    chk('made 는 작업의뢰서·메일·보고서 3건(캘린더는 made 에 안 들어감)', len(rep['made']) == 3, rep['made'])
    chk('marks 는 작업의뢰서 1건뿐(메일·보고서는 appenddoc 성공 뒤에만 표시되므로 여기 안 넣음)', len(rep['marks']) == 1, rep['marks'])

    wo = [m for m in rep['made'] if m['pick'] == '작업의뢰서'][0]
    mail_made = [m for m in rep['made'] if m['pick'] == '메일'][0]
    rpt_made = [m for m in rep['made'] if m['pick'] == '보고서'][0]
    chk('작업의뢰서 파일이 실제로 생김', os.path.isfile(wo['path']), wo)
    chk('메일은 made 에 상태만(경로 없음, 아직 시트에 안 들어감)', mail_made['path'] == '' and '시트' in mail_made['state'], mail_made)
    chk('보고서도 made 에 상태만', rpt_made['path'] == '' and '시트' in rpt_made['state'], rpt_made)

    print('--- 캘린더항목 : 이 스크립트는 만들지 않고 그대로 넘긴다(Routine 이 구글캘린더에 직접 등록)')
    cal = rep['캘린더항목']
    chk('캘린더항목 1건', len(cal) == 1, cal)
    chk('캘린더항목 : 현장·할일·기한(due) 이 시트 글 그대로', cal[0]['site'] == '조선호텔' and cal[0]['task'] == '화재감지기 재설치' and cal[0]['due'] == '2026-10-05', cal[0])
    chk('캘린더항목에 row 번호도 있음(성공하면 그 번호로 markqueue)', cal[0]['row'] == 6, cal[0])

    print('--- 메일추가·보고서추가 : appenddoc 에 그대로 먹일 수 있는 모양')
    ma = rep['메일추가']
    ra = rep['보고서추가']
    chk('메일추가 : tab=메일 · markTab=만들 차례', ma and ma['tab'] == '메일' and ma['markTab'] == '만들 차례', ma)
    chk('메일추가 : rows 1건, 날짜|현장|할일|본문 4칸', ma and len(ma['rows']) == 1 and len(ma['rows'][0]) == 4, ma)
    chk('메일추가 : 현장·할일 글이 그대로 있음(지어내지 않음)', ma and ma['rows'][0][1] == '양양 쏠비치' and ma['rows'][0][2] == '견적 재확인 메일 보내기', ma)
    chk('메일추가 : 본문이 범용·정중체(담당·회사 특정 안 함)', ma and ma['rows'][0][3].startswith('안녕하십니까,') and '견적 재확인 메일 보내기' in ma['rows'][0][3], ma)
    chk('메일추가 : marks 는 원래 행 번호(3행)로 미리 준비됨(appenddoc 이 성공할 때만 실제로 쓰임)', ma and ma['marks'] == [{'row': 3, 'text': '만듦(시트) · 메일 탭'}], ma)

    chk('보고서추가 : tab=보고서', ra and ra['tab'] == '보고서', ra)
    chk('보고서추가 : 현장·할일·일자가 그대로 있음', ra and ra['rows'][0][1] == '단양디캠프' and '이번 주 진행 상황 정리' in ra['rows'][0][3], ra)

    if rep['template'] and os.path.isfile(wo['path']):
        import openpyxl
        wb = openpyxl.load_workbook(wo['path'])
        ws = wb.worksheets[0]
        chk('작업의뢰서 : 현장(E12)이 채워짐', ws['E12'].value == '수유초등학교', ws['E12'].value)
        chk('작업의뢰서 : G4 = 배성윤', ws['G4'].value == '배성윤', ws['G4'].value)
        chk('작업의뢰서 : 본문 B19 에 할일 글', '화재 감지기' in str(ws['B19'].value or '') or '화재 감지기' in str(ws['B20'].value or ''), ws['B19'].value)
        chk('작업의뢰서 : O4(납기일) 는 비워 둠(차장님 몫)', ws['O4'].value in (None, ''), ws['O4'].value)
    else:
        print('  (참고) 원틀을 못 찾아 작업의뢰서 xlsx 자체 시험은 건너뜀 — 파일 있으면 위에서 이미 통과')

    print('--- 두 번째 실행(같은 입력) : 작업의뢰서만 「이미 있음」, 메일·보고서·캘린더항목은 판단 없이 그대로 다시 준비물을 냄')
    rep2 = T.run(PAYLOAD, out_root=tmp, today='260928')
    wo2 = [m for m in rep2['made'] if m['pick'] == '작업의뢰서'][0]
    chk('두 번째 : 작업의뢰서는 「이미 있음」(파일 안 덮어씀)', wo2['state'] in ('이미 있음', '원틀 없음 — 만들지 않음'), wo2)
    chk('두 번째 : 메일추가·보고서추가 도 다시 준비됨(실제 중복 방지는 시트 J열이 함 — 이 스크립트는 판단 안 함)', rep2['메일추가'] and rep2['보고서추가'], (rep2['메일추가'], rep2['보고서추가']))
finally:
    shutil.rmtree(tmp, ignore_errors=True)

print('--- 메일추가·보고서추가 가 비어 있을 때는 None(appenddoc 을 부를 필요 없다는 신호)')
tmp2 = tempfile.mkdtemp(prefix='km_t56_')
try:
    only_wo = {'ok': True, 'valueRanges': [{'range': "'만들 차례'!A1:J10", 'values': [
        HEAD, ['2026-09-28', '수유초등학교', '- 재설치', '', '작업의뢰서', '', '수유초등학교', '답요청', '2026-09-28', '']]}]}
    rep3 = T.run(only_wo, out_root=tmp2, today='260928')
    chk('메일추가 없으면 None', rep3['메일추가'] is None, rep3['메일추가'])
    chk('보고서추가 없으면 None', rep3['보고서추가'] is None, rep3['보고서추가'])
    chk('캘린더항목 없으면 빈 리스트', rep3['캘린더항목'] == [], rep3['캘린더항목'])
finally:
    shutil.rmtree(tmp2, ignore_errors=True)

print('\n합계 %d개 중 통과 %d · 실패 %d' % (OK + len(NG), OK, len(NG)))
if NG:
    print('실패 : ' + ', '.join(NG))
    sys.exit(1)
