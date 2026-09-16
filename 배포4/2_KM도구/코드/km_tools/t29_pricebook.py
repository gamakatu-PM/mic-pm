# -*- coding: utf-8 -*-
"""29. 단가장 채우기 - 배선도에서 CB 내부 모듈을 뽑아 단가장에 없는 줄을 찾아 넣는다. 토큰 0.

파이썬이 하는 것 (0원)
  1) 배선도(DXF/PDF)에서 모듈 형번과 수량을 뽑는다
  2) 단가장과 대조해 「단가장에 없는 모듈」을 찾아낸다
  3) 단가장 새 버전을 만들어 그 모듈을 주황색 줄로 추가한다 (실행가는 빈칸)
     견적가·예산가는 배수 수식으로 채워 둔다
  4) CB구성.csv 후보를 만들어 28번이 바로 쓰게 한다

파이썬이 못 하는 것
  * 단가(금액) 자체는 만들 수 없습니다. 매입가·개발원가는 사람이 정하는 값입니다.
    그래서 실행가 칸은 반드시 비워 두고 주황색으로 표시합니다.
  * 스캔·사진 배선도는 글자가 없어 못 읽습니다.
"""
import os, re, csv, shutil, collections
from common import *
import t27_drawing as D
import t28_cost as C

TOOL = '단가장채우기'

# CB 내부에서 실제로 쓰이는 낱말 (이게 든 글자만 모듈 후보로 본다)
MODWORD = ('MCB', 'RCBO', 'ELB', 'MCCB', 'TRANS', 'SMPS', 'SSR', 'EOCR',
           'RELAY', '릴레이', 'MAIN', 'EXIO', 'EX', 'SUB', '485', 'LAN',
           'HMS', 'RIU', 'FV', 'SA', 'TL', 'CB-', 'BCB', '퓨즈', 'FUSE',
           'DIMMER', 'LATCH', 'SCOM', 'CAN', 'AIRC', 'SPEAKER', 'CHIME',
           'POWER', 'CONVERT', 'DTC', 'DCU', 'FIP', 'BOARD', '보드', '단자')
CODE = re.compile(r'[A-Z]{2,}[-_]?[0-9]{2,}[A-Z0-9_]*|[0-9]{2,3}A[FT]')

# MAIN 보드의 커넥터·회로 이름들. 모듈이 아니므로 뺀다 (모듈무시.csv 로 늘리실 수 있습니다)
MODSKIP = ('DIMMER', 'LATCH', 'SCOM', 'RELAY5', 'RELAY8', 'RELAY5/8', 'CAN1', 'CAN2',
           'AIRC', 'AIRC1', 'AIRC2', '485', '12V', 'SP', 'IND', 'KS', 'DM', 'EXIO',
           'R1', 'R2', 'R3', 'R4', 'R5', 'A1', 'A2', 'V01', 'SW', 'GND', 'AC', 'DC')
SKIP_F = '모듈무시.csv'
SKIP_DEFAULT = [['# 모듈이 아닌 낱말. 여기 적으면 29번이 셈에서 뺍니다.', ''],
                ['무시할낱말', '비고'],
                ['DIMMER1', 'MAIN 보드 커넥터'],
                ['RELAY5/8', 'MAIN 보드 커넥터'],
                ['LATCH', 'MAIN 보드 커넥터'],
                ['SCOM1~4', 'MAIN 보드 커넥터'],
                ['LATCH 485', 'MAIN 보드 커넥터']]

def _skips():
    out = set(x.upper() for x in MODSKIP)
    for r in C.rows_of(SKIP_F):
        if r and r[0] and '무시' not in r[0]:
            out.add(r[0].strip().upper())
    return out

def is_mod(s, skips=None):
    u = re.sub(r'\s+', ' ', str(s or '')).strip().upper()
    if not u:
        return False
    if skips is None:
        skips = _skips()
    bare = re.sub(r'[^0-9A-Z가-힣/~]', '', u)
    if bare in skips or u in skips:
        return False
    return any(w in u for w in MODWORD)

# ---------------- 배선도 읽기 ----------------

def pick_files(folder=None):
    """도면 폴더에서 배선도로 보이는 것을 먼저 고른다."""
    root = folder or D.dwg_root()
    files = D.gather(folder) if folder else (D.gather(os.path.join(root, D.INBOX)) or D.gather(root))
    if not files:
        print('[도면이 없습니다] %s 에 배선도를 넣어주십시오.' % os.path.join(root, D.INBOX))
        open_folder(os.path.join(root, D.INBOX))
        return []
    want = [f for f in files if any(k in os.path.basename(f) for k in ('배선', '결선', 'CB', '계통'))]
    use = want or files
    print('본 곳 : %s' % root)
    if want:
        print('배선도로 보이는 파일 %d개를 먼저 봅니다.' % len(want))
    return use

def scan(files):
    """모듈 후보 : {이름: 나온 횟수}"""
    hit = collections.Counter()
    read = []
    skips = _skips()
    for f in files:
        if D.is_xref(f):
            continue
        e = os.path.splitext(f)[1].lower()
        nm = os.path.basename(f)
        texts, blocks = [], collections.Counter()
        if e in D.DXF_EXT:
            d, msg = D.read_dxf(f)
            if d:
                texts = d['texts']; blocks = d['blocks']
            read.append((nm, 'DXF', msg or ''))
        elif e in D.DWG_EXT:
            conv = D.dwg_to_dxf(f)
            if conv:
                d, msg = D.read_dxf(conv)
                if d:
                    texts = d['texts']; blocks = d['blocks']
                read.append((nm, 'DWG(변환)', ''))
            else:
                read.append((nm, 'DWG', '못 읽음 - DXF 로 주십시오'))
                continue
        elif e in D.PDF_EXT:
            d, msg = D.read_pdf(f)
            if d:
                texts = d['texts']
            for r in D.read_pdf_table(f):
                if is_mod(r['name'], skips) or is_mod(r['sym'], skips):
                    hit[r['name'].strip()] += r['qty']
            read.append((nm, 'PDF', msg or ''))
        else:
            read.append((nm, '사진/캡처', '파이썬으로는 못 읽습니다'))
            continue
        for b, n in blocks.items():
            if is_mod(b, skips):
                hit[str(b).strip()] += n
        for t in texts:
            s = re.sub(r'\\[A-Za-z][^;]*;', ' ', str(t or '')).strip()
            if not s or len(s) > 60:
                continue
            if is_mod(s, skips) and (CODE.search(s.upper()) or len(s) <= 24):
                hit[s] += 1
    return hit, read

# ---------------- 단가장 대조 ----------------

def in_pricebook(name, pb):
    mod, cost, why, cands = C.find_price(name, pb, [], '')
    return (mod, cost) if cost else (None, None)

def next_version(path):
    """CB모듈_단가장_v70.xlsx -> CB모듈_단가장_v71_260916.xlsx (덮어쓰지 않는다)"""
    d, b = os.path.split(path)
    stem, ext = os.path.splitext(b)
    m = re.search(r'v(\d+)', stem, re.I)
    if m:
        stem = stem[:m.start()] + 'v%d' % (int(m.group(1)) + 1) + stem[m.end():]
    else:
        stem += '_v2'
    return os.path.join(d, '%s_%s%s' % (stem, ymd6(), ext))

ORANGE = 'FFE0B2'

def add_rows_xlsx(src, rows, mult):
    """단가장 새 버전을 만들고 총괄 시트 끝에 주황색 줄을 넣는다. 실행가는 빈칸."""
    try:
        import openpyxl
        from openpyxl.styles import PatternFill, Font
    except ImportError:
        print('[부품 없음] openpyxl 이 없어 엑셀에 못 넣었습니다.')
        return None
    dst = next_version(src)
    shutil.copy(src, dst)
    wb = openpyxl.load_workbook(dst)
    names = sorted(wb.sheetnames, key=lambda n: ('총괄' not in n, n))
    ws = wb[names[0]]
    # 머리글 찾기
    hrow, ci = None, {}
    for r in range(1, min(ws.max_row, 30) + 1):
        cells = [('' if ws.cell(r, c).value is None else str(ws.cell(r, c).value).strip())
                 for c in range(1, min(ws.max_column, 12) + 1)]
        joined = ' '.join(cells)
        if '실행' in joined and any(k in joined for k in ('모듈', '형번', '품명', '품목')):
            for i, v in enumerate(cells, start=1):
                if '구분' in v and 'grp' not in ci: ci['grp'] = i
                if any(k in v for k in ('모듈', '형번', '품명', '품목')) and 'mod' not in ci: ci['mod'] = i
                if '실행' in v and 'cost' not in ci: ci['cost'] = i
                if '견적' in v and 'q' not in ci: ci['q'] = i
                if '예산' in v and 'b' not in ci: ci['b'] = i
            hrow = r
            break
    if not hrow or 'mod' not in ci or 'cost' not in ci:
        print('[머리글을 못 찾았습니다] 총괄 시트에 구분/모듈명/실행가 머리글이 있어야 합니다.')
        return None
    r = ws.max_row + 1
    fill = PatternFill('solid', fgColor=ORANGE)
    qm = mult.get('견적배수'); bm = mult.get('예산배수')
    for name, qty in rows:
        if 'grp' in ci: ws.cell(r, ci['grp'], '신규(확인)')
        ws.cell(r, ci['mod'], name)
        ws.cell(r, ci['cost'], None)                    # 단가는 사람이 넣는다
        cl = ws.cell(r, ci['cost']).column_letter
        if 'q' in ci and qm:
            ws.cell(r, ci['q'], '=IF(%s%d="","",ROUND(%s%d*%s,0))' % (cl, r, cl, r, qm))
        if 'b' in ci and bm:
            ws.cell(r, ci['b'], '=IF(%s%d="","",ROUND(%s%d*%s,0))' % (cl, r, cl, r, bm))
        last = max(ci.values())
        for c in range(1, last + 1):
            ws.cell(r, c).fill = fill
        note = ws.cell(r, last + 1, '배선도에서 %s회 나옴 · %s 추가 · 실행가 입력 필요' % (won(qty), today().isoformat()))
        note.font = Font(name='맑은 고딕', size=9)
        r += 1
    wb.save(dst)
    return dst

# ---------------- 실행 ----------------

def run(folder=None, site_hint=None):
    title('29. 단가장 채우기   (배선도 -> 없는 모듈 찾아 단가장에 줄 추가. 토큰 0)')
    pb = C.load_pricebook()
    if not pb:
        print('')
        print('[단가장이 없습니다] 3_공통사용\\단가장\\ 에 단가장 파일을 먼저 넣어주십시오.')
        print('  엑셀 그대로 됩니다. 이름에 「단가장」이나 「CB모듈」이 들어가면 찾습니다.')
        open_folder(C.root())
        return
    src = C._find_pb_file()
    C.seed(C.MULT, C.MULT_DEFAULT); C.seed(C.CBC, C.CBC_DEFAULT); C.seed(SKIP_F, SKIP_DEFAULT)
    mult = C.load_mult()
    print('단가장 : %s  (%d줄)' % (os.path.basename(src), len(pb)))
    print('-' * 74)

    files = pick_files(folder)
    if not files:
        return
    print('')
    hit, read = scan(files)
    for nm, kind, msg in read:
        print(' [%s] %-40s %s' % (kind, nm[:40], msg))
    if not hit:
        print('')
        print('[모듈로 보이는 글자를 못 찾았습니다]')
        print(' 배선도가 DXF 면 블록·글자를 다 읽습니다. PDF 는 글자가 살아 있어야 합니다.')
        print(' 사진·스캔이면 저(클로드)에게 그 파일을 주십시오.')
        return

    site = ask('\n현장명 > ', site_hint or D.guess_site(files, ''))
    print('')
    print('%-46s %7s  %s' % ('배선도에서 뽑은 모듈 후보', '횟수', '단가장'))
    print('-' * 78)
    have, none_ = [], []
    for name, n in hit.most_common(80):
        mod, cost = in_pricebook(name, pb)
        if cost:
            print('%-46s %7s  있음  %s (%s)' % (str(name)[:46], won(n), str(mod)[:24], won(cost)))
            have.append([name, n, mod, int(cost)])
        else:
            print('%-46s %7s  ★없음' % (str(name)[:46], won(n)))
            none_.append([name, n])

    od = outdir(TOOL)
    tag = '%s_%s' % (safe_name(site), ymd6())
    made = []
    made.append(write_csv(os.path.join(od, '%s_배선도모듈.csv' % tag),
                          [[a, b, c, d] for a, b, c, d in have] + [[a, b, '', ''] for a, b in none_],
                          ['배선도에서 뽑은 이름', '나온 횟수', '단가장에서 맞춘 이름', '실행가']))
    if none_:
        made.append(write_csv(os.path.join(od, '%s_단가장에없는모듈.csv' % tag), none_,
                              ['모듈 이름', '배선도에 나온 횟수']))
    # CB구성 후보
    cbc = [['확인', a, 1] for a, b in
           [(x[0], x[1]) for x in have] + none_]
    made.append(write_csv(os.path.join(od, '%s_CB구성_후보.csv' % tag), cbc,
                          ['구분', '모듈명(형번)', 'CB1대당수량']))

    newpb = None
    if none_:
        print('')
        print('단가장에 없는 모듈이 %d개입니다.' % len(none_))
        if ask('단가장 새 버전을 만들어 주황색 줄로 넣을까요? (y=엔터 / n) > ', 'y').lower().startswith('y'):
            if os.path.splitext(src)[1].lower() in ('.xlsx', '.xlsm'):
                newpb = add_rows_xlsx(src, none_, mult)
                if newpb:
                    print('만들었습니다 : %s' % newpb)
                    print('  -> 주황색 줄의 실행가만 넣으시면 됩니다. 견적가·예산가는 수식으로 채워 뒀습니다.')
                    made.append(newpb)
            else:
                p = os.path.join(od, '%s_단가장_추가할줄.csv' % tag)
                made.append(write_csv(p, [['신규(확인)', a, '', '', ''] for a, b in none_],
                                      ['구분', '모듈명(형번)', '실행가', '견적가', '예산가']))
                print('단가장이 csv 라서 추가할 줄만 따로 뽑았습니다 : %s' % p)

    blocks = [('단가장에 없는 모듈 (실행가를 넣어 주셔야 합니다)',
               [('red', '%s : 배선도에 %s회' % (a, won(b))) for a, b in none_]),
              ('단가장에 이미 있는 모듈',
               [('green', '%s -> %s (%s)' % (a, c, won(d))) for a, b, c, d in have]),
              ('읽은 파일', [('gray', '%s [%s] %s' % (a, b, c)) for a, b, c in read]),
              ('파이썬이 못 하는 것',
               [('yellow', '단가(금액) 자체는 만들 수 없습니다. 매입가·개발원가는 사람이 정하는 값입니다.'),
                ('yellow', '그래서 실행가 칸은 비워 두고 주황색으로 표시합니다.')])]
    f9 = write_html(os.path.join(od, '%s_단가장채우기.html' % tag),
                    '%s 단가장 채우기' % site, blocks, files=made)
    print('')
    print('파일을 만들었습니다. 결과 화면을 지금 띄웁니다.')
    for x in made + [f9]:
        print('  %s' % x)
    log(TOOL, '%s 모듈%d 없음%d' % (site, len(have) + len(none_), len(none_)))
    if not open_file(f9):
        open_folder(od)

if __name__ == '__main__':
    run(); pause()
