# -*- coding: utf-8 -*-
"""39. 자가 시험 - 도구 전체 흐름(36 오늘 한 방에)을 시험자료(합성 도면 + 가짜 단가장)로 임시 폴더에서 돌려
기대값과 대조한다. 토큰 0. 클로드는 zip 을 만들기 전에 반드시 이것을 통과시켜야 한다 (프로님: "한 달 뒤 또 이상한 짓" 방지).
  · 임시 폴더에 시험자료를 복사 → 36번(quiet) 실행 → 결과 파일·숫자 검사 → 통과/실패 표
  · 정답본 폴더(_원틀\\정답본)가 있으면 37 검수까지 같이 본다
"""
import os, sys, shutil, tempfile, subprocess, re, glob
from common import *

TOOL = '자가시험'
DATA = os.path.join(os.path.dirname(os.path.dirname(HERE)), '시험자료')

EXPECT = {
    '앵커_품목수': 7, '앵커_ENTRANCE': 320, '앵커_L6': 39, '앵커_이전판건너뜀': 2,
    '쏠비치_품목수': 2,
    '회의변경_앵커': 2, '받는함_옮김': 1,
}

def _worker(base):
    """자식 프로세스 : 경로를 임시 폴더로 바꾸고 36번을 돌린다"""
    import common
    common.DEFAULTS.update({'base': base, 'template': os.path.join(base, '_원틀'), 'out': os.path.join(base, '_도구결과'),
                            'price': os.path.join(base, '3_공통사용', '단가장'), 'drawing': os.path.join(base, '3_공통사용', '도면'),
                            'handover': os.path.join(base, '인수인계함'), 'plaud': os.path.join(base, 'plaud'),
                            'biseo': os.path.join(base, '_현장비서'), 'meeting_inbox': os.path.join(base, '받는함')})
    common.AUTO = True
    common.open_file = lambda *a, **k: False
    common.open_folder = lambda *a, **k: False
    import t98_update as U
    U.fetch_latest = lambda quiet=True: None
    U.newer_zip = lambda: None
    import t36_today
    t36_today.open_file = lambda *a, **k: False
    t36_today.run(quiet=True)

def prepare():
    base = tempfile.mkdtemp(prefix='km_selftest_')
    for d in ('3_공통사용', '_원틀', '인수인계함', 'plaud'):
        os.makedirs(os.path.join(base, d), exist_ok=True)
    for d in ('plaud', '_현장비서', '받는함'):
        if os.path.isdir(os.path.join(DATA, d)):
            shutil.rmtree(os.path.join(base, d), ignore_errors=True)
            shutil.copytree(os.path.join(DATA, d), os.path.join(base, d))
    shutil.copytree(os.path.join(DATA, '도면'), os.path.join(base, '3_공통사용', '도면'))
    shutil.copytree(os.path.join(DATA, '단가장'), os.path.join(base, '3_공통사용', '단가장'))
    g = os.path.join(cfg('template'), '정답본')
    if os.path.isdir(g):
        shutil.copytree(g, os.path.join(base, '_원틀', '정답본'))
    return base

def _csv_rows(p):
    out = []
    for i, line in enumerate(read_text(p).splitlines()):
        if i == 0 or not line.strip(): continue
        try:
            import csv; out.append([c.strip() for c in next(csv.reader([line]))])
        except Exception: pass
    return out

def check(base, log):
    R = []
    def ok(name, cond, note=''):
        R.append((name, bool(cond), note))
    ok('오류 없이 끝남 (Traceback 없음)', 'Traceback' not in log and '오늘 한 방에 끝' in log)
    o = os.path.join(base, '_도구결과')
    q = glob.glob(os.path.join(o, '도면수량', '*', '앵커호텔_도면에적힌수량표_*.csv'))
    ok('앵커 NOTE 수량표 파일 생성', bool(q), q[0] if q else '없음')
    if q:
        rows = _csv_rows(q[-1]); d = {r[1]: int(float(r[2])) for r in rows if len(r) >= 3 and r[2]}
        ok('앵커 품목 수 = %d (최신 판만)' % EXPECT['앵커_품목수'], len(rows) == EXPECT['앵커_품목수'], '실제 %d' % len(rows))
        ok('ENTRANCE INDICATOR = %d (Rev2 값)' % EXPECT['앵커_ENTRANCE'], d.get('ENTRANCE INDICATOR') == EXPECT['앵커_ENTRANCE'], '실제 %s' % d.get('ENTRANCE INDICATOR'))
        ok('LIGHT SWITCH bath 6 = %d' % EXPECT['앵커_L6'], d.get('LIGHT SWITCH bath 6') == EXPECT['앵커_L6'], '실제 %s' % d.get('LIGHT SWITCH bath 6'))
    ok('이전 판 %d개 건너뜀 표시' % EXPECT['앵커_이전판건너뜀'], ('이전 판이라 읽지 않은 파일 %d개' % EXPECT['앵커_이전판건너뜀']) in log)
    q2 = glob.glob(os.path.join(o, '도면수량', '*', '쏠비치양양_도면에적힌수량표_*.csv'))
    ok('쏠비치 수량표 %d품목' % EXPECT['쏠비치_품목수'], bool(q2) and len(_csv_rows(q2[-1])) == EXPECT['쏠비치_품목수'])
    # v33 회의 연결 (40)
    ch = glob.glob(os.path.join(o, '회의연결', '*', '앵커호텔_회의변경수량_*.csv'))
    ok('회의록 수량·규격 변경 앵커 %d줄' % EXPECT['회의변경_앵커'], bool(ch) and len(_csv_rows(ch[-1])) == EXPECT['회의변경_앵커'],
       '실제 %s' % (len(_csv_rows(ch[-1])) if ch else '파일 없음'))
    req = glob.glob(os.path.join(o, '완성품', '*', '_클로드부탁서_앵커호텔_*.md'))
    rq = read_text(req[-1]) if req else ''
    ok('부탁서에 「회의에서 바뀐 수량·규격」 절 + 330', ('회의에서 바뀐 수량' in rq) and ('330' in rq))
    mv = glob.glob(os.path.join(base, '_현장비서', '1.여기에_v10결과_넣기', '*.txt'))
    ok('받는함 txt %d개 -> _현장비서 대기함' % EXPECT['받는함_옮김'], len(mv) == EXPECT['받는함_옮김'], '실제 %d' % len(mv))
    bk = glob.glob(os.path.join(base, '인수인계함', '회의록코드_백업', '*', '코드', 'km_run.py'))
    ok('회의록 코드 백업 생성', bool(bk))
    dash = glob.glob(os.path.join(base, '인수인계함', '_현황판.md'))
    ok('현황판 ⑦ 회의 변경 블록', bool(dash) and '⑦ 회의에서 바뀐 수량' in read_text(dash[-1]) and '330' in read_text(dash[-1]))
    x = glob.glob(os.path.join(o, '단가붙이기', '*', '앵커호텔_*_실행산출_v1.xlsx'))
    ok('앵커 실행산출 xlsx 생성', bool(x))
    if x:
        try:
            import openpyxl
            wb = openpyxl.load_workbook(x[-1], data_only=True); ws = wb['1.입력판']
            cb = next((ws.cell(r, 2).value for r in range(1, 40) if str(ws.cell(r, 1).value or '').startswith('CONTROL BOX')), None)
            ok('CB 대수 기본값 = 320 (실당 품목 수량)', cb == 320, '실제 %s' % cb)
            ws3 = wb['3.견적↔실행 대조']
            tot = next((ws3.cell(r, 7).value for r in range(1, 60) if str(ws3.cell(r, 2).value or '').strip() == 'TOTAL'), None)
            ok('실행 TOTAL 값이 파일 안에 들어 있음', isinstance(tot, (int, float)) and tot > 0, '실제 %s' % tot)
        except Exception as e:
            ok('실행산출 열기', False, str(e))
    k = glob.glob(os.path.join(o, '단가붙이기', '*', '앵커호텔_*_실행견적_기구물.csv'))
    if k:
        rows = _csv_rows(k[-1]); priced = [r for r in rows if len(r) >= 4 and r[3]]
        ok('기구물 단가 매칭 6줄 이상 (가짜 단가장)', len(priced) >= 6, '실제 %d/%d' % (len(priced), len(rows)))
    ok('앵커 견적서 xlsx 생성', bool(glob.glob(os.path.join(o, '단가붙이기', '*', '앵커호텔_*_견적서_v1.xlsx'))))
    ok('부탁서 생성', bool(glob.glob(os.path.join(o, '완성품', '*', '_클로드부탁서_앵커호텔_*.md'))))
    ok('현황판 생성', os.path.exists(os.path.join(base, '_현황판.html')))
    if glob.glob(os.path.join(base, '_원틀', '정답본', '*.xlsx')):
        v = re.findall(r'대상\s*:\s*앵커호텔_\d+_실행산출_v1\.xlsx.*?검수 판정\s*:\s*(\S+)', log, re.S)
        ok('37 검수(실행산출) 「잘못」 아님', bool(v) and v[-1] != '잘못', '판정 %s' % (v[-1] if v else '없음'))
        v2 = re.findall(r'대상\s*:\s*앵커호텔_\d+_견적서_v1\.xlsx.*?검수 판정\s*:\s*(\S+)', log, re.S)
        ok('37 검수(견적서) 「잘못」 아님', bool(v2) and v2[-1] != '잘못', '판정 %s' % (v2[-1] if v2 else '없음'))
    else:
        ok('정답본 없음 → 37 검수 생략 (PC 에 정답본을 넣으면 같이 봅니다)', True)
    return R

def run(keep=False):
    title('39. 자가 시험   (시험자료로 36번 전체 흐름을 임시 폴더에서 돌려 기대값과 대조. 토큰 0)')
    if not os.path.isdir(DATA):
        print('[시험자료 없음] %s' % DATA); return False
    base = prepare()
    print('임시 폴더 : %s' % base)
    print('36번을 조용히 돌리는 중... (1분 안팎)')
    r = subprocess.run([sys.executable, os.path.abspath(__file__), '--worker', base], capture_output=True, text=True,
                       encoding='utf-8', errors='replace', timeout=600)
    log = (r.stdout or '') + (r.stderr or '')
    od = outdir(TOOL)
    lp = os.path.join(od, '자가시험_로그_%s.txt' % ymd6())
    import io as _io
    _io.open(lp, 'w', encoding='utf-8').write(log)
    R = check(base, log)
    print('')
    print('%-46s %s' % ('검사', '결과'))
    print('-' * 70)
    bad = 0
    for name, good, note in R:
        print('%-46s %s  %s' % (name[:46], '통과' if good else '실패', note if not good else ''))
        bad += 0 if good else 1
    print('-' * 70)
    verdict = '전부 통과 (%d/%d)' % (len(R) - bad, len(R)) if not bad else '실패 %d개 / %d' % (bad, len(R))
    print(verdict)
    rp = write_csv(os.path.join(od, '자가시험_%s.csv' % ymd6()), [[n, '통과' if g else '실패', t] for n, g, t in R], ['검사', '결과', '비고'])
    print('보고 : %s / 로그 : %s' % (rp, lp))
    log_line = '%s' % verdict
    log(TOOL, log_line) if False else None
    try:
        from common import log as _log; _log(TOOL, verdict)
    except Exception:
        pass
    if not keep:
        shutil.rmtree(base, ignore_errors=True)
    return bad == 0

if __name__ == '__main__':
    if '--worker' in sys.argv:
        _worker(sys.argv[sys.argv.index('--worker') + 1])
    else:
        run(); pause()
