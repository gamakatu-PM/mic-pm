# -*- coding: utf-8 -*-
"""56번 수금대장 자동 채움 시험. 가짜 현장·가짜 금액 (실제 아님). 임시 폴더에서만 돈다.
    python t56_collect_auto_시험.py
"""
from __future__ import print_function
import os, sys, io, json, csv, shutil, tempfile, datetime

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import common
tmp = tempfile.mkdtemp(prefix='t56_')
os.makedirs(os.path.join(tmp, '_도구결과'))
common.DEFAULTS.update({'base': tmp, 'out': os.path.join(tmp, '_도구결과'), 'template': os.path.join(tmp, '_원틀')})
common.AUTO = True
import facts as FX
import t56_collect_auto as T

R = []


def ok(name, cond, note=''):
    R.append((name, bool(cond), note))
    print('  %s %s %s' % ('OK  ' if cond else 'FAIL', name, '' if cond else note))


def rows(p):
    return T._read_csv(p)


def put_meta(d, fn, site, sched=(), todo=()):
    j = {'meta': {'site': site, 'ymd': '260925', 'hm': '10:00', 'person': '홍길동 소장'},
         'sec': {'1': {'현장': site, '일정': list(sched)}, '2': {'할 일': list(todo)}}}
    io.open(os.path.join(d, fn), 'w', encoding='utf-8').write(json.dumps(j, ensure_ascii=False))


def make_quote(site):
    import openpyxl
    d = os.path.join(common.cfg('out'), '단가붙이기', '260925'); os.makedirs(d, exist_ok=True)
    p = os.path.join(d, '%s_260925_견적서_v1.xlsx' % site)
    wb = openpyxl.Workbook(); g = wb.active; g.title = '갑지 '; ws = wb.create_sheet('내역서 ')
    ws['A4'] = ' 1. 중앙관리 시스템'
    ws['A5'] = 'PC 서버'; ws['K5'] = 1000000
    ws['A6'] = '소프트웨어'; ws['K6'] = 900000       # 감사 : 「소」 로 시작하는 품목이 빠지던 것
    ws['A14'] = '소   계'
    ws['A7'] = '2. 객실관리 시스템'
    ws['A8'] = 'CONTROL BOX 1'; ws['K8'] = 2000000
    ws['A9'] = '-제조사 : 한국마이크로닉'
    ws['A10'] = 'CONTROL BOX 외함 400*900'; ws['K10'] = 1000000
    ws['A11'] = '온도조절기'; ws['K11'] = 5000000
    ws['A12'] = '약전 결선'; ws['D12'] = 10; ws['G12'] = 100000      # K 없음 → 수량×노무단가
    ws['A13'] = '합성수지 커버'; ws['K13'] = 9000           # 감사 : 「합」 로 시작하는 품목
    ws['A16'] = '소   계'
    ws['A17'] = '합        계'
    wb.save(p)
    return p


def main():
    print('== 1. 도우미')
    ok('d6 260929', T.d6('260929') == datetime.date(2026, 9, 29))
    ok('d6 2026-09-29', T.d6('2026-09-29') == datetime.date(2026, 9, 29))
    ok('d6 엉터리 None', T.d6('다음주') is None)
    ok('money', T.money('12,300,000원') == 12300000 and T.money('') is None)
    ok('kind_of', T.kind_of('외함 납품') == '외함' and T.kind_of('CB 속판 출고') == '속판' and T.kind_of('챠임벨 기구 입고') == '기구물' and T.kind_of('그냥') == '')
    ok('is_site', T.is_site('가나호텔') and not T.is_site('복합회의') and not T.is_site('복합회의 (가나·다라)'))
    ok('same_site 정확 일치만', T.same_site('가나호텔', '가나 호텔') and not T.same_site('가나호텔', '다라기숙사') and not T.same_site('조선호텔', '조선호텔 리뉴얼'))
    ok('특수문자·? 는 같은 현장', T.same_site('가나호텔😀', '가나호텔?') and T.same_site('가나호텔😀', '가나호텔'))
    print('== 1-1. 감사 (v2) : 현장 고르기')
    ok('정확 일치 우선', T.pick_site('조선호텔 리뉴얼', ['조선호텔', '조선호텔 리뉴얼']) == '조선호텔 리뉴얼')
    ok('앞머리 일치는 후보가 하나일 때만', T.pick_site('동구로초', ['동구로초등학교', '수유초등학교']) == '동구로초등학교')
    ok('앞머리 후보가 둘이면 None', T.pick_site('가나', ['가나호텔', '가나기숙사']) is None)
    ok('짧은 이름(2자)은 앞머리 일치 안 함', T.pick_site('제천', ['제천 리솜']) is None or T.pick_site('SK', ['SK하이닉스']) is None)
    ok('money 숫자만', T.money('90,000,000원') == 90000000 and T.money('30일') == 30 and T.money('30일/60일') is None and T.money('9천만원') is None and T.money('선수금 30%, 중도금 40%') is None)

    print('== 2. 회의록 → 납품예정')
    md = os.path.join(tmp, 'meta'); os.makedirs(md)
    put_meta(md, 'a__meta.json', '가나호텔', sched=['- 납품 | 261006 | 외함 1차 반입', '- 설치 | 261013 | 기구물 설치'],
             todo=['- 261020 | 속판 납품 일정 확인 | 배성윤 → 제작'])
    put_meta(md, 'b__meta.json', '복합회의', sched=['- 납품 | 261001 | 외함'])
    put_meta(md, 'c__meta.json', '다라기숙사', sched=['- 납품 | 261002 | 자재 반입'])   # 구분 없음
    plans, odd = T.scan_meta(md)
    ok('납품 줄 2건 (일정 외함 · 할 일 속판). 설치 줄은 아님', sorted((s, k, T.iso(d)) for s, k, d, w in plans) == [('가나호텔', '속판', '2026-10-20'), ('가나호텔', '외함', '2026-10-06')] + [('다라기숙사', '', '2026-10-02')], plans)
    ok('복합회의는 현장 확인으로', len(odd) == 1 and odd[0][0] == '복합회의')

    print('== 3. 확정 대장')
    FX.add('가나호텔', '계약금액', '90,000,000', '계약서')
    FX.add('가나호텔', '납품 외함', '260930', '차장님 한 줄')
    FX.add('가나호텔', '계산서 속판', '260901', '경리')      # 납품일 없이 계산서만 → 납품일 = 계산서일
    FX.add('가나호텔', '결제조건', '30', '계약서')
    FX.add('마바현장', '입금 기구물', '2026-09-20', '통장')
    got, contract, term, bad = T.scan_facts()
    ok('납품·계산서·입금 3건', set(got.keys()) == {('가나호텔', '외함'), ('가나호텔', '속판'), ('마바현장', '기구물')}, got.keys())
    ok('계약금액·결제조건', contract['가나호텔'][0] == 90000000 and term['가나호텔'][0] == 30)

    print('== 4. 견적서 비율')
    q = make_quote('가나호텔')
    tot, why = T.quote_split(q)
    ok('외함 1,000,000 / 속판 3,900,000(소프트웨어 포함) / 기구물 6,009,000(합성수지 포함)', tot == {'외함': 1000000, '속판': 3900000, '기구물': 6009000}, tot)
    ok('find_quote 현장 맞춤', T.find_quote('가나 호텔') == q)

    print('== 5. 통째로')
    rep = T.run(md, today_ymd='261010', quiet=True)
    bk = rows(T.book_path())
    ok('머리 9칸', bk[0] == T.HEAD, bk[0])
    by = dict(((r[0], r[1]), r) for r in bk[1:])
    tot_q = 1000000 + 3900000 + 6009000
    exp_w = str(int(round(90000000 * round(1000000.0 / tot_q, 4))))
    exp_s = str(int(round(90000000 * round(3900000.0 / tot_q, 4))))
    ok('가나 외함 : 납품일 09-30 · 예정 10-06 · 금액 = 계약 × 외함비율 · 조건 30',
       by[('가나호텔', '외함')][2] == '2026-09-30' and by[('가나호텔', '외함')][7] == '2026-10-06' and by[('가나호텔', '외함')][3] == exp_w and by[('가나호텔', '외함')][6] == '30', (by.get(('가나호텔', '외함')), exp_w))
    ok('가나 속판 : 계산서 09-01 → 납품일도 09-01 (답 1-다) · 금액 = 계약 × 속판비율', by[('가나호텔', '속판')][4] == '2026-09-01' and by[('가나호텔', '속판')][2] == '2026-09-01' and by[('가나호텔', '속판')][3] == exp_s, (by.get(('가나호텔', '속판')), exp_s))
    ok('마바 기구물 : 입금만 · 금액 빈칸(계약금액 없음) · 이유 적힘', by[('마바현장', '기구물')][5] == '2026-09-20' and by[('마바현장', '기구물')][3] == '' and any('계약금액' in n for n in rep['못 채운 이유']))
    ok('다라 : 구분 없는 납품 줄은 이유만', not any(k[0] == '다라기숙사' for k in by) and any('다라기숙사' in n and '구분' in n for n in rep['못 채운 이유']))
    ok('복합회의는 대장에 없고 이유에 「현장 확인」', not any('복합' in k[0] for k in by) and any('현장 확인' in n for n in rep['못 채운 이유']))
    rt = rows(T.ratio_path())
    ok('수금비율.csv 자동 (세 비율 합 1)', rt[1][0] == '가나호텔' and abs(float(rt[1][1]) + float(rt[1][2]) + float(rt[1][3]) - 1) < 0.001 and rt[1][5] == '자동', rt)
    ok('경리 메일 : 외함(납품 O 계산서 X) 1건', rep['경리메일'] == ['계산서요청_경리_가나호텔_외함.txt'], rep['경리메일'])
    ok('독촉 : 속판 (09-01 발행 + 30일 < 10-10, 입금 없음) 1건', rep['독촉'] == ['독촉_가나호텔_속판.txt'], rep['독촉'])
    od = os.path.join(common.cfg('out'), '수금', '261010')
    m = io.open(os.path.join(od, rep['경리메일'][0]), encoding='utf-8').read()
    ok('경리 메일에 현장·구분·납품일·금액, 「N일 지남」 없음', '가나호텔' in m and '외함' in m and '2026-09-30' in m and T.won(int(exp_w)) in m and '지남' not in m)
    ok('결과 폴더가 --today 날짜', os.path.isdir(os.path.join(common.cfg('out'), '수금', '261010')) and os.path.exists(os.path.join(common.cfg('out'), '수금', '261010', '수금자동_261010.json')))
    dn = io.open(os.path.join(od, rep['독촉'][0]), encoding='utf-8').read()
    ok('독촉에 전화 첫마디 + 메일 + [담당자] 빈칸', '[전화 첫마디]' in dn and '[메일]' in dn and '[담당자]' in dn)
    hist = rows(os.path.join(common.cfg('out'), '수금', '수금대장_이력.csv'))
    ok('이력 : 줄·칸마다 남음', len(hist) > 5 and hist[0] == T.HIST_HEAD)

    print('== 6. 차장님 값은 안 건드림 · 손 비율 · 두 번 돌려도 같음')
    bk = rows(T.book_path())
    for r in bk[1:]:
        if r[0] == '가나호텔' and r[1] == '외함':
            r[3] = '8888888'          # 차장님이 금액을 손으로 고침
    common.write_csv(T.book_path(), bk[1:], T.HEAD)
    rt = rows(T.ratio_path()); rt[1][3] = '0.5'; rt[1][5] = '손'
    common.write_csv(T.ratio_path(), rt[1:], T.RATIO_HEAD)
    FX.add('가나호텔', '납품 외함', '261001', '다시')      # 새 확정이 와도 이미 값이 있으면 안 바꿈
    n_hist = len(rows(os.path.join(common.cfg('out'), '수금', '수금대장_이력.csv')))
    rep2 = T.run(md, today_ymd='261010', quiet=True)
    bk2 = dict(((r[0], r[1]), r) for r in rows(T.book_path())[1:])
    ok('손으로 고친 금액 그대로', bk2[('가나호텔', '외함')][3] == '8888888')
    ok('이미 있는 납품일은 새 확정으로 안 바뀜', bk2[('가나호텔', '외함')][2] == '2026-09-30')
    rt2 = rows(T.ratio_path())
    ok('상태=손 비율은 다시 계산 안 함', rt2[1][3] == '0.5' and rt2[1][5] == '손', rt2)
    ok('두 번째는 새 이력 0', len(rows(os.path.join(common.cfg('out'), '수금', '수금대장_이력.csv'))) == n_hist, rep2)

    print('== 6-1. 감사 (v2) : 오귀속 · 줄 안 지움 · 엉터리 값')
    bk = rows(T.book_path())[1:]
    bk.append(['조선호텔', '외함', '', '', '', '', '', '', ''])
    bk.append(['조선호텔 리뉴얼', '외함', '', '', '', '', '', '', ''])
    bk.append(['', '외함', '2026-09-01', '5000', '', '', '30', '', '현장 빈칸 줄'])
    bk.append(['예) 광희동1가', '외함', '2026-09-01', '', '', '', '30', '', ''])
    common.write_csv(T.book_path(), bk, T.HEAD)
    FX.add('조선호텔 리뉴얼', '납품 외함', '260930', '리뉴얼')
    FX.add('가나', '결제조건', '45', '앞머리 둘')          # 가나호텔·가나기숙사 둘 다 아님 → 아무 데도 안 붙어야
    FX.add('가나기숙사', '계약금액', '9천만원', '글자')      # 숫자 아님 → 이유
    FX.add('가나기숙사', '결제조건', '30일/60일', '글자')
    FX.add('복합회의', '납품 외함', '260930', '현장 아님')
    n_before = len(rows(T.book_path())) - 1
    rep4 = T.run(md, today_ymd='261010', quiet=True)
    bk4 = rows(T.book_path())[1:]
    by4 = dict(((r[0], r[1]), r) for r in bk4)
    ok('「조선호텔 리뉴얼」 확정은 리뉴얼 줄에만', by4[('조선호텔 리뉴얼', '외함')][2] == '2026-09-30' and by4[('조선호텔', '외함')][2] == '', (by4.get(('조선호텔', '외함')), by4.get(('조선호텔 리뉴얼', '외함'))))
    ok('앞머리 후보 둘(가나)인 결제조건은 아무 줄에도 안 붙음', by4[('가나호텔', '외함')][6] == '30' and not any(r[6] == '45' for r in bk4))
    ok('현장 빈칸 줄 · 「예)」 줄이 재저장 뒤에도 남음', any(r[0] == '' and r[8] == '현장 빈칸 줄' for r in bk4) and any(r[0].startswith('예)') for r in bk4), [r[0] for r in bk4])
    ok('줄 수가 줄지 않음', len(bk4) >= n_before)
    ok('「9천만원」「30일/60일」 은 이유로 남음', any('계약금액' in n and '숫자만' in n for n in rep4['못 채운 이유']) and any('결제조건' in n and '일수 숫자만' in n for n in rep4['못 채운 이유']))
    ok('확정 대장의 「복합회의」 는 대장에 안 생김', not any('복합' in r[0] for r in bk4))
    rep5 = T.run(md, today_ymd='엉터리', quiet=True)
    ok('--today 엉터리여도 안 죽고 오늘로', rep5['today'] == datetime.date.today().isoformat())
    print('== 6-2. 감사 (v2) : 특수문자 현장명이 매번 새 줄을 만들던 것')
    put_meta(md, 'e__meta.json', '사아호텔😀', sched=['- 납품 | 261101 | 외함'])
    T.run(md, today_ymd='261010', quiet=True)
    n1 = len(rows(T.book_path()))
    T.run(md, today_ymd='261010', quiet=True)
    n2 = len(rows(T.book_path()))
    ok('두 번 돌려도 줄 수 같음 (? 로 저장된 이름과 😀 이름을 같은 현장으로 봄)', n1 == n2, (n1, n2))
    print('== 6-3. 견적서 파일 이름에 _ 가 있어도 현장 전부로 맞춤')
    make_quote('가나_별관호텔')
    ok('find_quote 「가나」 로는 못 찾음(후보 둘) · 「가나_별관호텔」 은 찾음', T.find_quote('가나') is None and T.find_quote('가나_별관호텔') is not None and T.find_quote('가나 호텔') is not None)

    print('== 7. 회의록 없이 · 대장 없이')
    tmp2 = tempfile.mkdtemp(prefix='t56b_')
    os.makedirs(os.path.join(tmp2, '_도구결과'))      # 폴더가 없으면 cfg() 가 다른 곳을 찾아가므로 먼저 만든다
    common.DEFAULTS.update({'out': os.path.join(tmp2, '_도구결과')})
    rep3 = T.run(None, today_ymd='261010', quiet=True)
    ok('아무것도 없으면 줄 0, 오류 없음', rep3['줄'] == 0 and rep3['이력'] == 0)
    shutil.rmtree(tmp2, ignore_errors=True)

    shutil.rmtree(tmp, ignore_errors=True)
    bad = [r for r in R if not r[1]]
    print('')
    print('통과 %d / 실패 %d' % (len(R) - len(bad), len(bad)))
    return 1 if bad else 0


if __name__ == '__main__':
    sys.exit(main())
