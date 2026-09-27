# -*- coding: utf-8 -*-
"""
54. 회의록 → 작업의뢰서 초안  (km_tools / t54_workorder)

회의록(meta.json)의 「타부서 전달 사항」 중 줄 끝이 「→ 작업의뢰서」 인 것만 골라
회사 원틀(MB-004)에 값만 넣은 작업의뢰서 초안을 만든다. 클로드(AI)를 쓰지 않는다 → 사용량 0.

    python t54_workorder.py <meta폴더> [--day 260921 | --from 260915 --to 260922] [--out 폴더] [--template 원틀.xlsx]

왜 (2026-09-27 차장님 「가장 하고 싶었던 것 세 가지」 중 2번)
    회의 → 할 일 까지는 07:00 메일·「26년 회의록2」 가 한다. 그런데 부서는 작업의뢰서가 없으면 움직이지 않는다.
    PLAUD 회의록은 이미 「→ 작업의뢰서」 라고 표시해 두는데, 그 뒤가 끊겨 있었다 (15번 부서메일도
    「작업의뢰서를 별도로 발행하셔야 합니다」 로 끝났다 = 「만드세요」). 이 도구가 그 고리를 잇는다.

지키는 것 (km-30 확정 규칙 · km-work-order)
    · 원틀을 복사해 값만 넣는다. 새로 그리지 않는다. 원틀이 없으면 xlsx 는 안 만들고 복사용 txt 만 낸다
    · 자동 : C4 기안일자(회의한 날) · G4 배성윤 · 현장명(E12, G6=E12) · D16 상대 회사·이름직함(전화)
             · O5/O16 =IF($O$4="","",$O$4) · 본문(회의록 글 그대로)
    · 비워 둠(차장님이 채움) : 납기일(O4 노란칸) · 업체명 · 객실수 · 계약No · 체크박스 전부 · 수량·규격 [   ]
    · 현장명은 meta.site 그대로. 「복합회의」 처럼 현장이 아닌 것은 현장명 칸을 비우고 「현장 확인」 으로 표시
    · 같은 회의 · 같은 부서는 한 장(순번으로 나눔), 부서가 다르면 장을 나눈다
    · 덮어쓰지 않는다. 같은 이름 파일이 있으면 건너뛰고 「이미 있음」 으로 적는다
    · 본문은 회의록 글에서 옮긴다. 없는 말을 지어내지 않는다 (요지 첫 줄과 「수고하세요.」 만 정해진 틀)

만드는 것 (out 폴더)
    작업의뢰서초안_<현장>_<회의날>_<부서>.xlsx   한 장씩 (원틀이 있을 때)
    작업의뢰서초안_모음_<오늘>.txt                전부의 본문 (그룹웨어에 붙여 넣기용)
    작업의뢰서초안_대장_<오늘>.csv                날짜|현장|부서|요청|상대|회의록|파일|상태
"""
from __future__ import print_function
import os, sys, io, re, json, csv, glob, shutil, datetime
from copy import copy

VERSION = 'v2 2026-09-27'   # v2 : 독립 감사 지적 고침 (부서 글자 · 태그 변형 · 현장 아님 0 · 이름 충돌 · 행 높이 · 날짜 폴더)
# v1 2026-09-27

# 줄 끝 「→ 작업의뢰서」. v2 : 「-> 작업의뢰서」 「→ 작업의뢰서 필요」 「(→ 작업의뢰서)」 「→ 작업의뢰서.」 도 받는다 (감사 지적)
TAG_RE = re.compile(r'\(?\s*(?:→|->|=>|⇒)\s*작업\s*의뢰서\s*(?:필요|요청|발행|작성)?\s*[.)\]]*\s*$')
# 「1. 설계 / 내용」 「2. 제작 — 내용」 「3. 설계 · 내용」 「4. 설계: 내용」 「1. 설비(난방) — 내용」
# v2 : 구분자 뒤에 빈칸이 있어야 자른다 → 「설계/개발 · 매핑」 은 부서 「설계/개발」, 「전기·통신 — 확인」 은 「전기·통신」,
#      「설비(난방/급수) — 확인」 은 「설비(난방/급수)」 (감사 지적 : 부서 글자가 잘려 다른 부서 요청이 섞였다)
LINE_RE = re.compile(r'^\s*[-·•]?\s*(?:\d+\s*[.)]\s*)?(?P<dept>[^\s:][^:]{0,19}?)\s*(?:(?:—|–|/|·|-|\|)\s+|\s(?:—|–|/|·|-|\|)|:\s*)(?P<what>.+?)\s*$')

# 회사 서식 수신부서 칸에 있는 부서 (체크는 차장님이 하신다. 여기서는 「사내 부서인가」 표시만)
INHOUSE = {'설계': '디자인&설계', '디자인': '디자인&설계', '개발': '개발', '구매': '구매', '제작': '제작',
           '시공': '시공', '영업': '영업', 'AS': '고객지원', '고객지원': '고객지원'}

NOT_A_SITE = ('복합회의', '확인필요', '확인 필요', '현장미정', '미정', '수금채크', '수금체크', '삭제요망')
NOT_A_SITE_IN = ('복합회의', '확인필요', '현장미정', '삭제요망', '수금채크', '수금체크')   # v2 : 「복합회의 (앵커·연합)」 처럼 붙어 있어도

CHECK_CELLS = ('F5', 'C7', 'C8', 'C9', 'C10', 'C14')
LINE_LIMIT = 116          # km-work-order : 실작성 최대 폭. 132 넘으면 인쇄에서 잘린다
YELLOW = 'FFFFFF00'


# ── 회의록 읽기 ─────────────────────────────────────────────
def _txt(v):
    return re.sub(r'^\s*[-·•]\s*', '', str(v or '')).strip()


def _list(v):
    if v is None:
        return []
    if isinstance(v, str):
        return [x for x in (_txt(s) for s in v.split('\n')) if x]
    out = []
    for it in v:
        if isinstance(it, (list, tuple)):
            s = ' | '.join(_txt(c) for c in it if _txt(c))
        else:
            s = _txt(it)
        if s:
            out.append(s)
    return out


def clear_checks(s):
    """「( V )선투입자재」 「(V )  디자인&설계」 → 괄호 안 V 만 빈칸으로. 괄호 폭은 그대로."""
    return re.sub(r'\((\s*)[Vv✓](\s*)\)', lambda m: '(%s %s)' % (m.group(1), m.group(2)), s)


def is_site(name):
    n = re.sub(r'\s+', '', name or '')
    if not n or any(x in n for x in NOT_A_SITE_IN):
        return False
    return all(n != re.sub(r'\s+', '', x) for x in NOT_A_SITE)


def parse_line(line):
    """타부서 한 줄 → (부서, 요청) 또는 None (작업의뢰서 표시가 없는 줄)"""
    s = _txt(line)
    if not TAG_RE.search(s):
        return None
    body = TAG_RE.sub('', s).strip()
    m = LINE_RE.match(body)
    if not m:
        return ('확인필요', re.sub(r'^\d+\s*[.)]\s*', '', body))
    dept = m.group('dept').strip()
    return (dept, m.group('what').strip())


def dept_key(dept):
    """묶는 데만 쓴다 (글자는 원문대로 남긴다). 괄호 안과 빈칸만 뗀다.
    「설비(난방)」 → 「설비」 / 「설계/개발」 → 「설계/개발」 (v2 : 설계 장에 개발 요청이 섞이지 않게)"""
    k = re.sub(r'\s+', '', re.sub(r'\(.*?\)', '', dept or ''))
    return k or (dept or '').strip()


def inhouse(dept):
    k = dept_key(dept)
    for a, b in INHOUSE.items():
        if k.startswith(a):
            return b
    return ''


def phone_fmt(p):
    d = re.sub(r'\D', '', p or '')
    if len(d) == 11:
        return '%s-%s-%s' % (d[:3], d[3:7], d[7:])
    if len(d) == 10:
        return '%s-%s-%s' % (d[:3], d[3:6], d[6:])
    return p or ''


def counterpart(m):
    """D16 업체담당자/연락처. 회사 칸에 안건이 통째로 들어온 자료는 버린다 (t52 who 와 같은 기준)."""
    comp = (m.get('company') or '').strip()
    if len(comp) > 20 or ',' in comp:
        comp = ''
    name = (m.get('name') or '').strip()
    rank = (m.get('rank') or '').strip()
    base = (name + ' ' + rank).strip() if name else (m.get('person') or '').strip()
    if comp and comp == base:
        comp = ''
    who = ' '.join(x for x in (comp, base) if x).strip()
    ph = phone_fmt(m.get('phone'))
    return (who + (' (%s)' % ph if ph else '')).strip()


def read_one(path):
    with io.open(path, 'r', encoding='utf-8') as f:
        d = json.load(f)
    m = d.get('meta') or {}
    s1 = (d.get('sec') or {}).get('1') or {}
    s2 = (d.get('sec') or {}).get('2') or {}
    site = (m.get('site') or s1.get('현장') or '').strip()
    reqs, odd = [], []
    for ln in _list(s2.get('타부서 전달 사항')):
        p = parse_line(ln)
        if p:
            reqs.append(p)
        elif '의뢰서' in ln:
            odd.append(ln)          # v2 : 「의뢰서」 는 있는데 표시 모양이 달라 못 읽은 줄 — 조용히 버리지 않고 모음·대장에 알린다
    todos = []
    for ln in _list(s2.get('할 일')):
        parts = [x.strip() for x in ln.split('|')]
        if len(parts) >= 2:
            todos.append((parts[0], parts[1], parts[2] if len(parts) > 2 else ''))
    decisions = []
    for it in (s1.get('안건목록') or []):
        if isinstance(it, dict):
            dc = _txt(it.get('decision'))
            if dc and not re.match(r'^(없음|해당\s*없음|미정)(\s*[—\-–:]|$)', dc):
                decisions.append(dc)
    return {'file': os.path.basename(path), 'site': site, 'ymd': (m.get('ymd') or '').strip(),
            'hm': (m.get('hm') or '').strip(), 'who': counterpart(m), 'who_short': counterpart(dict(m, phone='')),
            'person': (m.get('person') or '').strip(), 'name': (m.get('name') or '').strip(),
            'reqs': reqs, 'odd': odd, 'todos': todos, 'decisions': decisions}


def collect(meta_dir, d_from=None, d_to=None):
    recs, skipped = [], []
    for root, _d, files in os.walk(meta_dir):
        for fn in sorted(files):
            if not fn.endswith('meta.json'):
                continue
            if fn.startswith('_삭제요망'):
                skipped.append((fn, '삭제요망'))
                continue
            try:
                r = read_one(os.path.join(root, fn))
            except Exception as e:
                skipped.append((fn, '읽기실패 %s' % e))
                continue
            if (d_from and (not r['ymd'] or r['ymd'] < d_from)) or (d_to and (not r['ymd'] or r['ymd'] > d_to)):
                continue
            recs.append(r)
    recs.sort(key=lambda r: (r['ymd'], r['hm'], r['site']))
    return recs, skipped


# ── 의뢰서 한 장 계획 ───────────────────────────────────────
def _md(ymd):
    return '%d월 %d일' % (int(ymd[2:4]), int(ymd[4:6])) if re.match(r'^\d{6}$', ymd or '') else ''


def plan(recs):
    """회의 × 부서 → 한 장. 순서는 회의 날·시각, 부서는 회의록에 처음 나온 순서."""
    out = []
    for r in recs:
        if not r['reqs']:
            continue
        order, bag = [], {}
        for dept, what in r['reqs']:
            k = dept_key(dept)
            if k not in bag:
                order.append(k)
                bag[k] = {'dept': dept, 'items': []}
            bag[k]['items'].append(what)
        for k in order:
            g = bag[k]
            # 같은 부서 할 일만. 「작업의뢰서 작성」 같은 이 문서 자체를 가리키는 할 일은 뺀다
            rel = [t for t in r['todos'] if k and k in t[2] and '의뢰서' not in t[1]]
            out.append({'rec': r, 'dept': g['dept'], 'key': k, 'items': g['items'], 'todos': rel,
                        'site_ok': is_site(r['site']), 'inhouse': inhouse(g['dept'])})
    return out


def body_lines(p):
    """[(순번 or None, 한 줄)] . 회의록 글을 옮기고, 수량·규격은 [   ] 로 비운다."""
    r = p['rec']
    site = r['site'] if p['site_ok'] else '[현장 확인]'
    head = '%s 현장 %s %s 협의 결과, 아래 작업을 요청 드립니다.' % (site, _md(r['ymd']), r['person'] or r['who'] or '')
    lines = [(1, re.sub(r'\s+', ' ', head).strip())]
    n = 1
    for what in p['items']:
        n += 1
        lines.append((n, '%s : %s' % (p['dept'], what)))
    lines.append((None, ' -수량 : [   ]    -규격·사양 : [   ]'))
    if p['todos']:
        lines.append((None, '관련 할 일 (회의록)'))
        for due, what, _who in p['todos']:
            d = '' if due in ('', '미정') else '  (기한 %s)' % due
            lines.append((None, ' -%s%s' % (what, d)))
    if r['decisions']:
        lines.append((None, '협의 결정 (회의록)'))
        for dc in r['decisions'][:5]:
            lines.append((None, ' -%s' % dc))
    lines.append((None, '근거 : %s 협의록' % ' '.join(x for x in (r['ymd'], r['hm'], r['who_short'] or r['person']) if x)))
    lines.append((None, '수고하세요.'))
    return lines


def _w(s):
    return sum(2 if ord(c) > 0x1100 else 1 for c in str(s))


def split_line(text, limit=LINE_LIMIT):
    """km-work-order fill_workorder.split_line 과 같은 규칙 (폭 116, 이어지는 줄은 두 칸 들여쓰기)."""
    text = str(text).rstrip()
    if _w(text) <= limit:
        return [text]
    out, rest = [], text
    while _w(rest) > limit:
        cut = len(rest)
        while _w(rest[:cut]) > limit:
            cut -= 1
        seg = rest[:cut]
        for sep in (' // ', ', ', ',', ' '):
            q = seg.rfind(sep)
            if q > cut * 0.45:
                cut = q + (len(sep) if sep != ' ' else 1)
                break
        out.append(rest[:cut].rstrip())
        rest = '  ' + rest[cut:].lstrip()
    if rest.strip():
        out.append(rest)
    return out


# ── 원틀 찾기 · 채우기 ─────────────────────────────────────
def find_template(explicit=None):
    """1) --template 2) 환경변수 KM_WO_TEMPLATE 3) _원틀 폴더(설정.ini) 4) km-work-order 스킬 자산.
    없으면 None → xlsx 를 만들지 않는다 (원틀을 새로 그리지 않는다)."""
    cands = [explicit, os.environ.get('KM_WO_TEMPLATE')]
    try:
        sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
        import common as C
        root = C.cfg('template')
        if root and os.path.isdir(root):
            for p in sorted(glob.glob(os.path.join(root, '**', '*.xlsx'), recursive=True)):
                b = os.path.basename(p)
                if ('작업의뢰서' in b or 'MB-004' in b or 'workorder' in b) and not b.startswith('~$'):
                    cands.append(p)
    except Exception:
        pass
    cands += sorted(glob.glob(os.path.expanduser('~/.claude/skills/**/km-work-order/assets/workorder_template.xlsx'),
                              recursive=True))
    cands += sorted(glob.glob('/root/.claude/skills/**/km-work-order/assets/workorder_template.xlsx', recursive=True))
    for c in cands:
        if c and os.path.isfile(c):
            return c
    return None


def _sheet(wb):
    for n in wb.sheetnames:
        if '그룹웨어' in n:
            return wb[n]
    for ws in wb.worksheets:
        if '작' in str(ws['A1'].value or '') and '서' in str(ws['A1'].value or ''):
            return ws
    return wb.worksheets[0]


def _safe(s, n=40):
    s = re.sub(r'[\\/:*?"<>|\[\]\s]+', '_', s or '').strip('_')
    return s[:n] or '미정'


def fill_xlsx(tpl, out_path, p, lines):
    import openpyxl
    from openpyxl.styles import Alignment, PatternFill
    from openpyxl.worksheet.properties import PageSetupProperties
    shutil.copy(tpl, out_path)
    wb = openpyxl.load_workbook(out_path)
    ws = _sheet(wb)
    for n in list(wb.sheetnames):             # 작성예시 등 다른 시트는 뺀다 (원틀 시트는 그대로)
        if wb[n] is not ws:
            del wb[n]
    r = p['rec']
    ws.title = _safe('%s(%s)' % (r['site'] if p['site_ok'] else '현장확인', p['dept']), 31)
    if re.match(r'^\d{6}$', r['ymd']):
        ws['C4'] = datetime.datetime(2000 + int(r['ymd'][:2]), int(r['ymd'][2:4]), int(r['ymd'][4:6]))
    ws['G4'] = '배성윤'
    ws['O4'] = None
    ws['O4'].fill = PatternFill('solid', fgColor=YELLOW)      # 납기일 = 차장님이 넣는다
    ws['O5'] = '=IF($O$4="","",$O$4)'
    ws['O16'] = '=IF($O$4="","",$O$4)'
    ws['E12'] = r['site'] if p['site_ok'] else None
    ws['G6'] = '=E12' if p['site_ok'] else None     # v2 : 현장이 아닐 때 =E12 는 「0」 으로 보인다 (감사 지적)
    if r['who']:
        ws['D16'] = r['who']
    # 체크박스(F5·C7·C8·C9·C10·C14) : 전부 빈 괄호로. 차장님이 V 를 넣으신다 (km-30)
    #   ★ 원틀 파일에 예전 작성분의 V 가 남아 있다 (2026-09-27 렌더 검수에서 발견 : 기타·선발행·계약·유상·개발·구매·택배).
    #     그대로 두면 차장님이 정하지 않은 체크가 결재에 올라간다 → 괄호 안 V 만 지운다 (글자 배치는 그대로)
    for k in CHECK_CELLS:
        v = ws[k].value
        if isinstance(v, str):
            ws[k] = clear_checks(v)

    flat = []
    for no, txt in lines:
        parts = split_line(txt)
        flat.append((no, parts[0]))
        flat.extend((None, x) for x in parts[1:])

    last_body = 35
    if len(flat) > 15:                        # km-work-order make_workorder 와 같은 순서 (병합 → 삽입 → 서식 → 로고)
        extra = len(flat) - 15                # v2 : 본문 뒤 빈 줄 2개도 남게 (감사 지적)
        for mg in [mg for mg in ws.merged_cells.ranges if mg.min_row >= 19]:
            ws.unmerge_cells(str(mg))
        # v2 : insert_rows 는 행 높이를 안 옮긴다 → 35행부터 아래 높이를 적어 두었다가 밀린 자리에 다시 준다
        below = dict((rr, ws.row_dimensions[rr].height) for rr in range(35, ws.max_row + 1))
        ws.insert_rows(35, extra)
        for rr, hgt in sorted(below.items(), reverse=True):
            ws.row_dimensions[rr + extra].height = hgt
        src = 34
        for rr in range(35, 35 + extra):
            ws.row_dimensions[rr].height = ws.row_dimensions[src].height or 25.2
            for col in range(1, 19):
                s, d = ws.cell(src, col), ws.cell(rr, col)
                d.font = copy(s.font); d.border = copy(s.border)
                d.fill = copy(s.fill); d.alignment = copy(s.alignment)
        for img in ws._images:
            img.anchor._from.row += extra; img.anchor.to.row += extra
        last_body = 35 + extra
        for rr in range(19, last_body + 1):
            ws.merge_cells('B%d:R%d' % (rr, rr))

    row = 19
    for no, txt in flat:
        if no:
            ws['A%d' % row] = no
        ws['B%d' % row] = txt
        ws['B%d' % row].alignment = Alignment(horizontal='left', vertical='center')
        if '[   ]' in txt:
            ws['B%d' % row].fill = PatternFill('solid', fgColor=YELLOW)
        row += 1

    first_del = row + 2
    if first_del <= last_body:
        n_del = last_body - first_del + 1
        for mg in [mg for mg in ws.merged_cells.ranges if mg.min_row >= first_del]:
            ws.unmerge_cells(str(mg))
        ws.delete_rows(first_del, n_del)
        for img in ws._images:
            img.anchor._from.row -= n_del
            img.anchor.to.row -= n_del
        tail = first_del
    else:
        tail = last_body + 1
    ws.print_area = 'A1:R%d' % tail
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    wb.save(out_path)
    return tail


# ── 실행 ───────────────────────────────────────────────────
def run(meta_dir, d_from=None, d_to=None, out=None, template=None, today=None):
    today = today or datetime.date.today().strftime('%y%m%d')
    out = out or os.path.join(os.getcwd(), '작업의뢰서초안', today)
    os.makedirs(out, exist_ok=True)
    recs, skipped = collect(meta_dir, d_from, d_to)
    plans = plan(recs)
    tpl = find_template(template)
    try:
        import openpyxl  # noqa
    except Exception:
        tpl = None
    rows, txt = [], []
    made = skipped_same = 0
    used = set()
    for p in plans:
        r = p['rec']
        lines = body_lines(p)
        site_lbl = r['site'] if p['site_ok'] else '현장확인'
        # 같은 현장·같은 날·같은 부서 회의가 여러 건일 수 있다 (2026-09-21 앵커호텔 설계 3건) → 시각·상대로 가른다
        # v2 : 이름이 회의마다 늘 같게 (감사 지적 : -2·-3 을 순서로 붙이면 회의가 늘 때 다른 회의 파일을 「이미 있음」 으로 봤다)
        #      시각이 있으면 시각, 없으면 상대 이름 + 회의록 파일 표식 4자리(회의록 파일 이름에서 나옴 → 늘 같다)
        fid = '%04x' % (sum(ord(c) * (i + 1) for i, c in enumerate(r['file'])) % 65536)
        hm = re.sub(r'\D', '', r['hm'])
        tag = hm or '%s_%s' % (_safe(r['name'] or r['person'].split(' ')[0], 10), fid)
        name = '작업의뢰서초안_%s_%s_%s_%s' % (_safe(site_lbl, 20), r['ymd'], tag, _safe(p['key'], 10))
        if name in used:                       # 같은 시각·같은 부서 회의 두 건 (중복 올림 등) → 표식을 붙인다
            name = '%s_%s' % (name, fid)
        used.add(name)
        xp = os.path.join(out, name + '.xlsx')
        # v2 : 결과 폴더가 날짜별이라 다음 날 다시 돌리면 같은 초안을 또 만들었다 → 옆 날짜 폴더도 본다
        before = [q for q in glob.glob(os.path.join(os.path.dirname(os.path.abspath(out)), '*', name + '.xlsx'))
                  if os.path.dirname(os.path.abspath(q)) != os.path.abspath(out)]
        state = ''
        if tpl:
            if os.path.exists(xp) or before:
                state = '이미 있음'
                skipped_same += 1
            else:
                fill_xlsx(tpl, xp, p, lines)
                state = '만듦'
                made += 1
        else:
            state = '원틀 없음 (본문만)'
        flags = []
        if not p['site_ok']:
            flags.append('현장 확인 (%s)' % (r['site'] or '빈칸'))
        if not p['inhouse']:
            flags.append('받는 부서 확인 (「%s」 — 서식 수신부서 칸에 없는 이름)' % p['dept'])
        rows.append([r['ymd'], r['site'], p['dept'], ' / '.join(p['items']), r['who'], r['file'],
                     os.path.basename(xp) if tpl else '', state, '; '.join(flags)])
        txt.append('━━ %s · %s · %s  (%s)' % (site_lbl, _md(r['ymd']), p['dept'], state))
        if flags:
            txt.append('   ※ ' + ' / '.join(flags))
        txt.append('   납기일 : [   ]   업체명 : [   ]   객실수 : [   ]   체크박스 : 차장님')
        txt.append('   업체담당자 : %s' % (r['who'] or '[   ]'))
        for no, t in lines:
            for i, part in enumerate(split_line(t)):
                txt.append('%3s  %s' % (no if (no and i == 0) else '', part))
        txt.append('')
    odd = [(r, ln) for r in recs for ln in r.get('odd', [])]
    if odd:
        txt.append('━━ 「의뢰서」 는 있는데 표시 모양이 달라 초안을 안 만든 줄 %d개 — 보시고 필요하면 말씀해 주십시오' % len(odd))
        for r, ln in odd:
            txt.append('   %s %s : %s' % (r['ymd'], r['site'], ln))
            rows.append([r['ymd'], r['site'], '', ln, r['who'], r['file'], '', '못 읽음 (표시 모양)', '「→ 작업의뢰서」 모양이 아님'])
        txt.append('')
    head = ['[KM] 작업의뢰서 초안 %s  —  회의 %d건 중 「→ 작업의뢰서」 %d장' % (today, len(recs), len(plans)),
            '원틀 : %s' % (os.path.basename(tpl) if tpl else '없음 → xlsx 안 만듦. 아래 본문을 그룹웨어 원틀에 붙여 넣으십시오'),
            '넣으실 것 : 납기일(O4 노란칸 하나만 치면 출고요청일·현장납기일이 따라옵니다) · 업체명 · 객실수 · 체크박스 V · 수량·규격 [   ]',
            '']
    tp = os.path.join(out, '작업의뢰서초안_모음_%s.txt' % today)
    io.open(tp, 'w', encoding='utf-8').write('\n'.join(head + txt) + '\n')
    cp = os.path.join(out, '작업의뢰서초안_대장_%s.csv' % today)
    with io.open(cp, 'w', encoding='utf-8-sig', newline='') as f:
        w = csv.writer(f)
        w.writerow(['회의날', '현장', '부서', '요청', '상대', '회의록', '파일', '상태', '확인할 것'])
        w.writerows(rows)
    rep = {'version': VERSION, 'today': today, 'meetings': len(recs), 'sheets': len(plans), 'made': made, 'odd': len(odd),
           'same': skipped_same, 'template': tpl or '', 'skipped': skipped, 'out': out, 'txt': tp, 'csv': cp}
    io.open(os.path.join(out, '작업의뢰서초안_자가진단_%s.json' % today), 'w', encoding='utf-8').write(
        json.dumps(rep, ensure_ascii=False, indent=1))
    return rep


def main(argv):
    if not argv or argv[0] in ('-h', '--help'):
        print(__doc__)
        return 0
    meta_dir = argv[0]
    opt = {}
    i = 1
    while i < len(argv):
        if argv[i].startswith('--') and i + 1 < len(argv):
            opt[argv[i][2:]] = argv[i + 1]
            i += 2
        else:
            i += 1
    d_from, d_to = opt.get('from'), opt.get('to')
    if opt.get('day'):
        d_from = d_to = opt['day']
    rep = run(meta_dir, d_from, d_to, opt.get('out'), opt.get('template'), opt.get('today'))
    print('54. 작업의뢰서 초안 %s' % VERSION)
    print('회의 %d건 → 의뢰서 %d장 (새로 만듦 %d · 이미 있음 %d)' % (rep['meetings'], rep['sheets'], rep['made'], rep['same']))
    print('원틀 : %s' % (rep['template'] or '없음 (본문 txt 만)'))
    print('모음 : %s' % rep['txt'])
    print('대장 : %s' % rep['csv'])
    return 0


if __name__ == '__main__':
    sys.exit(main(sys.argv[1:]))
