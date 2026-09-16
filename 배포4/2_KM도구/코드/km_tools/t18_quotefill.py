# -*- coding: utf-8 -*-
"""12. 견적서 채우기 - 수량표를 회사 견적서 원틀에 부어 넣는다.
원틀은 절대 새로 그리지 않고 복사해 값만 바꾼다 (6.9). 단가는 프로님이 넣으십니다."""
import os, sys
from common import *
import t07_delta

def fill(tpl, rows, site, out):
    import openpyxl
    from copy import copy
    wb = openpyxl.load_workbook(tpl)
    ws = None
    for w in wb.worksheets:
        if '내역' in w.title:
            ws = w; break
    ws = ws or wb.worksheets[-1]
    head = None
    for r in ws.iter_rows(min_row=1, max_row=40):
        txt = ' '.join(str(c.value) for c in r if c.value)
        if ('품명' in txt or '품 명' in txt) and ('수량' in txt or '규격' in txt):
            head = r[0].row; break
    if head is None:
        return None, '원틀에서 품명/수량 머리글 행을 못 찾았습니다'
    start = head + 1
    sample = [copy(c._style) for c in ws[start]]
    for i, (name, qty) in enumerate(rows):
        r = start + i
        ws.cell(row=r, column=1, value=name)
        for col in range(1, 12):
            v = str(ws.cell(row=head, column=col).value or '')
            if '수량' in v:
                ws.cell(row=r, column=col, value=qty)
                ws.cell(row=r, column=col).number_format = '#,##0'
            if '단위' in v and not ws.cell(row=r, column=col).value:
                ws.cell(row=r, column=col, value='EA')
        for col, st in enumerate(sample, 1):
            try:
                ws.cell(row=r, column=col)._style = st
            except Exception:
                pass
    wb.save(out)
    return out, None

def run(qty_file=None):
    title('12. 견적서 채우기')
    tpl = find_template('견적')
    if not tpl:
        no_template('견적서 원틀(xlsx)', '견적')
        return
    print('쓸 원틀 : %s' % os.path.basename(tpl))
    qty_file = qty_file or ask('수량표 경로 (csv/xlsx) > ')
    if not os.path.isfile(qty_file):
        print('수량표를 찾지 못했습니다.'); return
    site = ask('현장명 > ', '현장미정')
    m = t07_delta.to_map(t07_delta.read_rows(qty_file))
    rows = [(k, int(v) if float(v) == int(v) else v) for k, v in m.items()]
    if not rows:
        print('수량표에서 품목/수량을 못 읽었습니다. 1열=품명, 뒤쪽 칸=수량 형태여야 합니다.'); return
    out = os.path.join(outdir('견적서'), '%s_견적서_%s.xlsx' % (safe_name(site), ymd6()))
    p, err = fill(tpl, rows, site, out)
    if err:
        print('[실패] %s' % err); return
    print('품목 %s개를 넣었습니다.' % won(len(rows)))
    print('파일 : %s' % p)
    print('* 단가/금액은 비어 있습니다. 단가장에서 프로님이 넣으십니다.')
    print('* 소계 범위는 눈으로 한 번 검산하십시오 (구간 누락 사고 사례 있음).')
    log('견적서', '%s %d품목' % (site, len(rows)))

if __name__ == '__main__':
    a = sys.argv[1:]
    run(a[0] if a else None); pause()
