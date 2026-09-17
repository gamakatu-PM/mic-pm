# -*- coding: utf-8 -*-
"""33. 현황판 - 모든 도구가 낸 것을 한 장으로. 프로님도 보고, 새 창의 클로드도 읽는다. 토큰 0.

만드는 것
  {base}\\_현황판.html        프로님용 한 장 (더블클릭)
  {인수인계함}\\_현황판.md    클로드용 (새 창이 이것부터 읽는다)
담는 것
  ① 오늘 결정할 것(결정대기 ★급함)  ② 현장별 도면 판·증감  ③ 클로드 부탁서 대기
  ④ 납기 경보·수금  ⑤ 단가장 상태(미확정 주황 줄)  ⑥ 최근 도구 실행 기록
"""
import os, re, csv, glob, io as _io, datetime, collections
from common import *
import sitebook, t13_brief, t24_handover
import t27_drawing as D
import t28_cost as C
import t31_intake as I

TOOL = '현황판'

def urgent_rows():
    x = t24_handover.latest_xlsx(cfg('handover'))
    out = []
    if not x:
        return out, None
    try:
        import openpyxl
        ws = openpyxl.load_workbook(x, data_only=True)['결정대기']
        for r in ws.iter_rows(min_row=5, values_only=True):
            if not r or not r[0]:
                continue
            ans = r[8] if len(r) > 8 else ''
            done = r[9] if len(r) > 9 else ''
            if done or (ans and str(ans).strip()):
                continue
            out.append((r[0], str(r[1] or ''), str(r[2] or ''), str(r[3] or '')))
    except Exception:
        pass
    out.sort(key=lambda t: (t[1] != '★급함', t[0]))
    return out, x

def site_rows():
    out = []
    for name, p in I.sites():
        led = I.load_ledger(p)
        revs = [int(r['판']) for r in led if str(r['판']).isdigit()]
        rev = max(revs) if revs else 0
        last = max((r['읽은날'] for r in led), default='')
        od = os.path.join(cfg('out'), I.TOOL, safe_name(name))
        dif = sorted(glob.glob(os.path.join(od, '*_증감_*.csv')), key=os.path.getmtime, reverse=True)
        nd = 0
        if dif:
            nd = max(0, len([l for l in read_text(dif[0]).splitlines() if l.strip()]) - 1)
        nfiles = len(I.drawing_files(p))
        out.append((name, nfiles, rev, last, nd, dif[0] if dif else None))
    return out

def request_rows():
    od = os.path.join(cfg('out'), '완성품')
    reqs = sorted(glob.glob(os.path.join(od, '*', '_클로드부탁서_*.md')), key=os.path.getmtime, reverse=True)
    seen, out = set(), []
    for p in reqs:
        m = re.search(r'_클로드부탁서_(.+?)_\d{6}\.md$', os.path.basename(p))
        site = m.group(1) if m else '?'
        if site in seen:
            continue
        seen.add(site)
        n = len(re.findall(r'^## \d+\.', read_text(p), re.M))
        out.append((site, n, p, datetime.date.fromtimestamp(os.path.getmtime(p)).isoformat()))
    ad = os.path.join(C.root(), '받은답')
    pending = [f for f in glob.glob(os.path.join(ad, '*.csv')) if not os.path.basename(f).startswith('_처리')]
    return out, pending

def price_rows():
    p = C._find_pb_file()
    if not p:
        return None, 0, 0
    pb = C.load_pricebook() or []
    orange = 0
    if p.lower().endswith(('.xlsx', '.xlsm')):
        try:
            import openpyxl
            wb = openpyxl.load_workbook(p, data_only=True)
            ws = wb[sorted(wb.sheetnames, key=lambda n: ('총괄' not in n, n))[0]]
            for row in ws.iter_rows():
                c = row[1] if len(row) > 1 else row[0]
                if c.fill and c.fill.fgColor and str(c.fill.fgColor.rgb).endswith('FFE0B2'):
                    orange += 1
        except Exception:
            pass
    return p, len(pb), orange

def recent_log(n=10):
    p = os.path.join(cfg('out'), '_실행기록.txt')
    if not os.path.exists(p):
        return []
    lines = [l for l in read_text(p).splitlines() if l.strip()]
    return lines[-n:][::-1]

def build(quiet=False):
    sitebook.sync_from_folders()
    urg, xfile = urgent_rows()
    red, yel = t13_brief.due_rows()
    unbilled, waiting = t13_brief.money_rows()
    srows = site_rows()
    reqs, pending = request_rows()
    pbf, pbn, orange = price_rows()
    logs = recent_log()
    t0 = today()

    B1 = [('red' if u == '★급함' else 'yellow', '[%s] %s · %s' % (u, k, w[:70])) for no, u, k, w in urg[:12]]
    B2 = []
    for name, nf, rev, last, nd, dif in srows:
        col = 'blue' if nd else ('green' if rev else 'gray')
        B2.append((col, '%s : 도면 %d개 · r%d · 마지막 읽음 %s · 최근 판 변경 %d품목' % (name, nf, rev, last or '-', nd)))
    if not B2:
        B2 = [('gray', '현장 폴더 없음 - 3_공통사용\\도면\\{현장명}\\ 을 만들고 도면을 넣으십시오')]
    B3 = [('yellow', '%s : 부탁서 %d항목 (%s) - 대화창에 끌어다 넣으십시오' % (s, n, d)) for s, n, p, d in reqs]
    if pending:
        B3.append(('blue', '받은답 %d개가 아직 반영 전 - 30번을 누르십시오' % len(pending)))
    if not B3:
        B3 = [('green', '클로드에게 넘길 것 없음')]
    B4 = [('red', '%s · %s 기한 %s (%d일 지남)' % (s, n, d.isoformat(), -dd)) for s, n, d, dd in sorted(red, key=lambda x: x[3])]
    B4 += [('red', '계산서 미발행 : %s %s · %s원 (납품 %d일 경과)' % (s, k, won(a), g)) for s, k, a, g in unbilled]
    B4 += [('yellow', '%s · %s 까지 %s (D%+d)' % (s, n, d.isoformat(), dd)) for s, n, d, dd in sorted(yel, key=lambda x: x[3])]
    B4 += [('yellow', '입금 대기 : %s %s · %s원 (D%+d)' % (s, k, won(a), g)) for s, k, a, g in waiting]
    import t05_schedule
    for s_ in sitebook.load():
        if not s_['due']:
            continue
        try:
            due = datetime.datetime.strptime(s_['due'], '%Y-%m-%d').date()
        except Exception:
            continue
        nxt = [(n, d) for n, d, w in t05_schedule.back(due) if d >= t0]
        if nxt:
            n, d = nxt[0]
            B4.append(('blue', '%s : 다음 단계 「%s」 %s (D%+d) · 준공 %s'
                       % (s_['site'], n, d.isoformat(), (d - t0).days, s_['due'])))
    if not B4:
        B4 = [('green', '경보 없음 (현장대장에 준공일이 없으면 여기가 비어 보입니다)')]
    if pbf:
        B5 = [('green' if not orange else 'yellow', '%s · %s줄 · 미확정(주황) %d줄' % (os.path.basename(pbf), won(pbn), orange))]
    else:
        B5 = [('red', '단가장이 PC에 없습니다 - 3_공통사용\\단가장\\ 에 드라이브 CB모듈_단가장 xlsx 를 넣으십시오')]
    B6 = [('gray', l) for l in logs]
    B7 = []
    try:
        import t40_meeting
        for r in t40_meeting.changes_for('', days=30)[:15]:
            B7.append(('yellow', '%s · %s · %s : %s' % (r[0], r[1], r[2], r[4][:80])))
    except Exception:
        pass
    if not B7:
        B7 = [('green', '최근 30일 회의록에 수량·규격 변경 없음')]

    blocks = [('① 오늘 결정할 것 (결정대기)', B1),
              ('② 현장별 도면 (판 · 증감)', B2),
              ('③ 클로드에게 넘길 것', B3),
              ('④ 납기 경보 · 수금', B4),
              ('⑤ 단가장', B5),
              ('⑥ 최근 도구 실행', B6),
              ('⑦ 회의에서 바뀐 수량·규격 (최근 30일, PLAUD 회의록)', B7)]
    links = [x for x in [xfile] + [p for s, n, p, d in reqs] + [d for *_, d in srows if d] + [pbf] if x]
    od = outdir(TOOL)
    f_out = write_html(os.path.join(od, '현황판_%s.html' % ymd6()), 'KM 현황판', blocks, files=links)
    # base 에 한 장
    top = os.path.join(cfg('base'), '_현황판.html')
    try:
        import shutil
        shutil.copy(f_out, top)
    except Exception:
        top = f_out
    # 클로드용 md
    md = ['# KM 현황판  (%s 생성 · 도구 %s)' % (t0.isoformat(), VERSION), '',
          '새 창의 클로드는 이 파일을 가장 먼저 읽는다. 아래 숫자는 프로님 PC 도구가 방금 낸 것이다.', '']
    for name, rows in blocks:
        md.append('## ' + name)
        for col, line in rows:
            md.append('- [%s] %s' % (col, line))
        md.append('')
    md.append('## 파일')
    for x in links:
        md.append('- %s' % x)
    hp = cfg('handover')
    mdp = os.path.join(hp if os.path.isdir(hp) else od, '_현황판.md')
    _io.open(mdp, 'w', encoding='utf-8').write('\n'.join(md))
    if not quiet:
        print('현황판 : %s' % top)
        print('클로드용 : %s' % mdp)
    return top, mdp, blocks

def run(quiet=False):
    title('33. 현황판   (모든 도구 결과를 한 장으로. 토큰 0)')
    top, mdp, blocks = build(quiet)
    for name, rows in blocks:
        print('')
        print(name)
        for col, line in rows[:6]:
            print('   %-6s %s' % (col, line[:90]))
    log(TOOL, '생성')
    if not quiet and not open_file(top):
        open_folder(os.path.dirname(top))
    return top

if __name__ == '__main__':
    run(); pause()
