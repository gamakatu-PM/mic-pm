# -*- coding: utf-8 -*-
"""54번 작업의뢰서 초안 시험. 가짜 회의록(실제 현장·사람 아님)으로 돌린다.
    python t54_workorder_시험.py
"""
from __future__ import print_function
import os, sys, io, json, shutil, tempfile

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import t54_workorder as T

R = []


def ok(name, cond, note=''):
    R.append((name, bool(cond), note))
    print('  %s %s %s' % ('OK  ' if cond else 'FAIL', name, '' if cond else note))


def meta(site, ymd, hm, dept_lines, todos=(), decisions=(), company='가나건설', name='홍길동', rank='소장', phone='01012345678'):
    return {'meta': {'site': site, 'ymd': ymd, 'hm': hm, 'company': company, 'name': name, 'rank': rank,
                     'person': '%s %s' % (name, rank), 'phone': phone},
            'sec': {'1': {'현장': site, '안건목록': [{'title': 't', 'bullets': [], 'decision': d, 'actions': []} for d in decisions]},
                    '2': {'타부서 전달 사항': list(dept_lines), '할 일': list(todos)}}}


def _fid(fn):
    return '%04x' % (sum(ord(c) * (i + 1) for i, c in enumerate(fn)) % 65536)


def fake_template(path):
    """원틀과 같은 칸 구조만 흉내 낸 시험용 서식 (회사 원틀 아님)."""
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '작업의뢰서 (그룹웨어)'
    ws['A1'] = '작  업  의  뢰  서'
    for r in range(19, 36):
        ws.merge_cells('B%d:R%d' % (r, r))
    ws['A37'] = '양식 MB- 004    REV.0     보존년한: 영구'
    ws['F5'] = '( )선투입자재 ( )목업룸 (V)기타'            # 회사 원틀처럼 예전 V 가 남아 있는 상태
    ws['C7'] = '     (V)  선발행                 ( )  후발행'
    ws['C8'] = '     ( )  계약 전(샘플)       ( V ) 계약'
    ws['C9'] = '     ( )  무상                   (V)  유상 ( \\                 )   '
    ws['C10'] = '     ( )  영업    ( )  디자인&설계    (V)  개발   (V)  구매'
    ws['C14'] = '     ( )  직납          (V)  택배'
    wb.create_sheet('작성예시 (가짜)')
    wb.save(path)


def main():
    print('== 1. 한 줄 읽기')
    ok('「/」 구분', T.parse_line('1. 설계 / 회로 비교표 작성 → 작업의뢰서') == ('설계', '회로 비교표 작성'))
    ok('「—」 구분', T.parse_line('2. 제작 — 스위치함 용량 검토 → 작업의뢰서') == ('제작', '스위치함 용량 검토'))
    ok('「·」 구분', T.parse_line('3. 설계 · 도면 노트 문구 반영 → 작업의뢰서') == ('설계', '도면 노트 문구 반영'))
    ok('「:」 구분', T.parse_line('4. 설계: 타입별 수량 정리 → 작업의뢰서') == ('설계', '타입별 수량 정리'))
    ok('괄호 부서', T.parse_line('1. 설비(난방) — 센서 매설 위치 확정 → 작업의뢰서') == ('설비(난방)', '센서 매설 위치 확정'))
    ok('「- 」 로 시작하는 줄', T.parse_line('- 5. 시공 — 절체 순서 조사 → 작업의뢰서') == ('시공', '절체 순서 조사'))
    ok('표시 없는 줄은 None', T.parse_line('3. 전기 — 차단기 용량 확인') is None)
    ok('「작업 의뢰서」 띄어쓰기', T.parse_line('1. 개발 — 매핑 확인 → 작업 의뢰서') == ('개발', '매핑 확인'))
    ok('내용 속 「-」 는 안 자름', T.parse_line('1. 설계 — RS-485 변환 검토 → 작업의뢰서') == ('설계', 'RS-485 변환 검토'))
    ok('부서 묶음 키 (괄호만 뗌)', T.dept_key('설비(난방)') == '설비' and T.dept_key('설계/개발') == '설계/개발' and T.dept_key(' 설계 ') == '설계')
    print('== 1-1. 감사 지적 (v2)')
    ok('「설계/개발 · 매핑」 부서 = 설계/개발', T.parse_line('3. 설계/개발 · 매핑 → 작업의뢰서') == ('설계/개발', '매핑'), T.parse_line('3. 설계/개발 · 매핑 → 작업의뢰서'))
    ok('「전기·통신 — 확인」 부서 = 전기·통신', T.parse_line('7. 전기·통신 — 확인 → 작업의뢰서') == ('전기·통신', '확인'))
    ok('「설비(난방/급수) — 확인」', T.parse_line('13. 설비(난방/급수) — 확인 → 작업의뢰서') == ('설비(난방/급수)', '확인'))
    ok('「2) 제작: 외함 400*900*90」', T.parse_line('2) 제작: 외함 400*900*90 → 작업의뢰서') == ('제작', '외함 400*900*90'))
    ok('「A/B 비교」 내용 속 / 는 안 자름', T.parse_line('1. 설계 — A/B 비교 → 작업의뢰서') == ('설계', 'A/B 비교'))
    for v in ('→ 작업의뢰서 필요', '→ 작업의뢰서.', '-> 작업의뢰서', '(→ 작업의뢰서)', '→작업의뢰서 작성'):
        ok('태그 변형 받음 : %s' % v, T.parse_line('1. 설계 — 회로 비교 ' + v) == ('설계', '회로 비교'), T.parse_line('1. 설계 — 회로 비교 ' + v))
    ok('v3 태그 뒤 괄호 설명을 내용으로', T.parse_line('- 제작 — 외함 → 작업의뢰서 (70~80개, 약 1파렛트)') == ('제작', '외함 (70~80개, 약 1파렛트)'), T.parse_line('- 제작 — 외함 → 작업의뢰서 (70~80개, 약 1파렛트)'))
    ok('v3 「작업의뢰서 불요」 는 의뢰 아님', T.parse_line('- 없음 (확인 단계, 작업의뢰서 불요)') is None and T.NOT_REQ_RE.search('- 없음 (확인 단계, 작업의뢰서 불요)'))
    ok('붙어 있는 복합회의도 현장 아님', not T.is_site('복합회의 (앵커·연합)') and not T.is_site('_삭제요망_가나'))
    ok('사내 부서', T.inhouse('설계') == '디자인&설계' and T.inhouse('AS') == '고객지원' and T.inhouse('통신') == '')
    ok('현장 아님', not T.is_site('복합회의') and not T.is_site('확인 필요') and T.is_site('가나호텔'))
    ok('전화 모양', T.phone_fmt('01012345678') == '010-1234-5678')
    ok('상대 한 줄', T.counterpart({'company': '가나건설', 'name': '홍길동', 'rank': '소장', 'phone': '01012345678'}) == '가나건설 홍길동 소장 (010-1234-5678)')
    ok('회사==이름이면 한 번', T.counterpart({'company': '홍길동', 'name': '홍길동', 'phone': ''}) == '홍길동')
    ok('회사 칸에 안건이 들어온 자료는 버림', T.counterpart({'company': '가, 나, 다', 'name': '홍길동'}) == '홍길동')

    tmp = tempfile.mkdtemp(prefix='t54_')
    md = os.path.join(tmp, 'meta')
    os.makedirs(md)

    def put(fn, j):
        io.open(os.path.join(md, fn), 'w', encoding='utf-8').write(json.dumps(j, ensure_ascii=False))

    put('가나호텔__260921_a__meta.json', meta('가나호텔', '260921', '17:26',
        ['1. 설계 / 회로 비교표 작성 → 작업의뢰서', '2. 제작 — 스위치함 용량 검토 → 작업의뢰서',
         '3. 전기 — 차단기 용량 확인', '4. 설계 · 타입별 수량 정리 → 작업의뢰서'],
        todos=['- 미정 | 회로 비교표 작성 | 배성윤 → 설계', '- 미정 | 비교표 작업의뢰서 작성 | 배성윤 → 설계', '- 260923 | 스위치함 확인 | 배성윤 → 제작', '- 미정 | 차단기 | 배성윤 → 전기'],
        decisions=['댐퍼 전원은 도면에 재반영하기로 함', '없음 — 확인 후 재협의 예정']))
    put('복합회의__260922_b__meta.json', meta('복합회의', '260922', '09:00', ['1. 개발 — 매핑 확인 → 작업의뢰서']))
    put('다라기숙사__260915_c__meta.json', meta('다라기숙사', '260915', '10:00', ['1. 통신 — 카탈로그 확보 → 작업의뢰서']))
    put('마바현장__260921_d__meta.json', meta('마바현장', '260921', '11:00', ['1. 전기 — 확인만']))
    put('_삭제요망__가나호텔__260921_e__meta.json', meta('가나호텔', '260921', '12:00', ['1. 설계 — 지운 것 → 작업의뢰서']))
    put('가나호텔__260921_g__meta.json', meta('가나호텔', '260921', '19:05', ['1. 설계 · 댐퍼 도면 반영 → 작업의뢰서'], name='김철수'))
    put('가나호텔__260921_h__meta.json', meta('가나호텔', '260921', '', ['1. 설계 · 시각 없는 회의 → 작업의뢰서'], name='이영희', phone=''))
    long_lines = ['%d. 설계 — 긴 항목 %d 번째 내용입니다 → 작업의뢰서' % (i, i) for i in range(1, 21)]
    put('사아현장__260920_f__meta.json', meta('사아현장', '260920', '15:00', long_lines))

    print('== 2. 모으기 · 나누기')
    recs, skipped = T.collect(md)
    ok('삭제요망 건너뜀', any('삭제요망' in s[1] for s in skipped) and len(recs) == 7, '%d %s' % (len(recs), skipped))
    ok('날짜 거르기', len(T.collect(md, '260921', '260921')[0]) == 4)
    ps = T.plan(recs)
    ga = [p for p in ps if p['rec']['site'] == '가나호텔' and p['rec']['hm'] == '17:26']
    ok('같은 회의 · 부서별로 장 나눔 (설계·제작)', [p['key'] for p in ga] == ['설계', '제작'], [p['key'] for p in ga])
    ok('설계 장에 2건 (회의록 순서)', ga[0]['items'] == ['회로 비교표 작성', '타입별 수량 정리'])
    ok('표시 없는 전기 줄은 안 들어감', all(p['key'] != '전기' for p in ps))
    ok('작업의뢰서 표시 없는 회의는 장 없음', all(p['rec']['site'] != '마바현장' for p in ps))
    ok('관련 할 일은 같은 부서만', [t[1] for t in ga[0]['todos']] == ['회로 비교표 작성'] and [t[1] for t in ga[1]['todos']] == ['스위치함 확인'])
    bx = [p for p in ps if p['rec']['site'] == '복합회의'][0]
    ok('복합회의 = 현장 아님 표시', not bx['site_ok'])
    tong = [p for p in ps if p['rec']['site'] == '다라기숙사'][0]
    ok('통신 = 사내 부서 아님 표시', tong['inhouse'] == '')

    print('== 3. 본문')
    L = T.body_lines(ga[0])
    txt = [t for _n, t in L]
    ok('첫 줄 요지 (현장 · 날짜 · 상대)', L[0] == (1, '가나호텔 현장 9월 21일 홍길동 소장 협의 결과, 아래 작업을 요청 드립니다.'), L[0])
    ok('요청 순번 2·3', L[1] == (2, '설계 : 회로 비교표 작성') and L[2] == (3, '설계 : 타입별 수량 정리'))
    ok('수량·규격은 빈칸 [   ]', any('수량 : [   ]' in t and '규격·사양 : [   ]' in t for t in txt))
    ok('「없음 — 」 결정은 안 옮김', ' -댐퍼 전원은 도면에 재반영하기로 함' in txt and not any('재협의' in t for t in txt))
    ok('마지막 줄 수고하세요.', L[-1] == (None, '수고하세요.'))
    ok('근거 줄 (전화 번호 없이)', '근거 : 260921 17:26 가나건설 홍길동 소장 협의록' in txt)
    ok('「의뢰서 작성」 할 일은 관련 할 일에서 뺌', all('의뢰서' not in t for t in txt[1:]))
    ok('clear_checks', T.clear_checks('( V )a (V)b ( )c (v )d') == '(   )a ( )b ( )c (  )d')
    ok('복합회의 요지는 [현장 확인]', T.body_lines(bx)[0][1].startswith('[현장 확인] 현장'))
    ok('기한 있는 할 일은 (기한)', ' -스위치함 확인  (기한 260923)' in [t for _n, t in T.body_lines(ga[1])])
    ok('긴 줄 나눔 폭 116 이하', all(T._w(x) <= 116 for x in T.split_line('가' * 100)) and len(T.split_line('가' * 100)) == 2)

    print('== 4. 원틀 채우기 (시험용 서식)')
    tpl = os.path.join(tmp, '작업의뢰서_원틀_시험.xlsx')
    fake_template(tpl)
    out = os.path.join(tmp, 'out')
    rep = T.run(md, out=out, template=tpl, today='260927')
    ok('의뢰서 장 수 = 7 (가나 2+1+1 · 복합 1 · 다라 1 · 사아 1)', rep['sheets'] == 7 and rep['made'] == 7, rep)
    fs = sorted(f for f in os.listdir(out) if f.startswith('작업의뢰서초안_가나호텔_260921') and f.endswith('.xlsx'))
    ok('같은 현장·날·부서 회의 3건 = 파일 3개 (시각·상대로 가름)', fs == ['작업의뢰서초안_가나호텔_260921_1726_설계.xlsx', '작업의뢰서초안_가나호텔_260921_1726_제작.xlsx',
        '작업의뢰서초안_가나호텔_260921_1905_설계.xlsx', '작업의뢰서초안_가나호텔_260921_이영희_%s_설계.xlsx' % _fid('가나호텔__260921_h__meta.json')], fs)
    import openpyxl
    xp = os.path.join(out, '작업의뢰서초안_가나호텔_260921_1726_설계.xlsx')
    ok('파일 이름', os.path.exists(xp), os.listdir(out))
    wb = openpyxl.load_workbook(xp)
    ws = wb.worksheets[0]
    ok('다른 시트(작성예시) 뺌 · 원틀 시트 이름 바꿈', wb.sheetnames == ['가나호텔(설계)'], wb.sheetnames)
    ok('C4 기안일자 = 회의한 날', str(ws['C4'].value)[:10] == '2026-09-21')
    ok('G4 배성윤', ws['G4'].value == '배성윤')
    ok('O4 납기일 빈칸 + 노란칸', ws['O4'].value is None and ws['O4'].fill.fgColor.rgb == T.YELLOW)
    ok('O5·O16 은 O4 를 따라감 (IF)', ws['O5'].value == '=IF($O$4="","",$O$4)' and ws['O16'].value == ws['O5'].value)
    ok('현장명 E12 · G6=E12', ws['E12'].value == '가나호텔' and ws['G6'].value == '=E12')
    ok('D16 상대 · 전화', ws['D16'].value == '가나건설 홍길동 소장 (010-1234-5678)')
    ok('업체명 · 객실수 · 계약No 비움', ws['C6'].value is None and ws['Q6'].value is None and ws['C5'].value is None)
    ok('원틀에 남은 V 를 전부 지움', all('V' not in str(ws[k].value) for k in ('F5', 'C7', 'C8', 'C9', 'C10', 'C14')), [ws[k].value for k in ('F5', 'C10')])
    ok('글자 배치는 그대로 (괄호 폭 유지)', ws['F5'].value == '( )선투입자재 ( )목업룸 ( )기타' and ws['C8'].value == '     ( )  계약 전(샘플)       (   ) 계약')
    ok('유상 ( \\ ) 금액 괄호는 안 건드림', ws['C9'].value.endswith('( \\                 )   '))
    ok('A19=1 · B19 요지', ws['A19'].value == 1 and str(ws['B19'].value).startswith('가나호텔 현장'))
    ok('본문 왼쪽 정렬', ws['B20'].alignment.horizontal == 'left')
    ok('수량·규격 줄 노란칸', any('[   ]' in str(ws['B%d' % r].value or '') and ws['B%d' % r].fill.fgColor.rgb == T.YELLOW for r in range(19, 36)))
    ok('빈 본문 행 정리 → 꼬리 바로 위', ws.print_area and ws.print_area.endswith('$R$%d' % (19 + len(T.body_lines(ga[0])) + 2)), ws.print_area)
    ok('꼬리(양식 MB-004) 살아 있음', any('MB- 004' in str(c.value or '') for row in ws.iter_rows() for c in row))
    xb = os.path.join(out, '작업의뢰서초안_현장확인_260922_0900_개발.xlsx')
    wsb = openpyxl.load_workbook(xb).worksheets[0]
    ok('복합회의 : 현장명 비움 · 시트 현장확인', wsb['E12'].value is None and wsb.title == '현장확인(개발)')
    ok('복합회의 : G6 도 비움 (=E12 → 0 으로 안 보이게)', wsb['G6'].value is None)
    xl = os.path.join(out, '작업의뢰서초안_사아현장_260920_1500_설계.xlsx')
    wsl = openpyxl.load_workbook(xl).worksheets[0]
    n_long = len(T.body_lines([p for p in ps if p['rec']['site'] == '사아현장'][0]))
    ok('17줄 넘으면 행을 늘림 (20건 전부 들어감)', wsl['B%d' % (19 + n_long - 1)].value == '수고하세요.', n_long)
    ok('늘린 행도 B:R 병합', any(str(m) == 'B%d:R%d' % (19 + n_long - 1, 19 + n_long - 1) for m in wsl.merged_cells.ranges))
    ok('늘려도 본문 뒤 빈 줄 2개', wsl.print_area.endswith('$R$%d' % (19 + n_long + 2)), wsl.print_area)

    print('== 5. 덮어쓰지 않음 · 모음 · 대장')
    rep2 = T.run(md, out=out, template=tpl, today='260927')
    ok('두 번째는 전부 「이미 있음」 (덮어쓰지 않음)', rep2['made'] == 0 and rep2['same'] == 7, rep2)
    t = io.open(rep['txt'], encoding='utf-8').read()
    ok('모음 txt 에 부서·요청·수고하세요', '가나호텔 · 9월 21일 · 설계' in t and '회로 비교표 작성' in t and '수고하세요.' in t)
    ok('모음 txt 에 확인할 것 표시', '현장 확인 (복합회의)' in t and '받는 부서 확인 (「통신」' in t)
    import csv as _csv
    rows = list(_csv.reader(io.open(rep['csv'], encoding='utf-8-sig')))
    ok('대장 머리 + 7줄', rows[0][:4] == ['회의날', '현장', '부서', '요청'] and len(rows) == 8, len(rows))
    ok('자가진단 json', os.path.exists(os.path.join(out, '작업의뢰서초안_자가진단_260927.json')))

    print('== 5-1. 감사 지적 : 이름이 순서에 안 흔들림 · 날짜 폴더 · 못 읽은 줄')
    md2 = os.path.join(tmp, 'meta2'); os.makedirs(md2)
    def put2(fn, j):
        io.open(os.path.join(md2, fn), 'w', encoding='utf-8').write(json.dumps(j, ensure_ascii=False))
    put2('b__meta.json', meta('가나호텔', '260921', '', ['1. 설계 — 비 → 작업의뢰서'], phone=''))
    put2('c__meta.json', meta('가나호텔', '260921', '', ['1. 설계 — 씨 → 작업의뢰서'], phone=''))
    base2 = os.path.join(tmp, 'o2')
    r1 = T.run(md2, out=os.path.join(base2, '260927'), template=tpl, today='260927')
    put2('a0__meta.json', meta('가나호텔', '260921', '', ['1. 설계 — 에이 → 작업의뢰서', '2. 제작 → 의뢰서 검토 부탁'], phone=''))
    r2 = T.run(md2, out=os.path.join(base2, '260928'), template=tpl, today='260928')
    ok('첫날 2장 만듦', r1['made'] == 2, r1)
    ok('다음 날 새 회의 a0 만 만듦 (b·c 는 어제 폴더에 있어 「이미 있음」)', r2['made'] == 1 and r2['same'] == 2, r2)
    new = [f for f in os.listdir(os.path.join(base2, '260928')) if f.endswith('.xlsx')]
    ok('새로 만든 것이 a0 의 것', len(new) == 1 and _fid('a0__meta.json') in new[0], new)
    t2 = io.open(r2['txt'], encoding='utf-8').read()
    ok('「의뢰서」 는 있는데 모양이 다른 줄은 모음에 알림', '표시 모양이 달라' in t2 and '의뢰서 검토 부탁' in t2 and r2['odd'] == 1)

    print('== 6. 원틀이 없을 때')
    out3 = os.path.join(tmp, 'out3')
    T.find_template = lambda explicit=None: None
    rep3 = T.run(md, out=out3, today='260927')
    ok('원틀 없으면 xlsx 안 만듦', not any(f.endswith('.xlsx') for f in os.listdir(out3)))
    ok('그래도 본문 txt 는 나옴', '원틀 없음' in io.open(rep3['txt'], encoding='utf-8').read() and '수고하세요.' in io.open(rep3['txt'], encoding='utf-8').read())

    shutil.rmtree(tmp, ignore_errors=True)
    bad = [r for r in R if not r[1]]
    print('')
    print('통과 %d / 실패 %d' % (len(R) - len(bad), len(bad)))
    return 1 if bad else 0


if __name__ == '__main__':
    sys.exit(main())
