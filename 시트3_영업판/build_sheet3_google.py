# 시트3 영업판 「구글판」 빌더 — 구글시트 전용 배열 수식(MAP·LAMBDA·UNIQUE·FILTER·SORT) 판
# 한 칸 수식이 열 전체를 계산하므로 파일이 작다 → 드라이브에 바로 올려 구글시트로 변환해 쓴다.
# 엑셀·LibreOffice 에서는 계산되지 않는다(엑셀용은 build_sheet3.py 의 v2).
# 사용 : python3 build_sheet3_google.py 출력.xlsx [--no-seed] [--seed-json 조사.json] [--verify-tab]
import sys, json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule

VER = "v3(구글판)"; TODAY = "2026-09-27"
N = 1000; E = N + 1          # 붙여넣기 최대 줄 / 2.현장 마지막 행
DS = 300; CS = 200; VR = 500
FMT = 150   # 날짜·숫자 표시 서식을 거는 줄 수(파일 크기 때문에 앞쪽만)
WITH_SEED = "--no-seed" not in sys.argv
SEED_JSON = sys.argv[sys.argv.index("--seed-json") + 1] if "--seed-json" in sys.argv else None
VERIFY = "--verify-tab" in sys.argv
OUT = sys.argv[1]

HF = PatternFill("solid", fgColor="1F3864"); HFont = Font(bold=True, color="FFFFFF")
IN = PatternFill("solid", fgColor="FFF2CC"); AUTO = PatternFill("solid", fgColor="E2EFDA")
T = Font(bold=True, size=14); B = Font(bold=True); GREY = Font(italic=True, color="808080")
WRAP = Alignment(wrap_text=True, vertical="top")
thin = Side(style="thin", color="BFBFBF"); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
DATE = "yyyy-mm-dd"

wb = Workbook()
def sheet(name, first=False):
    ws = wb.active if first else wb.create_sheet(); ws.title = name; return ws
def hdr(ws, row, labels, col=1):
    for i, l in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=l); c.fill = HF; c.font = HFont
        c.alignment = Alignment(wrap_text=True, vertical="center"); c.border = BOX
def widths(ws, m):
    for k, v in m.items(): ws.column_dimensions[k].width = v

P1 = "'1.산군_붙여넣기'"; P2 = "'2.현장'"; P3 = "'3.설계사'"; P3B = "'3b.시공사'"; P4 = "'4.방문기록'"; P5 = "'5.이번주_갈곳'"
ST = "설정"; MEMO = "업체_메모"; RS = "현장_조사"
def c2(col): return f"{P2}!${col}$2:${col}${E}"          # 2.현장 열 범위
def c3(col): return f"{P3}!${col}$4:${col}${3+DS}"
def c3b(col): return f"{P3B}!${col}$4:${col}${3+CS}"
def c4(col): return f"{P4}!${col}$3:${col}${2+VR}"
def mp(ranges, names, body):  # MAP(범위들, LAMBDA(이름들, 식))
    return f"=MAP({','.join(ranges)},LAMBDA({','.join(names)},{body}))"

# ---------------- 설정 행 번호 (엑셀판 v2 와 같은 위치) ----------------
fields = ["주소", "건물명", "연면적", "주용도", "허가구분", "설계사", "시공사", "감리", "건축주", "공사단계", "허가일", "착공일"]
REQUIRED = ["주소", "건물명", "연면적", "주용도", "설계사", "공사단계"]
FROW = {f: 4 + i for i, f in enumerate(fields)}
guess = {"주소": "소재지", "건물명": "현장명", "연면적": "연면적㎡", "주용도": "용도", "허가구분": "구분", "설계사": "건축설계",
         "시공사": "열J(미확인)", "감리": "", "건축주": "", "공사단계": "단계", "허가일": "날짜1(추정:허가일)", "착공일": "날짜2(추정:착공일)"}
num_rows = [
 ("minA", "연면적 최소(㎡) — 이보다 작으면 ✕", 1500, "신규 현장 레이더 숙박 기준(프로님 확정값)"),
 ("rev", "재방문 간격(일)", 14, "0번 목록 32 「다녀오신 지 14일」"),
 ("reg", "이번 주 지역(비우면 전국) — 시·도 첫 단어, 예: 서울특별시", None, "5번 탭을 이 지역만 보기"),
 ("fee", "산군 월 요금(원)", 90000, "2026-09-27 검색 확인"),
 ("start", "산군 사용 시작일", None, "넣으시면 6.본전계산이 누적 비용을 냅니다(9/4 통화: 결제 완료 확인)"),
 ("today", "오늘 날짜", "=TODAY()", "자동"),
 ("year", "산군 연결제 금액(원)", 864000, "2026-09-27 검색 확인(연 20% 할인)"),
 ("star", "★ 연면적 기준(㎡) — 이 이상이고 표류가 아니면 ★", 4000, "9/14 캡처 판정(2,596㎡=△, 4,482㎡=★)에서 역산한 추정값 — 고쳐 주십시오"),
 ("drift", "표류 기준(일) — 허가→착공(착공 없으면 허가→오늘)이 이보다 길면 △", 1200, "9/14 캡처 판정 8건과 모두 맞는 경계값(△ 1,295·1,464일 / ★ 1,083·1,096일) — 추정"),
 ("recent", "착공 「최근」 기준(개월)", 6, "이 안에 착공한 현장은 아직 기구물 사양을 넣을 여지가 있다고 보고 점수를 줌"),
]
wt_rows = [
 ("w_design", "설계단계 현장 1건", 10), ("w_recent", "최근 착공 현장 1건", 6), ("w_old", "오래된 착공 현장 1건", 2),
 ("w_late", "마감·준공 현장 1건", 0), ("w_star", "★ 현장 1건 추가", 5), ("w_new", "안 가본 곳 가산", 30),
 ("w_day", "마지막 방문 뒤 1일당(최대 60일)", 0.5), ("w_area", "최대 연면적 1,000㎡당(최대 20점)", 1), ("w_appt", "7일 안 약속", 100),
]
NUM_START = 19; WT_START = NUM_START + len(num_rows) + 2
KW_START = WT_START + len(wt_rows) + 2; STG_START = KW_START + 12; KWN = 10; STN = 8
NUM = {}
for i, (k, *_r) in enumerate(num_rows): NUM[k] = f"{ST}!$B${NUM_START + i}"
for i, (k, *_r) in enumerate(wt_rows): NUM[k] = f"{ST}!$B${WT_START + i}"
TD = NUM["today"]
KW = f"{ST}!$A${KW_START}:$A${KW_START+KWN-1}"
SK = f"{ST}!$A${STG_START}:$A${STG_START+STN-1}"; SV = f"{ST}!$B${STG_START}:$B${STG_START+STN-1}"

# ================= 검증(시험판에만) =================
if VERIFY:
    vt = sheet("검증", first=True)

# ================= 사용법 =================
u = sheet("사용법", first=not VERIFY)
u["A1"] = f"시트 3번 「영업판」 {VER} ({TODAY}) — 산군 + 설계사 영업 전용 · 구글시트에서 바로 계산"; u["A1"].font = T
lines = [
 "■ 이 파일은 구글시트 전용입니다(열마다 수식 한 칸이 전체를 계산). 폰 구글시트 앱에서도 그대로 보고 고칠 수 있습니다.",
 "■ 프로님 손은 1번 : 매주 월요일 산군 엑셀을 받아 「1.산군_붙여넣기」 탭 A1 칸을 누르고 통째로 붙여넣기(전체 덮어쓰기). 여러 조건으로 받으면 아래로 이어 붙여도 됩니다(같은 주소+현장명은 아래 줄이 이김).",
 "   → 2.현장(★△✕) · 3.설계사 · 3b.시공사 · 5.이번주_갈곳 · 7.문안 · 6.본전계산 이 저절로 바뀝니다.",
 "■ 5.이번주_갈곳 의 「지도」 를 누르면 네이버 지도, 「검색」 을 누르면 사무소 이름+전화 네이버 검색이 열립니다. 「확인할 것」 칸에 이 줄에서 아직 못 믿는 정보가 적혀 있습니다.",
 "■ 처음 한 번만 : 「설정」 탭 B4~B15 에 산군 엑셀 1행의 열 이름을 똑같이 적기. 지금은 9/14 산군 화면 캡처의 열 이름입니다. 「0.자가진단」 1번이 「통과」 면 끝.",
 "■ 지금 1번 탭의 11줄 = 2026-09-14 산군 캡처 실제 자료(드라이브 산군/산군_관심현장 v1). 설계사 이름 2곳은 캡처에서 잘려 「…」 로 끝납니다.",
 "■ 다녀오시면 : 「4.방문기록」 에 한 줄(날짜·업체·결과). 「관심 없음」 이면 갈 곳 목록에서 빠집니다.",
 "■ 담당자·전화 : 「업체_메모」 에 한 줄. 확신도 칸을 「확정」 으로 바꾸시면 5번 「확인할 것」 에서 빠집니다.",
 "■ 현장 뉴스 : 「현장_조사」 탭. 객실 수는 확신도가 「확정」 일 때만 7.문안 메일에 들어갑니다(추정 숫자가 대외로 안 나가게).",
 "■ 판정 : ✕ = 용도 아님 또는 연면적 미달 / △ = 허가→착공 표류 또는 ★ 기준 미만 / ★ = 나머지. 기준·점수 가중치는 설정 탭 노란 칸.",
 "■ 초록 칸(자동)은 지우거나 덮어쓰지 마십시오. 한 칸 수식이 아래 줄까지 채웁니다. 지웠으면 「0.자가진단」 이 실패로 알려 줍니다 → 판 기록에서 되돌리기.",
 "■ 이 시트는 회의록·견적·공정·수금과 섞지 않습니다. 기존 레이더·도구 47번과도 따로입니다.",
]
for i, l in enumerate(lines): u.cell(row=3 + i, column=1, value=l).alignment = WRAP
widths(u, {"A": 130})
d = sheet("0.자가진단")

# ================= 1.산군_붙여넣기 =================
p = sheet("1.산군_붙여넣기")
seed_hdr = ["No", "현장명", "소재지", "공종", "단계", "날짜1(추정:허가일)", "날짜2(추정:착공일)", "건축설계", "열J(미확인)", "용도", "구분", "연면적㎡", "구조"]
for i, h in enumerate(seed_hdr): p.cell(row=1, column=i + 1, value=h).font = B
if WITH_SEED:
    seed = [
     [1, "길상면 선두리 976-7 생활숙박시설", "인천광역시 강화군 길상면 선두리 976-7", "건축", "착공", "2026-06-23", "2026-07-02", "", "", "숙박시설", "신축", 460, "경량철골구조"],
     [2, "범방동 1898-7 생활숙박시설", "부산광역시 강서구 범방동 1898-7번지", "건축", "착공", "2022-04-15", "2025-10-31", "", "", "숙박시설", "신축", 22409, ""],
     [3, "대소면 성본리 115-5 숙박시설", "충청북도 음성군 대소면 성본리 산115-5", "건축", "착공", "2023-09-14", "2026-09-14", "태건축사사무소", "", "숙박시설", "신축", 2596, "철근콘크리트구조"],
     [4, "조양동 1558-3 생활숙박시설", "강원특별자치도 속초시 조양동 1558-3번지", "건축", "착공", "2023-09-27", "2026-09-14", "(주)파인드건축사사무소", "", "숙박시설", "신축", 14868, "철근콘크리트구조"],
     [5, "명동2가 95-1 생활숙박시설 신축공사", "서울특별시 중구 남대문로 64 (명동2가)", "건축", "착공", "2023-08-18", "2026-08-17", "(주)야촌건축사사무소", "(주)시경개발", "숙박시설", "신축", 6095, "철근콘크리트구조"],
     [6, "부산시 중구 남포동3가 생활형숙박시설", "부산광역시 중구 남포길 16 (남포동3가)", "건축", "착공", "2023-07-31", "2026-07-31", "(주)상지이앤에이건축…", "(주)평광이앤씨", "숙박시설", "신축", 14453, "철근콘크리트구조"],
     [7, "와룡동 37 관광숙박시설", "서울특별시 종로구 와룡동 37번지", "건축", "착공", "2026-04-01", "2026-07-30", "오파드건축연구소", "", "숙박시설", "신축", 301, "철근콘크리트구조"],
     [8, "경서동 경서3구역 11블록 2로트 숙박시설", "인천광역시 서구 경서동 블록", "건축", "착공", "2022-07-20", "2026-07-23", "건축사무소 고강", "", "숙박시설", "신축", 1995, "철근콘크리트구조"],
     [9, "종로5가 193-17,3번지 관광숙박시설", "서울특별시 종로구 종로 228 (종로5가)", "건축", "착공", "2026-01-29", "2026-07-15", "(주)아이디어키텍…", "태재연구재단", "숙박시설", "신축", 4482, "철근콘크리트구조"],
     [10, "부전동 숙박시설", "부산광역시 부산진구 중앙대로680번가길", "건축", "착공", "2025-10-20", "2026-07-06", "이진건축사사무소", "", "숙박시설", "신축", 2000, "철근콘크리트구조"],
     [11, "서면 중방대리 106 일반숙박시설", "강원특별자치도 홍천군 서면 고두개길 112", "건축", "착공", "2024-08-08", "2026-06-30", "(주)세움건축사사무소", "", "숙박시설", "신축", 97, "일반목구조"],
    ]
    for r, row in enumerate(seed):
        for c, v in enumerate(row): p.cell(row=2 + r, column=c + 1, value=v)
widths(p, {"A": 5, "B": 30, "C": 34, "D": 6, "E": 6, "F": 12, "G": 12, "H": 22, "I": 14, "J": 10, "K": 6, "L": 9, "M": 14})

# ================= 설정 =================
s = sheet(ST)
s["A1"] = "설정 — 노란 칸만 고치십시오"; s["A1"].font = T
hdr(s, 3, ["항목", "산군 엑셀의 열 이름(1행 글자 그대로, 없으면 비움)", "찾은 열 번호", "상태"])
for f in fields:
    r = FROW[f]
    s.cell(row=r, column=1, value=f + (" ★필수" if f in REQUIRED else ""))
    s.cell(row=r, column=2, value=guess[f] or None).fill = IN
    s.cell(row=r, column=3, value=f'=IF(B{r}="","",IFERROR(MATCH(B{r},{P1}!$1:$1,0),""))').fill = AUTO
    s.cell(row=r, column=4, value=f'=IF(B{r}="","안 씀(비움)",IF(C{r}="","못 찾음 — B칸을 산군 엑셀 1행 글자와 똑같이","정상"))').fill = AUTO
s[f"A{NUM_START-1}"] = "숫자 입력칸"; s[f"A{NUM_START-1}"].font = B
for i, (k, label, v, note) in enumerate(num_rows):
    r = NUM_START + i
    s.cell(row=r, column=1, value=label); c = s.cell(row=r, column=2, value=v)
    c.fill = AUTO if k == "today" else IN
    s.cell(row=r, column=3, value=note).font = GREY
    if k in ("start", "today"): c.number_format = DATE
    if k in ("fee", "year", "minA", "star"): c.number_format = "#,##0"
s[f"A{WT_START-1}"] = "순위 점수 가중치 (5.이번주_갈곳 순서를 정함)"; s[f"A{WT_START-1}"].font = B
for i, (k, label, v) in enumerate(wt_rows):
    r = WT_START + i
    s.cell(row=r, column=1, value=label); s.cell(row=r, column=2, value=v).fill = IN
    s.cell(row=r, column=3, value="Claude 초안값 — 쓰시면서 고쳐 주십시오").font = GREY
s[f"A{KW_START-1}"] = f"대상 용도 키워드(주용도·현장명에 이 글자가 있으면 대상) — {KWN}칸"; s[f"A{KW_START-1}"].font = B
for i, k in enumerate(["숙박", "호텔", "리조트", "콘도", "기숙사", "수련", "연수", "노유자", "실버", "스테이"]):
    s.cell(row=KW_START + i, column=1, value=k).fill = IN
s[f"A{STG_START-1}"] = f"공사단계 키워드 → 누구를 만나나(위에서부터 먼저 맞는 것) — {STN}칸. 안 맞으면 착공일 없음=설계사 / 있음=시공사"; s[f"A{STG_START-1}"].font = B
for i, (k, v) in enumerate([("착공전", "설계사"), ("설계", "설계사"), ("인허가", "설계사"), ("허가", "설계사"), ("착공", "시공사"), ("골조", "시공사"), ("마감", "늦음"), ("준공", "늦음")]):
    s.cell(row=STG_START + i, column=1, value=k).fill = IN; s.cell(row=STG_START + i, column=2, value=v).fill = IN
widths(s, {"A": 56, "B": 30, "C": 60, "D": 40})
dv = DataValidation(type="list", formula1='"설계사,시공사,늦음"', allow_blank=True); s.add_data_validation(dv); dv.add(f"B{STG_START}:B{STG_START+STN-1}")

# ================= 2.현장 (열마다 수식 한 칸) =================
h = sheet("2.현장")
cols = ["원본행", "주소", "현장명", "연면적(㎡)", "주용도", "허가구분", "설계사", "시공사", "감리", "건축주", "공사단계", "허가일", "착공일",
        "시·도", "유효", "키", "최신", "용도맞음", "면적맞음", "대상", "누구를 만나나", "설계사(정리)", "(안 씀)", "허가일원문", "착공일원문", "대상키",
        "판정", "현장점수", "시공사(정리)", "연면적원문"]
hdr(h, 1, cols)
RAWR = f"{P1}!$A$2:$AZ${E}"
def rawcol(f): fr = FROW[f]; return f'=IF({ST}!$C${fr}="","",ARRAYFORMULA(IFERROR(INDEX({RAWR},0,{ST}!$C${fr})&"","")))'
def rng(col): return f"${col}$2:${col}${E}"
h["A2"] = f"=SEQUENCE({N},1,2)"
for col, f in (("B", "주소"), ("C", "건물명"), ("E", "주용도"), ("F", "허가구분"), ("G", "설계사"), ("H", "시공사"), ("I", "감리"),
               ("J", "건축주"), ("K", "공사단계"), ("X", "허가일"), ("Y", "착공일"), ("AD", "연면적")):
    h[f"{col}2"] = rawcol(f)
h["D2"] = mp([rng("AD")], ["x_a"], 'IFERROR(VALUE(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(x_a,",",""),"㎡",""),"m2","")),"")')
pdate = ('IF(x_t="","",IF(AND(LEN(x_t)=8,ISNUMBER(x_t*1)),DATE(LEFT(x_t,4),MID(x_t,5,2),RIGHT(x_t,2)),'
         'IFERROR(DATEVALUE(SUBSTITUTE(SUBSTITUTE(x_t,".","-"),"/","-")),IFERROR(x_t*1,""))))')
h["L2"] = mp([rng("X")], ["x_t"], pdate)
h["M2"] = mp([rng("Y")], ["x_t"], pdate)
h["N2"] = mp([rng("B")], ["x_b"], 'IFERROR(LEFT(x_b,FIND(" ",x_b)-1),x_b)')
h["O2"] = mp([rng("B"), rng("C")], ["x_b", "x_cc"], f'AND(x_b<>"",x_b<>{ST}!$B${FROW["주소"]},x_cc<>{ST}!$B${FROW["건물명"]})')
h["P2"] = mp([rng("B"), rng("C"), rng("O")], ["x_b", "x_cc", "x_o"], 'IF(x_o,x_b&"|"&x_cc,"")')
h["Q2"] = mp([rng("P"), rng("A")], ["x_k", "x_rw"], f'IF(x_k="",FALSE,SUMPRODUCT(({rng("P")}=x_k)*({rng("A")}>x_rw))=0)')
h["R2"] = mp([rng("E"), rng("C")], ["x_e", "x_cc"], f'SUMPRODUCT(({KW}<>"")*ISNUMBER(SEARCH({KW},x_e&" "&x_cc)))>0')
h["S2"] = mp([rng("D")], ["x_d"], f'OR(x_d="",x_d>={NUM["minA"]})')
h["T2"] = mp([rng("Q"), rng("R"), rng("S")], ["x_q", "x_u", "x_s"], "AND(x_q,x_u,x_s)")
h["U2"] = mp([rng("T"), rng("K"), rng("M")], ["x_g", "x_k", "x_m"],
             f'IF(NOT(x_g),"",LET(x_hit,IFERROR(INDEX(FILTER({SV},({SK}<>"")*ISNUMBER(SEARCH({SK},x_k))),1),""),IF(x_hit<>"",x_hit,IF(x_m="","설계사","시공사"))))')
norm = 'SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(x_n,"(주)",""),"㈜",""),"주식회사","")," ","")'
h["V2"] = mp([rng("G")], ["x_n"], norm)
h["Z2"] = mp([rng("T"), rng("V"), rng("U")], ["x_g", "x_v", "x_w"], 'IF(x_g,x_v&"|"&x_w,"")')
h["AA2"] = mp([rng("Q"), rng("R"), rng("S"), rng("D"), rng("L"), rng("M")], ["x_q", "x_u", "x_s", "x_d", "x_l", "x_m"],
              f'IF(NOT(x_q),"",IF(NOT(AND(x_u,x_s)),"✕",IF(OR(AND(x_l<>"",x_m<>"",N(x_m)-N(x_l)>{NUM["drift"]}),'
              f'AND(x_l<>"",x_m="",{TD}-N(x_l)>{NUM["drift"]}),AND(x_d<>"",N(x_d)<{NUM["star"]})),"△","★")))')
h["AB2"] = mp([rng("T"), rng("U"), rng("M"), rng("AA")], ["x_g", "x_w", "x_m", "x_j"],
              f'IF(NOT(x_g),"",IF(x_w="설계사",{NUM["w_design"]},IF(x_w="시공사",IF(OR(x_m="",{TD}-N(x_m)<={NUM["recent"]}*30.4),'
              f'{NUM["w_recent"]},{NUM["w_old"]}),{NUM["w_late"]}))+IF(x_j="★",{NUM["w_star"]},0))')
h["AC2"] = mp([rng("H")], ["x_n"], norm)
for col in "BCDEFGHIJKLMNOPQRSTUVZ": h[f"{col}2"].fill = AUTO
for col in ["AA", "AB", "AC", "AD", "X", "Y"]: h[f"{col}2"].fill = AUTO
for r in range(2, 2 + FMT):
    h[f"L{r}"].number_format = DATE; h[f"M{r}"].number_format = DATE; h[f"D{r}"].number_format = "#,##0"
h.freeze_panes = "C2"
widths(h, {"A": 6, "B": 30, "C": 28, "D": 10, "E": 12, "F": 6, "G": 22, "H": 16, "I": 8, "J": 10, "K": 8, "L": 11, "M": 11,
           "N": 14, "U": 10, "V": 18, "AA": 6, "AB": 8})
for c in ["O", "P", "Q", "R", "S", "T", "W", "X", "Y", "Z", "AC", "AD"]: h.column_dimensions[c].hidden = True
h.conditional_formatting.add(f"AA2:AA{E}", FormulaRule(formula=['AA2="★"'], fill=PatternFill("solid", fgColor="FFE699")))
h.conditional_formatting.add(f"AA2:AA{E}", FormulaRule(formula=['AA2="✕"'], font=Font(color="A6A6A6")))

# ================= 업체_메모 · 현장_조사 · 4.방문기록 (손으로 넣는 탭) =================
m = sheet(MEMO)
m["A1"] = "업체 담당자 메모 — 설계사·시공사 모두. 이름은 3번·3b번 탭 글자와 똑같이"; m["A1"].font = B
hdr(m, 2, ["업체명(시트 표기)", "담당자", "직함", "전화", "이메일", "주소·지역", "메모", "확신도", "출처"])
m["J2"] = "← 노란 칸 대신 2행 아래 아무 줄에나 입력"; m["J2"].font = GREY
widths(m, {"A": 26, "B": 10, "C": 8, "D": 15, "E": 22, "F": 24, "G": 44, "H": 8, "I": 40})
rs = sheet(RS)
rs["A1"] = "현장 조사노트 — 웹검색 결과(출처 포함). 현장명은 2.현장의 현장명과 똑같이. 확신도 : 확정 / 추정 / 미확인"; rs["A1"].font = B
hdr(rs, 2, ["현장명(산군 표기)", "사업명·브랜드", "발주처(시행)", "시공사", "객실 수", "준공 예정", "최신 상황", "확신도", "출처", "조사일"])
widths(rs, {"A": 30, "B": 24, "C": 18, "D": 16, "E": 8, "F": 10, "G": 44, "H": 8, "I": 44, "J": 11})
conf = DataValidation(type="list", formula1='"확정,추정,미확인"', allow_blank=True)
m.add_data_validation(conf); conf.add("H3:H302")
conf2 = DataValidation(type="list", formula1='"확정,추정,미확인"', allow_blank=True)
rs.add_data_validation(conf2); conf2.add("H3:H202")
v = sheet("4.방문기록")
v["A1"] = "방문·통화 기록 — 다녀오시면 한 줄 (날짜·업체·결과만 넣어도 됩니다). 설계사·시공사 모두"; v["A1"].font = B
hdr(v, 2, ["날짜", "업체(설계사·시공사)", "관련 현장", "만난 사람", "들은 것", "다음 약속일", "다음 할 것", "결과", "키(자동)", "목록에없음(자동)"])
for r in range(3, 3 + FMT):
    v[f"A{r}"].number_format = DATE; v[f"F{r}"].number_format = DATE
v["I3"] = mp([f"A3:A{2+VR}", f"B3:B{2+VR}"], ["x_a", "x_b"], 'IF(x_a="","",x_b&"|"&(x_a*1))')
v["J3"] = mp([f"B3:B{2+VR}"], ["x_b"], f'IF(x_b="",0,IF(COUNTIF({c3("B")},x_b)+COUNTIF({c3b("B")},x_b)=0,1,0))')
res = ["설계의뢰 받음", "도면 받음", "자료 요청받음", "재방문", "부재", "관심 없음"]
dvr = DataValidation(type="list", formula1='"' + ",".join(res) + '"', allow_blank=True); v.add_data_validation(dvr); dvr.add(f"H3:H{2+VR}")
for ws_, rg in ((v, f"B3:B{2+VR}"), (m, "A3:A302")):
    dvn = DataValidation(type="list", formula1=f"={c3('B')}", allow_blank=True, showErrorMessage=False); ws_.add_data_validation(dvn); dvn.add(rg)
v.freeze_panes = "A3"
widths(v, {"A": 11, "B": 26, "C": 24, "D": 12, "E": 40, "F": 11, "G": 24, "H": 14})
v.column_dimensions["I"].hidden = True; v.column_dimensions["J"].hidden = True

# 공통 식 조각 (업체명 x_n)
LASTV = f'IF(COUNTIF({c4("B")},x_n)=0,"",MAXIFS({c4("A")},{c4("B")},x_n))'
NEXTA = f'IF(MAXIFS({c4("F")},{c4("B")},x_n)=0,"",MAXIFS({c4("F")},{c4("B")},x_n))'
def memo(col): return f'IFERROR(VLOOKUP(x_n,{MEMO}!$A$3:$I$302,{col},FALSE)&"","")'

# ================= 3.설계사 =================
g = sheet("3.설계사")
g["A1"] = "설계사무소별 모음 (자동) — 산군 대상 현장을 설계사무소 이름으로 묶음"; g["A1"].font = T
g["A2"] = f'="설계사무소 "&COUNTIF(B4:B{3+DS},"?*")&"곳 · 대상 현장 "&COUNTIF({c2("T")},TRUE)&"건 · ★ "&COUNTIF({c2("AA")},"★")&"건"'
hdr(g, 3, ["번호", "설계사무소", "대상 현장", "설계단계", "시공단계", "★ 수", "현장점수", "최대 연면적", "최근 허가일", "최근 착공일", "대표 현장", "지역",
           "담당자", "전화", "마지막 방문", "경과일", "다음 약속", "최근 결과", "상태", "우선점수", "대표 주소", "전화 확신도", "왜 가야 하나", "확인할 것", "전화 첫마디", "지도 주소", "검색 주소", "객실 수(조사)"])
GB = f"B4:B{3+DS}"
def g_rng(col): return f"{col}4:{col}{3+DS}"
V2, T2, U2, Z2 = c2("V"), c2("T"), c2("U"), c2("Z")
g["B4"] = f'=IFERROR(UNIQUE(FILTER({V2},{T2}=TRUE,{V2}<>"")),"")'
g["A4"] = mp([GB, f"SEQUENCE({DS})"], ["x_n", "x_i"], 'IF(x_n="","",x_i)')
g["C4"] = mp([GB], ["x_n"], f'IF(x_n="","",COUNTIFS({V2},x_n,{T2},TRUE))')
g["D4"] = mp([GB], ["x_n"], f'IF(x_n="","",COUNTIFS({V2},x_n,{U2},"설계사"))')
g["E4"] = mp([GB], ["x_n"], f'IF(x_n="","",COUNTIFS({V2},x_n,{U2},"시공사"))')
g["F4"] = mp([GB], ["x_n"], f'IF(x_n="","",COUNTIFS({V2},x_n,{T2},TRUE,{c2("AA")},"★"))')
g["G4"] = mp([GB], ["x_n"], f'IF(x_n="","",SUMIFS({c2("AB")},{V2},x_n,{T2},TRUE))')
g["H4"] = mp([GB], ["x_n"], f'IF(x_n="","",MAXIFS({c2("D")},{V2},x_n,{T2},TRUE))')
g["I4"] = mp([GB], ["x_n"], f'IF(x_n="","",IF(MAXIFS({c2("L")},{V2},x_n,{T2},TRUE)=0,"",MAXIFS({c2("L")},{V2},x_n,{T2},TRUE)))')
g["J4"] = mp([GB], ["x_n"], f'IF(x_n="","",IF(MAXIFS({c2("M")},{V2},x_n,{T2},TRUE)=0,"",MAXIFS({c2("M")},{V2},x_n,{T2},TRUE)))')
REPI = f'IFERROR(MATCH(x_n&"|설계사",{Z2},0),IFERROR(MATCH(x_n&"|시공사",{Z2},0),MATCH(x_n&"|늦음",{Z2},0)))'
g["K4"] = mp([GB], ["x_n"], f'IF(x_n="","",IFERROR(INDEX({c2("C")},{REPI}),""))')
g["L4"] = mp([GB], ["x_n"], f'IF(x_n="","",IFERROR(INDEX({c2("N")},{REPI}),""))')
g["U4"] = mp([GB], ["x_n"], f'IF(x_n="","",IFERROR(INDEX({c2("B")},{REPI}),""))')
g["M4"] = mp([GB], ["x_n"], f'IF(x_n="","",{memo(2)})')
g["N4"] = mp([GB], ["x_n"], f'IF(x_n="","",{memo(4)})')
g["V4"] = mp([GB], ["x_n"], f'IF(x_n="","",{memo(8)})')
g["O4"] = mp([GB], ["x_n"], f'IF(x_n="","",{LASTV})')
g["P4"] = mp([g_rng("O")], ["x_o"], f'IF(x_o="","",{TD}-x_o)')
g["Q4"] = mp([GB], ["x_n"], f'IF(x_n="","",{NEXTA})')
g["R4"] = mp([GB, g_rng("O")], ["x_n", "x_o"], f'IF(x_o="","",IFERROR(INDEX({c4("H")},MATCH(x_n&"|"&(x_o*1),{c4("I")},0))&"",""))')
g["S4"] = mp([GB, g_rng("O"), g_rng("P"), g_rng("R")], ["x_n", "x_o", "x_p", "x_rr"],
             f'IF(x_n="","",IF(x_rr="관심 없음","제외",IF(x_o="","안 가봄",IF(x_p>={NUM["rev"]},x_p&"일 지남","최근 방문"))))')
g["T4"] = mp([GB, g_rng("S"), g_rng("L"), g_rng("Q"), g_rng("O"), g_rng("P"), g_rng("G"), g_rng("H")],
             ["x_n", "x_s", "x_l", "x_q", "x_o", "x_p", "x_g", "x_h"],
             f'IF(x_n="","",IF(x_s="제외","",IF(AND({NUM["reg"]}<>"",x_l<>{NUM["reg"]}),"",'
             f'IF(AND(x_q<>"",x_q>N(x_o),x_q<={TD}+7),{NUM["w_appt"]}+7-(x_q-{TD}),'
             f'IF(AND(x_g>0,OR(x_s="안 가봄",RIGHT(x_s,2)="지남")),'
             f'x_g+IF(x_o="",{NUM["w_new"]},MIN(x_p,60)*{NUM["w_day"]})+MIN(N(x_h)/1000*{NUM["w_area"]},20),"")))))')
RV = lambda col: f'VLOOKUP(x_k,{RS}!$A$3:$J$202,{col},FALSE)'
g["AB4"] = mp([GB, g_rng("K")], ["x_n", "x_k"], f'IF(x_n="","",IFERROR(IF({RV(5)}&""="","",{RV(5)}&IF({RV(8)}="확정",""," ("&{RV(8)}&")")),""))')
g["W4"] = mp([GB, g_rng("T"), g_rng("Q"), g_rng("O"), g_rng("P"), g_rng("D"), g_rng("E"), g_rng("F")],
             ["x_n", "x_t", "x_q", "x_o", "x_p", "x_d", "x_e", "x_f"],
             f'IF(x_n="","",IF(N(x_t)>={NUM["w_appt"]},"약속 "&TEXT(x_q,"m/d")&" · ",IF(x_o="","아직 안 가봄 · ","마지막 방문 "&x_p&"일 전 · "))'
             f'&IF(x_d>0,"설계단계 "&x_d&"건 ","")&IF(x_e>0,"착공 "&x_e&"건 ","")&IF(x_f>0,"★"&x_f,""))')
g["X4"] = mp([GB, g_rng("N"), g_rng("V"), g_rng("AB")], ["x_n", "x_ph", "x_cf", "x_rm"],
             f'IF(x_n="","",REGEXREPLACE(TRIM(IF(ISNUMBER(SEARCH("…",x_n)),"이름 잘림(산군 상세에서 전체 이름) · ","")'
             f'&IF(x_ph="","전화 없음(검색) · ",IF(x_cf<>"확정","전화 "&x_cf&" · ",""))'
             f'&IF(x_rm="","객실 수 미확인 · ",IF(ISNUMBER(SEARCH("(",x_rm)),"객실 수 추정 · ",""))'
             f'&IF(COUNTIFS({V2},x_n,{T2},TRUE,{c2("H")},"")>0,"시공사 미상","")),"\\s*·\\s*$",""))')
g["Y4"] = mp([GB, g_rng("K"), g_rng("D")], ["x_n", "x_k", "x_d"],
             'IF(x_n="","","한국마이크로닉 배성윤 차장입니다. "&x_k&IF(x_d>0," 설계 진행하고 계신"," 설계하신")&" 걸로 알고 연락드렸습니다. 호텔 객실관리(키센서·온도조절기·조명 제어) 도면과 특기시방서를 무상으로 지원해 드리고 있어서, 이번 주에 10분만 찾아뵈어도 될까요?")')
g["Z4"] = mp([GB, g_rng("U")], ["x_n", "x_u"], 'IF(x_n="","","https://map.naver.com/p/search/"&ENCODEURL(x_u))')
g["AA4"] = mp([GB, g_rng("L")], ["x_n", "x_l"], 'IF(x_n="","","https://search.naver.com/search.naver?query="&ENCODEURL(SUBSTITUTE(x_n,"…","")&" "&x_l&" 전화"))')
for c in ["W", "X", "Y", "Z", "AA", "AB"]: g[f"{c}4"].fill = AUTO
for c in "ABCDEFGHIJKLMNOPQRSTUV": g[f"{c}4"].fill = AUTO
for r in range(4, 4 + FMT):
    for c in "IJOQ": g[f"{c}{r}"].number_format = DATE
    g[f"H{r}"].number_format = "#,##0"; g[f"T{r}"].number_format = "0.0"
g.freeze_panes = "C4"
widths(g, {"A": 5, "B": 26, "C": 7, "D": 7, "E": 7, "F": 6, "G": 7, "H": 10, "I": 11, "J": 11, "K": 26, "L": 14, "M": 10, "N": 14,
           "O": 11, "P": 6, "Q": 11, "R": 12, "S": 10, "T": 8, "U": 30, "V": 8, "W": 30, "X": 30, "Y": 40, "Z": 10, "AA": 10, "AB": 10})

# ================= 3b.시공사 =================
c3s = sheet("3b.시공사")
c3s["A1"] = "시공사(또는 발주처)별 모음 (자동) — 착공 현장 영업용. 산군 캡처의 「열J」 가 시공사인지 발주처인지는 미확인"; c3s["A1"].font = T
c3s["A2"] = f'="시공사 "&COUNTIF(B4:B{3+CS},"?*")&"곳"'
hdr(c3s, 3, ["번호", "시공사", "대상 현장", "★ 수", "최대 연면적", "최근 착공일", "대표 현장", "지역", "설계사(같은 현장)", "담당자", "전화", "마지막 방문", "경과일", "최근 결과", "상태"])
CB = f"B4:B{3+CS}"; AC2 = c2("AC")
def cb_rng(col): return f"{col}4:{col}{3+CS}"
c3s["B4"] = f'=IFERROR(UNIQUE(FILTER({AC2},{T2}=TRUE,{AC2}<>"")),"")'
c3s["A4"] = mp([CB, f"SEQUENCE({CS})"], ["x_n", "x_i"], 'IF(x_n="","",x_i)')
c3s["C4"] = mp([CB], ["x_n"], f'IF(x_n="","",COUNTIFS({AC2},x_n,{T2},TRUE))')
c3s["D4"] = mp([CB], ["x_n"], f'IF(x_n="","",COUNTIFS({AC2},x_n,{T2},TRUE,{c2("AA")},"★"))')
c3s["E4"] = mp([CB], ["x_n"], f'IF(x_n="","",MAXIFS({c2("D")},{AC2},x_n,{T2},TRUE))')
c3s["F4"] = mp([CB], ["x_n"], f'IF(x_n="","",IF(MAXIFS({c2("M")},{AC2},x_n,{T2},TRUE)=0,"",MAXIFS({c2("M")},{AC2},x_n,{T2},TRUE)))')
CI = f'MATCH(1,({AC2}=x_n)*({T2}=TRUE),0)'
c3s["G4"] = mp([CB], ["x_n"], f'IF(x_n="","",IFERROR(INDEX({c2("C")},{CI}),""))')
c3s["H4"] = mp([CB], ["x_n"], f'IF(x_n="","",IFERROR(INDEX({c2("N")},{CI}),""))')
c3s["I4"] = mp([CB], ["x_n"], f'IF(x_n="","",IFERROR(INDEX({c2("G")},{CI}),""))')
c3s["J4"] = mp([CB], ["x_n"], f'IF(x_n="","",{memo(2)})')
c3s["K4"] = mp([CB], ["x_n"], f'IF(x_n="","",{memo(4)})')
c3s["L4"] = mp([CB], ["x_n"], f'IF(x_n="","",{LASTV})')
c3s["M4"] = mp([cb_rng("L")], ["x_o"], f'IF(x_o="","",{TD}-x_o)')
c3s["N4"] = mp([CB, cb_rng("L")], ["x_n", "x_o"], f'IF(x_o="","",IFERROR(INDEX({c4("H")},MATCH(x_n&"|"&(x_o*1),{c4("I")},0))&"",""))')
c3s["O4"] = mp([CB, cb_rng("L"), cb_rng("M"), cb_rng("N")], ["x_n", "x_o", "x_p", "x_rr"],
               f'IF(x_n="","",IF(x_rr="관심 없음","제외",IF(x_o="","안 가봄",IF(x_p>={NUM["rev"]},x_p&"일 지남","최근 방문"))))')
for c in "ABCDEFGHIJKLMNO": c3s[f"{c}4"].fill = AUTO
for r in range(4, 4 + FMT):
    for c in "FL": c3s[f"{c}{r}"].number_format = DATE
    c3s[f"E{r}"].number_format = "#,##0"
c3s.freeze_panes = "C4"
widths(c3s, {"A": 5, "B": 22, "C": 7, "D": 6, "E": 10, "F": 11, "G": 28, "H": 14, "I": 22, "J": 10, "K": 14, "L": 11, "M": 6, "N": 12, "O": 10})

# ================= 5.이번주_갈곳 (이름 목록은 한 칸 수식, 나머지는 줄마다 — 링크가 눌리게) =================
w = sheet("5.이번주_갈곳")
w["A1"] = (f'="이번 주 찾아갈 설계사무소 "&COUNTA(FILTER({c3("T")},{c3("T")}<>""))&"곳 · 지역: "&IF({NUM["reg"]}="","전국",{NUM["reg"]})'
           f'&" · 기준일 "&TEXT({TD},"yyyy-mm-dd")')
w["A1"].font = T
w["A2"] = "순서 : ①7일 안 약속 ②현장점수(설계단계>최근 착공>오래된 착공, ★ 가산)+안 가본 곳 ③오래 안 간 곳. 「지도」·「검색」 을 누르면 네이버가 열립니다. 「확인할 것」 = 아직 못 믿는 정보."
hdr(w, 3, ["순위", "설계사무소", "왜 가야 하나", "대표 현장", "지역", "최대 연면적", "객실 수(조사)", "담당자", "전화", "지도", "검색", "확인할 것", "전화 첫마디(복사)", "_행"])
w["B4"] = f'=IFERROR(ARRAY_CONSTRAIN(SORT(FILTER({{{c3("B")},{c3("T")},{c3("A")}}},{c3("T")}<>""),2,FALSE,3,TRUE),30,1),"")'
w["B4"].fill = AUTO
for r in range(4, 34):
    k = r - 3
    w[f"N{r}"] = f'=IF(B{r}="","",MATCH(B{r},{c3("B")},0))'
    def G(col): return f'INDEX({c3(col)},N{r})'
    w[f"A{r}"] = f'=IF(B{r}="","",ROW()-3)'
    for col, src in (("C", "W"), ("D", "K"), ("E", "L"), ("F", "H"), ("G", "AB"), ("H", "M"), ("I", "N"), ("L", "X"), ("M", "Y")):
        w[f"{col}{r}"] = f'=IF(B{r}="","",{G(src)})'
    w[f"J{r}"] = f'=IF(B{r}="","",HYPERLINK({G("Z")},"지도"))'
    w[f"K{r}"] = f'=IF(B{r}="","",HYPERLINK({G("AA")},"검색"))'
    w[f"F{r}"].number_format = "#,##0"
    for c in "CLM": w[f"{c}{r}"].alignment = WRAP
w.column_dimensions["N"].hidden = True
w.freeze_panes = "C4"
widths(w, {"A": 5, "B": 24, "C": 30, "D": 26, "E": 12, "F": 10, "G": 9, "H": 9, "I": 14, "J": 6, "K": 6, "L": 34, "M": 60})

# ================= 6.본전계산 =================
b6 = sheet("6.본전계산")
b6["A1"] = "산군 본전 계산 — 기준 숫자는 프로님이 넣으십시오(노란 칸)"; b6["A1"].font = T
hdr(b6, 3, ["항목", "값", "메모"])
rows6 = [
 ("산군 사용 시작일", f'=IF({NUM["start"]}="","",{NUM["start"]})', "설정 탭"),
 ("쓴 개월 수", f'=IF({NUM["start"]}="","시작일을 넣으십시오",DATEDIF({NUM["start"]},{TD},"m")+1)', ""),
 ("누적 비용(원)", f'=IF(ISNUMBER(B5),B5*{NUM["fee"]},"")', "월 요금 × 개월"),
 ("월결제 1년 vs 연결제 차액(원)", f'={NUM["fee"]}*12-{NUM["year"]}', "연결제로 바꾸면 아끼는 돈"),
 ("산군으로 잡힌 설계사무소(곳)", f'=COUNTIF({c3("B")},"?*")', "자동"),
 ("그중 한 번이라도 간 곳", f'=COUNTIF({c3("O")},">0")', "자동"),
 ("방문·통화 건수", f"=COUNT({c4('A')})", "자동"),
 ("설계의뢰 받음", f'=COUNTIF({c4("H")},"설계의뢰 받음")', "자동"),
 ("도면 받음", f'=COUNTIF({c4("H")},"도면 받음")', "자동"),
 ("방문 1건당 산군 비용(원)", f'=IF(AND(ISNUMBER(B6),B10>0),B6/B10,"")', ""),
 ("★ 현장 수(지금 붙여넣은 자료)", f'=COUNTIF({c2("AA")},"★")', "자동"),
]
for i, (a, f, note) in enumerate(rows6):
    r = 4 + i; b6[f"A{r}"] = a; b6[f"B{r}"] = f; b6[f"B{r}"].fill = AUTO; b6[f"C{r}"] = note
b6["B4"].number_format = DATE
for r in (6, 7, 13): b6[f"B{r}"].number_format = "#,##0"
b6["A17"] = "유지 판단 기준(3개월 기준 권장)"; b6["A17"].font = B
hdr(b6, 18, ["기준", "넣으실 숫자", "실제", "통과"])
for i, (a, ref) in enumerate([("간 설계사무소 최소(곳)", "B9"), ("설계의뢰 받음 최소(건)", "B11"), ("도면 받음 최소(건)", "B12")]):
    r = 19 + i; b6[f"A{r}"] = a; b6[f"B{r}"].fill = IN; b6[f"C{r}"] = f"={ref}"
    b6[f"D{r}"] = f'=IF(B{r}="","",IF(C{r}>=B{r},"통과","미달"))'
b6["A23"] = "판정"; b6["A23"].font = B
b6["B23"] = '=IF(COUNTA(B19:B21)=0,"기준 숫자를 넣으시면 판정합니다",IF(COUNTIF(D19:D21,"미달")=0,"유지","미달 있음 — 해지 또는 조건 변경 검토"))'
widths(b6, {"A": 34, "B": 30, "C": 26, "D": 10})

# ================= 7.문안 =================
t = sheet("7.문안")
t["A1"] = "보낼 문안 — B3 에서 업체를 고르면 채워집니다(비우면 5번 1순위). [ ] 칸만 고쳐서 복사"; t["A1"].font = T
t["A3"] = "업체 고르기"; t["B3"].fill = IN
dvt = DataValidation(type="list", formula1=f"={c3('B')}", allow_blank=True, showErrorMessage=False); t.add_data_validation(dvt); dvt.add("B3")
t["A5"] = "쓰는 이름"; t["B5"] = f'=IF(B3<>"",B3,IFERROR({P5}!B4&"",""))'
rowi = f"MATCH(B5,{c3('B')},0)"
t["A6"] = "대표 현장"; t["B6"] = f'=IFERROR(INDEX({c3("K")},{rowi})&"",IFERROR(INDEX({c3b("G")},MATCH(B5,{c3b("B")},0))&"","[현장명]"))'
mm = lambda col: f'IFERROR(VLOOKUP(B5,{MEMO}!$A$3:$I$302,{col},FALSE)&"","")'
t["A7"] = "담당자"; t["B7"] = f'=IF({mm(2)}="","[담당자]",{mm(2)})'
t["A8"] = "객실 수"; t["B8"] = f'=IFERROR(IF(AND(VLOOKUP(B6,{RS}!$A$3:$J$202,5,FALSE)&""<>"",VLOOKUP(B6,{RS}!$A$3:$J$202,8,FALSE)="확정"),VLOOKUP(B6,{RS}!$A$3:$J$202,5,FALSE)&"","[객실 수]"),"[객실 수]")'
t["C8"] = "현장_조사 확신도가 「확정」 일 때만 자동으로 넣음(추정 숫자가 대외 메일로 나가지 않게)"; t["C8"].font = GREY
t["A9"] = "전화(업체)"; t["B9"] = f'=IF({mm(4)}="","[업체 전화]",{mm(4)})'
for r in range(5, 10): t[f"B{r}"].fill = AUTO
NL = "CHAR(10)"
t["A11"] = "전화 첫마디"
t["B11"] = '="한국마이크로닉 배성윤 차장입니다. "&B6&" 설계하신 걸로 알고 연락드렸습니다. 호텔 객실관리(키센서·온도조절기·조명 제어) 도면과 특기시방서를 무상으로 지원해 드리고 있어서, [날짜]에 10분만 찾아뵈어도 될까요?"'
t["A12"] = "문자(짧게)"
t["B12"] = '="[한국마이크로닉 배성윤 차장] "&B7&"님, "&B6&" 객실관리시스템 설계(도면·특기시방서)를 무상 지원해 드립니다. [날짜] 잠시 찾아뵈어도 될지 여쭙습니다. [내 전화]"'
t["A13"] = "메일 제목"; t["B13"] = '=B6&" 객실관리시스템 설계 지원 안내 — 한국마이크로닉(주)"'
t["A14"] = "메일 본문"
t["B14"] = ('=B5&" "&B7&"님께"&' + NL + '&' + NL + '&"한국마이크로닉(주) 배성윤 차장입니다."&' + NL +
            '&"진행 중이신 "&B6&"("&B8&"실)의 객실관리시스템(RCU·CB함·키센서·온도조절기·조명스위치) 설계를 무상으로 지원해 드립니다."&' + NL + '&' + NL +
            '&"- 단위세대 평면도를 주시면 성급 기준으로 기구물 배치와 예상 수량표를 드립니다."&' + NL +
            '&"- 특기시방서(한글 파일)를 현장에 맞춰 드립니다."&' + NL +
            '&"- 전기설계업체용 설계 요청서도 함께 드립니다."&' + NL + '&' + NL +
            '&"[날짜] 중 편하신 시간에 찾아뵙겠습니다."&' + NL + '&' + NL + '&"한국마이크로닉(주) 배성윤 드림 / [내 전화]"')
t["A15"] = "방문 후 감사 문자"
t["B15"] = '="[한국마이크로닉 배성윤] 오늘 시간 내주셔서 감사합니다. 말씀하신 "&B6&" 자료는 [날짜]까지 보내 드리겠습니다."'
for r in range(11, 16): t[f"B{r}"].alignment = WRAP; t[f"B{r}"].fill = AUTO; t[f"A{r}"].font = B
t.row_dimensions[14].height = 230; t.row_dimensions[11].height = 60; t.row_dimensions[12].height = 45
widths(t, {"A": 18, "B": 100, "C": 50})

# ================= 0.자가진단 =================
d["A1"] = f"자가진단 — 전부 「통과」 면 정상 ({VER})"; d["A1"].font = T
hdr(d, 3, ["번호", "검사", "값", "판정"])
req_cnt = "+".join(f'N({ST}!$C${FROW[f]}<>"")' for f in REQUIRED)
wt_blank = f'COUNTBLANK({ST}!B{WT_START}:B{WT_START+len(wt_rows)-1})'
AAr = c2("AA")
chk = [
 ("필수 열 6개(주소·현장명·연면적·용도·설계사·단계) 중 찾은 수", f"={req_cnt}", 'IF(C{r}=6,"통과","실패 — 설정 B4~B15 를 산군 엑셀 1행과 맞추기")'),
 ("붙여넣은 줄 수(1행 제외)", f"=MAX(COUNTA({P1}!B:B)-1,0)", f'IF(C{{r}}<={N},"통과","실패 — {N}줄 넘음, 아래는 안 읽힘")'),
 ("읽은 현장 줄", f"=COUNTIF({c2('O')},TRUE)", 'IF(C{r}=C5,"통과","확인 — 빈 줄·제목줄 섞임")'),
 ("중복 뺀 현장", f"=COUNTIF({c2('Q')},TRUE)", '"참고"'),
 ("★ + △ + ✕ = 중복 뺀 현장", f'=COUNTIF({AAr},"★")+COUNTIF({AAr},"△")+COUNTIF({AAr},"✕")', 'IF(C{r}=C7,"통과","실패")'),
 ("대상(용도·면적 맞음) = ★ + △", f"=COUNTIF({c2('T')},TRUE)", f'IF(C{{r}}=COUNTIF({AAr},"★")+COUNTIF({AAr},"△"),"통과","실패")'),
 ("설계사행 + 시공사행 + 늦음 = 대상", f'=COUNTIF({c2("U")},"설계사")+COUNTIF({c2("U")},"시공사")+COUNTIF({c2("U")},"늦음")', 'IF(C{r}=C9,"통과","실패 — 설정 단계 행선지 확인")'),
 ("설계사무소 수(목록) = 대상 현장의 설계사 이름 수", f'=COUNTIF({c3("B")},"?*")', f'IF(C{{r}}=IFERROR(COUNTA(UNIQUE(FILTER({c2("V")},{c2("T")}=TRUE,{c2("V")}<>""))),0),IF(C{{r}}<{DS},"통과","실패 — {DS}곳 넘음"),"실패")'),
 ("시공사 수(목록) = 대상 현장의 시공사 이름 수", f'=COUNTIF({c3b("B")},"?*")', f'IF(C{{r}}=IFERROR(COUNTA(UNIQUE(FILTER({c2("AC")},{c2("T")}=TRUE,{c2("AC")}<>""))),0),IF(C{{r}}<{CS},"통과","실패 — {CS}곳 넘음"),"실패")'),
 ("설계사 이름 빈 대상 현장", f'=COUNTIFS({c2("T")},TRUE,{c2("V")},"")', 'IF(C{r}=0,"통과","확인 — 산군에 설계사 없는 현장(2.현장 필터)")'),
 ("방문기록 중 목록에 없는 업체", f"=SUM({c4('J')})", 'IF(C{r}=0,"통과","확인 — 이름이 3·3b번 탭과 다름")'),
 ("점수 가중치 빈칸", f"={wt_blank}", 'IF(C{r}=0,"통과","실패 — 설정 가중치 칸을 채우기")'),
 ("잘린 이름(…) 남음", f'=COUNTIF({c2("G")},"*…*")', 'IF(C{r}=0,"통과","알림 — 캡처에서 잘린 설계사 이름. 엑셀 다운로드로 바꾸면 사라짐")'),
 ("자동 수식 칸 살아 있음(2.현장·3번·3b번)", f'=SUMPRODUCT(--ISFORMULA({P2}!B2:AD2))+ISFORMULA({P3}!B4)+ISFORMULA({P3}!T4)+ISFORMULA({P3B}!B4)', 'IF(C{r}>=27,"통과","실패 — 초록 수식 칸이 지워짐. 판 기록에서 되돌리기")'),
]
for i, (name, val, judge) in enumerate(chk):
    r = 4 + i
    d[f"A{r}"] = i + 1; d[f"B{r}"] = name; d[f"C{r}"] = val; d[f"D{r}"] = "=" + judge.format(r=r)
LAST = 4 + len(chk) - 1
d.conditional_formatting.add(f"D4:D{LAST}", FormulaRule(formula=['LEFT(D4,2)="통과"'], fill=PatternFill("solid", fgColor="C6EFCE")))
d.conditional_formatting.add(f"D4:D{LAST}", FormulaRule(formula=['LEFT(D4,2)="실패"'], fill=PatternFill("solid", fgColor="FFC7CE")))
HR = LAST + 3
d[f"A{HR}"] = "변경 이력"; d[f"A{HR}"].font = B
hdr(d, HR + 1, ["날짜", "판", "내용", "이전 값"])
hist = [
 (TODAY, "v1", "최초 작성 — 엑셀판 11개 탭, [예시] 가짜 7줄", "없음"),
 (TODAY, "v2", "엑셀판 — 9/14 산군 캡처 11줄, ★△✕, 가중치 입력칸, 3b.시공사·현장_조사·업체_메모", "v1"),
 (TODAY, "v3(구글판)", "B: 계산 방식 → 구글시트 배열 수식(MAP·UNIQUE·SORT). 드라이브에 바로 쓰는 구글시트로 올림. 탭·열·규칙·숫자는 v2 와 같음", "엑셀판 줄마다 수식(v2)"),
 (TODAY, "v3(구글판)", "A: 5.이번주_갈곳 에 지도·검색 링크, 「확인할 것」 칸 / 3.설계사 에 대표 주소·전화 확신도 / 자가진단 14번(수식 칸 살아 있음)", "—"),
]
for i, row in enumerate(hist):
    for c, vv in enumerate(row): d.cell(row=HR + 2 + i, column=c + 1, value=vv)
widths(d, {"A": 12, "B": 44, "C": 70, "D": 40})

# ================= 조사 씨앗 =================
if SEED_JSON:
    J = json.load(open(SEED_JSON))
    for i, row in enumerate(J.get("memo", [])):
        for c, vv in enumerate(row): m.cell(row=3 + i, column=c + 1, value=vv)
    for i, row in enumerate(J.get("research", [])):
        for c, vv in enumerate(row): rs.cell(row=3 + i, column=c + 1, value=vv)

# ================= 검증 탭 (시험판에만 — 값을 한 줄 글자로 모아 CSV 로 읽는다) =================
if VERIFY:
    vt["A1"] = f'=TEXTJOIN("|",TRUE,{P5}!B4:B33)'
    vt["A2"] = ('=TEXTJOIN(";",TRUE,' + f'MAP({c3("B")},{c3("C")},{c3("D")},{c3("E")},{c3("F")},{c3("G")},{c3("H")},{c3("K")},{c3("L")},{c3("S")},{c3("T")},'
                'LAMBDA(x_n,x_c,x_d,x_e,x_f,x_g,x_h,x_k,x_l,x_s,x_t,IF(x_n="","",x_n&":"&x_c&":"&x_d&":"&x_e&":"&x_f&":"&x_g&":"&x_h&":"&x_k&":"&x_l&":"&x_s&":"&IF(x_t="","",TEXT(x_t,"0.000"))))))')
    vt["A3"] = f'=COUNTIF({AAr},"★")&","&COUNTIF({AAr},"△")&","&COUNTIF({AAr},"✕")'
    vt["A4"] = f'=TEXTJOIN(";",TRUE,MAP({c3b("B")},{c3b("C")},LAMBDA(x_n,x_c,IF(x_n="","",x_n&":"&x_c))))'
    vt["A5"] = f"=TEXTJOIN(\"|\",TRUE,'0.자가진단'!D4:D17)"
    vt["A6"] = f"=TEXTJOIN(\"|\",TRUE,{P5}!L4:L10)"

order = (["검증"] if VERIFY else []) + ["사용법", "0.자가진단", "5.이번주_갈곳", "1.산군_붙여넣기", "3.설계사", "3b.시공사", "4.방문기록", "7.문안",
                                        "업체_메모", "현장_조사", "2.현장", "6.본전계산", "설정"]
wb._sheets = [wb[n] for n in order]
wb.active = order.index("5.이번주_갈곳")
wb.save(OUT); print("saved", OUT)
