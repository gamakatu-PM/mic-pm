# 구글판 xlsx(build_sheet3_google.py 결과) → 구글시트 API 요청(JSON) 으로 바꾼다.
# 파일 업로드 없이 Zapier 의 Google Sheets API 요청으로 시트를 한 칸씩 쓰기 위함.
# 사용 : python3 gen_api.py 구글판.xlsx 출력폴더
import sys, json, os, re
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, column_index_from_string
SRC, OUTD = sys.argv[1], sys.argv[2]; os.makedirs(OUTD, exist_ok=True)
wb = load_workbook(SRC)
names = wb.sheetnames
SID = {n: (0 if i == 0 else 100 + i) for i, n in enumerate(names)}
ROWS = {"1.산군_붙여넣기": 1010, "2.현장": 1010, "4.방문기록": 510}
COLS = {"1.산군_붙여넣기": 52, "2.현장": 30, "3.설계사": 28}
COPY_ROWS = {"5.이번주_갈곳": (4, 5, 33)}   # 원본 행, 붙여넣기 시작~끝
def rgb(hexs):
    h = hexs[-6:]; return {"red": round(int(h[0:2], 16) / 255, 3), "green": round(int(h[2:4], 16) / 255, 3), "blue": round(int(h[4:6], 16) / 255, 3)}
def q(n): return "'" + n.replace("'", "''") + "'"

# 1) 구조 : 시트 이름·순서·크기·고정
struct = []
for i, n in enumerate(names):
    ws = wb[n]
    fr = fc = 0
    if ws.freeze_panes:
        m = re.match(r"([A-Z]+)(\d+)", ws.freeze_panes); fc = column_index_from_string(m.group(1)) - 1; fr = int(m.group(2)) - 1
    gp = {"rowCount": ROWS.get(n, 1000), "columnCount": COLS.get(n, 26), "frozenRowCount": fr, "frozenColumnCount": fc}
    props = {"sheetId": SID[n], "title": n, "index": i, "gridProperties": gp}
    if i == 0:
        struct.append({"updateSheetProperties": {"properties": props, "fields": "title,index,gridProperties"}})
    else:
        struct.append({"addSheet": {"properties": props}})
json.dump({"requests": struct}, open(f"{OUTD}/1_구조.json", "w"), ensure_ascii=False, separators=(",", ":"))

# 2) 값·수식 : 시트마다 한 파일 (행 단위 범위, 빈칸은 null = 건너뜀)
for n in names:
    ws = wb[n]; data = []
    skip = COPY_ROWS.get(n)
    for row in ws.iter_rows():
        cells = [(c.column, c.value) for c in row if c.value is not None]
        if not cells: continue
        r = row[0].row
        if skip and skip[1] <= r <= skip[2]: continue
        c1, c2 = cells[0][0], cells[-1][0]
        vals = [None] * (c2 - c1 + 1)
        for col, v in cells:
            if hasattr(v, "isoformat"): v = v.strftime("%Y-%m-%d")
            vals[col - c1] = v
        data.append({"range": f"{q(n)}!{get_column_letter(c1)}{r}:{get_column_letter(c2)}{r}", "values": [vals]})
    json.dump({"valueInputOption": "USER_ENTERED", "data": data}, open(f"{OUTD}/2_값_{names.index(n):02d}.json", "w"), ensure_ascii=False, separators=(",", ":"))

# 3) 서식 : 채우기·글꼴·줄바꿈·숫자서식 (같은 서식 가로 연속 → 세로 연속으로 묶음), 열 너비·숨김, 드롭다운, 조건부 서식, 줄 복사
def fmt_of(c):
    f = {}
    if c.fill is not None and c.fill.fgColor is not None and c.fill.fill_type == "solid" and isinstance(c.fill.fgColor.rgb, str) and c.fill.fgColor.rgb[-6:] != "000000":
        f["backgroundColor"] = rgb(c.fill.fgColor.rgb)
    tf = {}
    if c.font is not None:
        if c.font.b: tf["bold"] = True
        if c.font.i: tf["italic"] = True
        if c.font.sz and float(c.font.sz) != 11: tf["fontSize"] = int(float(c.font.sz))
        if c.font.color is not None and isinstance(c.font.color.rgb, str) and c.font.color.rgb[-6:] not in ("000000",):
            tf["foregroundColor"] = rgb(c.font.color.rgb)
    if tf: f["textFormat"] = tf
    if c.alignment is not None and c.alignment.wrap_text: f["wrapStrategy"] = "WRAP"
    if c.number_format and c.number_format != "General":
        nf = c.number_format
        f["numberFormat"] = {"type": "DATE", "pattern": "yyyy-mm-dd"} if "yy" in nf else {"type": "NUMBER", "pattern": nf}
    return json.dumps(f, sort_keys=True) if f else None
reqs = []
for n in names:
    ws = wb[n]; sid = SID[n]
    runs = {}   # (c1,c2,fmt) -> list of rows
    for row in ws.iter_rows():
        cur = None
        for c in row:
            k = fmt_of(c)
            if cur and k == cur[2] and c.column == cur[1] + 1:
                cur[1] = c.column
            else:
                if cur and cur[2]: runs.setdefault((cur[0], cur[1], cur[2]), []).append(row[0].row)
                cur = [c.column, c.column, k]
        if cur and cur[2]: runs.setdefault((cur[0], cur[1], cur[2]), []).append(row[0].row)
    for (c1, c2, k), rows_ in runs.items():
        rows_.sort(); start = prev = rows_[0]
        for r in rows_[1:] + [None]:
            if r is not None and r == prev + 1: prev = r; continue
            reqs.append({"repeatCell": {"range": {"sheetId": sid, "startRowIndex": start - 1, "endRowIndex": prev, "startColumnIndex": c1 - 1, "endColumnIndex": c2},
                                        "cell": {"userEnteredFormat": json.loads(k)}, "fields": "userEnteredFormat(" + ",".join(json.loads(k).keys()) + ")"}})
            if r is not None: start = prev = r
    cols_ = sorted((column_index_from_string(k) - 1, int(dm.width * 7) if dm.width else None, bool(dm.hidden)) for k, dm in ws.column_dimensions.items())
    grp = []
    for ci, wd, hd in cols_:
        if wd is None and not hd: continue
        if grp and grp[-1][1] == ci and grp[-1][2] == wd and grp[-1][3] == hd: grp[-1][1] = ci + 1
        else: grp.append([ci, ci + 1, wd, hd])
    for a, b_, wd, hd in grp:
        props = {}; fields = []
        if wd: props["pixelSize"] = wd; fields.append("pixelSize")
        if hd: props["hiddenByUser"] = True; fields.append("hiddenByUser")
        reqs.append({"updateDimensionProperties": {"range": {"sheetId": sid, "dimension": "COLUMNS", "startIndex": a, "endIndex": b_}, "properties": props, "fields": ",".join(fields)}})
    for key, dim in ws.row_dimensions.items():
        if dim.height:
            reqs.append({"updateDimensionProperties": {"range": {"sheetId": sid, "dimension": "ROWS", "startIndex": key - 1, "endIndex": key}, "properties": {"pixelSize": int(dim.height * 1.33)}, "fields": "pixelSize"}})
    for dv in ws.data_validations.dataValidation:
        f1 = dv.formula1
        if f1.startswith('"'):
            cond = {"type": "ONE_OF_LIST", "values": [{"userEnteredValue": x} for x in f1.strip('"').split(",")]}
        else:
            cond = {"type": "ONE_OF_RANGE", "values": [{"userEnteredValue": f1 if f1.startswith("=") else "=" + f1}]}
        for rg in str(dv.sqref).split():
            m = re.match(r"([A-Z]+)(\d+)(?::([A-Z]+)(\d+))?", rg)
            c1, r1 = column_index_from_string(m.group(1)), int(m.group(2))
            c2, r2 = (column_index_from_string(m.group(3)), int(m.group(4))) if m.group(3) else (c1, r1)
            reqs.append({"setDataValidation": {"range": {"sheetId": sid, "startRowIndex": r1 - 1, "endRowIndex": r2, "startColumnIndex": c1 - 1, "endColumnIndex": c2},
                                               "rule": {"condition": cond, "strict": bool(dv.showErrorMessage), "showCustomUi": True}}})
# 조건부 서식 (빌더와 같은 규칙)
def cf(sheet, col, r1, r2, formula, fmt):
    ci = column_index_from_string(col) - 1
    return {"addConditionalFormatRule": {"index": 0, "rule": {"ranges": [{"sheetId": SID[sheet], "startRowIndex": r1 - 1, "endRowIndex": r2, "startColumnIndex": ci, "endColumnIndex": ci + 1}],
            "booleanRule": {"condition": {"type": "CUSTOM_FORMULA", "values": [{"userEnteredValue": formula}]}, "format": fmt}}}}
reqs.append(cf("2.현장", "AA", 2, 1001, '=AA2="★"', {"backgroundColor": rgb("FFE699")}))
reqs.append(cf("2.현장", "AA", 2, 1001, '=AA2="✕"', {"textFormat": {"foregroundColor": rgb("A6A6A6")}}))
reqs.append(cf("0.자가진단", "D", 4, 17, '=LEFT(D4,2)="통과"', {"backgroundColor": rgb("C6EFCE")}))
reqs.append(cf("0.자가진단", "D", 4, 17, '=LEFT(D4,2)="실패"', {"backgroundColor": rgb("FFC7CE")}))
# 5번 탭 : 4행 수식을 5~33행에 붙여넣기(상대 참조가 줄마다 바뀜)
for n, (src, a, b) in COPY_ROWS.items():
    for c1, c2 in ((0, 1), (2, 14)):   # B열(이름 목록)은 한 칸 수식이 아래로 채우므로 복사하지 않음
        reqs.append({"copyPaste": {"source": {"sheetId": SID[n], "startRowIndex": src - 1, "endRowIndex": src, "startColumnIndex": c1, "endColumnIndex": c2},
                                   "destination": {"sheetId": SID[n], "startRowIndex": a - 1, "endRowIndex": b, "startColumnIndex": c1, "endColumnIndex": c2},
                                   "pasteType": "PASTE_NORMAL"}})
json.dump({"requests": reqs}, open(f"{OUTD}/3_서식.json", "w"), ensure_ascii=False, separators=(",", ":"))
print("sheets", len(names), "format reqs", len(reqs))
for f in sorted(os.listdir(OUTD)): print(f, os.path.getsize(f"{OUTD}/{f}"))
