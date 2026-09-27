# 시트3 영업판 빌더 — 산군 엑셀 → 현장 판정(★△✕) → 설계사·시공사별 집계 → 이번 주 갈 곳 (수식만, 매크로 없음)
# 사용 : python3 build_sheet3.py 출력.xlsx [--no-seed] [--seed-json 조사.json]
import sys, json
from datetime import date
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.workbook.properties import CalcProperties

VER = "v2"; TODAY = "2026-09-27"
N = 1000            # 붙여넣기 최대 줄 수
R2 = N + 1          # 2.현장 마지막 행
DS = 300            # 설계사 최대 수
CS = 200            # 시공사 최대 수
VR = 500            # 방문기록 줄 수
WITH_SEED = "--no-seed" not in sys.argv
SEED_JSON = sys.argv[sys.argv.index("--seed-json") + 1] if "--seed-json" in sys.argv else None
OUT = sys.argv[1]

HF = PatternFill("solid", fgColor="1F3864"); HFont = Font(bold=True, color="FFFFFF")
IN = PatternFill("solid", fgColor="FFF2CC")   # 입력칸 노랑
AUTO = PatternFill("solid", fgColor="E2EFDA") # 자동 초록
T = Font(bold=True, size=14); B = Font(bold=True); GREY = Font(italic=True, color="808080")
WRAP = Alignment(wrap_text=True, vertical="top")
thin = Side(style="thin", color="BFBFBF"); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
DATE = "yyyy-mm-dd"

wb = Workbook()
def sheet(name, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = name; return ws
def hdr(ws, row, labels, col=1):
    for i, l in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=l); c.fill = HF; c.font = HFont
        c.alignment = Alignment(wrap_text=True, vertical="center"); c.border = BOX
def widths(ws, ws_w):
    for k, v in ws_w.items(): ws.column_dimensions[k].width = v

P1 = "'1.산군_붙여넣기'"; P2 = "'2.현장'"; P3 = "'3.설계사'"; P3B = "'3b.시공사'"; P4 = "'4.방문기록'"
ST = "설정"; MEMO = "업체_메모"; RS = "현장_조사"

# ================= 설정 (다른 탭이 참조하므로 행 번호를 먼저 정한다) =================
fields = ["주소", "건물명", "연면적", "주용도", "허가구분", "설계사", "시공사", "감리", "건축주", "공사단계", "허가일", "착공일"]
REQUIRED = ["주소", "건물명", "연면적", "주용도", "설계사", "공사단계"]
FROW = {f: 4 + i for i, f in enumerate(fields)}
# 2026-09-14 산군 화면 캡처(드라이브 산군_관심현장 v1)의 열 이름. 엑셀 다운로드의 열 이름은 아직 미확인.
guess = {"주소": "소재지", "건물명": "현장명", "연면적": "연면적㎡", "주용도": "용도", "허가구분": "구분", "설계사": "건축설계",
         "시공사": "열J(미확인)", "감리": "", "건축주": "", "공사단계": "단계", "허가일": "날짜1(추정:허가일)", "착공일": "날짜2(추정:착공일)"}
NUM = {}   # 이름 → 설정 셀 주소
num_rows = [
 ("minA", "연면적 최소(㎡) — 이보다 작으면 ✕", 1500, "신규 현장 레이더 숙박 기준(프로님 확정값)"),
 ("rev", "재방문 간격(일)", 14, "0번 목록 32 「다녀오신 지 14일」"),
 ("reg", "이번 주 지역(비우면 전국) — 시·도 첫 단어, 예: 서울특별시", None, "5번 탭을 이 지역만 보기"),
 ("fee", "산군 월 요금(원)", 90000, "2026-09-27 검색 확인"),
 ("start", "산군 사용 시작일", None, "넣으시면 6.본전계산이 누적 비용을 냅니다(9/4 통화: 결제 완료 확인)"),
 ("today", "오늘 날짜", "=TODAY()", "자동"),
 ("year", "산군 연결제 금액(원)", 864000, "2026-09-27 검색 확인(연 20% 할인)"),
 ("star", "★ 연면적 기준(㎡) — 이 이상이고 표류가 아니면 ★", 4000, "9/14 캡처 판정(2,596㎡=△, 4,482㎡=★)에서 역산한 추정값 — 고쳐 주십시오"),
 ("drift", "표류 기준(일) — 허가→착공(착공 없으면 허가→오늘)이 이보다 길면 △", 1200, "9/14 캡처 판정 8건과 모두 맞는 경계값(△ 1,295·1,464일 / ★ 1,083·1,096일) — 추정"),
 ("recent", "착공 「최근」 기준(개월)", 6, "이 안에 착공한 현장은 아직 기구물 사양을 넣을 여지가 있다고 보고 점수를 줌"),
]
wt_rows = [
 ("w_design", "설계단계 현장 1건", 10), ("w_recent", "최근 착공 현장 1건", 6), ("w_old", "오래된 착공 현장 1건", 2),
 ("w_late", "마감·준공 현장 1건", 0), ("w_star", "★ 현장 1건 추가", 5), ("w_new", "안 가본 곳 가산", 30),
 ("w_day", "마지막 방문 뒤 1일당(최대 60일)", 0.5), ("w_area", "최대 연면적 1,000㎡당(최대 20점)", 1), ("w_appt", "7일 안 약속", 100),
]
NUM_START = 19; WT_START = NUM_START + len(num_rows) + 2
KW_START = WT_START + len(wt_rows) + 2; STG_START = KW_START + 12
for i, (k, *_r) in enumerate(num_rows): NUM[k] = f"{ST}!$B${NUM_START + i}"
for i, (k, *_r) in enumerate(wt_rows): NUM[k] = f"{ST}!$B${WT_START + i}"
KWN = 10; STN = 8
TD = NUM["today"]

# ================= 사용법 =================
u = sheet("사용법", first=True)
u["A1"] = f"시트 3번 「영업판」 {VER} ({TODAY}) — 산군 + 설계사 영업 전용"; u["A1"].font = T
lines = [
 "■ 프로님 손은 1번 : 매주 월요일 산군 엑셀을 받아 「1.산군_붙여넣기」 탭 A1 칸을 누르고 통째로 붙여넣기(전체 덮어쓰기). 여러 조건으로 받으면 아래로 이어 붙여도 됩니다(같은 주소+현장명은 아래 줄이 이김).",
 "   → 2.현장(★△✕ 판정) · 3.설계사 · 3b.시공사 · 5.이번주_갈곳 · 7.문안 · 6.본전계산 이 저절로 바뀝니다.",
 "■ 처음 한 번만 : 「설정」 탭 B4~B15 에 산군 엑셀 1행의 열 이름을 똑같이 적기. 지금은 9/14 산군 화면 캡처의 열 이름입니다. 「0.자가진단」 1번이 「통과」 면 끝.",
 "■ 지금 1번 탭에 들어 있는 11줄 = 2026-09-14 산군 캡처 실제 자료(드라이브 산군/산군_관심현장 v1). 설계사 이름 2곳은 캡처에서 잘려 「…」 로 끝납니다.",
 "■ 다녀오시면 : 「4.방문기록」 에 한 줄(날짜·업체·결과). 결과를 「관심 없음」 으로 고르면 갈 곳 목록에서 빠집니다. 설계사·시공사 모두 여기에 적습니다.",
 "■ 담당자·전화 : 「업체_메모」 에 한 줄. 웹검색으로 찾은 대표전화는 출처와 함께 미리 넣어 두었습니다(확신도 칸 확인).",
 "■ 현장 뉴스 : 「현장_조사」 탭 = 발주처·시공사·객실 수·최신 상황(웹검색, 출처 포함). 객실 수는 7.문안에 자동으로 들어갑니다.",
 "■ 전화·메일 문안 : 「7.문안」 B3 에서 업체를 고르면(비우면 5번 1순위) 전화 첫마디·문자·메일·감사 문자가 채워집니다.",
 "■ 판정 : ✕ = 용도 아님 또는 연면적 미달 / △ = 허가→착공 표류 또는 ★ 기준 미만 / ★ = 나머지. 기준 숫자는 설정 탭.",
 "■ 누구를 만나나 : 허가·설계 단계 → 설계사무소 / 착공 이후 → 시공사(설계사도 최근 착공이면 점수 있음) / 마감·준공 → 늦음.",
 "■ 순위 점수 가중치는 전부 「설정」 탭 노란 칸입니다. 노란 칸 = 프로님 입력 / 초록 칸 = 자동(건드리지 않음).",
 "■ 이 시트는 회의록·견적·공정·수금과 섞지 않습니다. 기존 레이더·도구 47번과도 따로입니다(계획서 9절).",
 "■ 바꿀 때는 덮어쓰지 않고 새 판으로 냅니다. 변경 이력은 「0.자가진단」 아래.",
]
for i, l in enumerate(lines): u.cell(row=3 + i, column=1, value=l).alignment = WRAP
widths(u, {"A": 130})

d = sheet("0.자가진단")

# ================= 1.산군_붙여넣기 =================
p = sheet("1.산군_붙여넣기")
seed_hdr = ["No", "현장명", "소재지", "공종", "단계", "날짜1(추정:허가일)", "날짜2(추정:착공일)", "건축설계", "열J(미확인)", "용도", "구분", "연면적㎡", "구조"]
for i, h in enumerate(seed_hdr): p.cell(row=1, column=i + 1, value=h).font = B
if WITH_SEED:
    seed = [
     [1, "길상면 선두리 976-7 생활숙박시설", "인천광역시 강화군 길상면 선두리 976-7", "건축", "착공", "2026-06-23", "2026-07-02", "", "", "숙박시설", "신축", 460, "경량철골구조"],
     [2, "범방동 1898-7 생활숙박시설", "부산광역시 강서구 범방동 1898-7번지", "건축", "착공", "2022-04-15", "2025-10-31", "", "", "숙박시설", "신축", 22409, ""],
     [3, "대소면 성본리 115-5 숙박시설", "충청북도 음성군 대소면 성본리 산115-5", "건축", "착공", "2023-09-14", "2026-09-14", "태건축사사무소", "", "숙박시설", "신축", 2596, "철근콘크리트구조"],
     [4, "조양동 1558-3 생활숙박시설", "강원특별자치도 속초시 조양동 1558-3번지", "건축", "착공", "2023-09-27", "2026-09-14", "(주)파인드건축사사무소", "", "숙박시설", "신축", 14868, "철근콘크리트구조"],
     [5, "명동2가 95-1 생활숙박시설 신축공사", "서울특별시 중구 남대문로 64 (명동2가)", "건축", "착공", "2023-08-18", "2026-08-17", "(주)야촌건축사사무소", "(주)시경개발", "숙박시설", "신축", 6095, "철근콘크리트구조"],
     [6, "부산시 중구 남포동3가 생활형숙박시설", "부산광역시 중구 남포길 16 (남포동3가)", "건축", "착공", "2023-07-31", "2026-07-31", "(주)상지이앤에이건축…", "(주)평광이앤씨", "숙박시설", "신축", 14453, "철근콘크리트구조"],
     [7, "와룡동 37 관광숙박시설", "서울특별시 종로구 와룡동 37번지", "건축", "착공", "2026-04-01", "2026-07-30", "오파드건축연구소", "", "숙박시설", "신축", 301, "철근콘크리트구조"],
     [8, "경서동 경서3구역 11블록 2로트 숙박시설", "인천광역시 서구 경서동 블록", "건축", "착공", "2022-07-20", "2026-07-23", "건축사무소 고강", "", "숙박시설", "신축", 1995, "철근콘크리트구조"],
     [9, "종로5가 193-17,3번지 관광숙박시설", "서울특별시 종로구 종로 228 (종로5가)", "건축", "착공", "2026-01-29", "2026-07-15", "(주)아이디어키텍…", "태재연구재단", "숙박시설", "신축", 4482, "철근콘크리트구조"],
     [10, "부전동 숙박시설", "부산광역시 부산진구 중앙대로680번가길", "건축", "착공", "2025-10-20", "2026-07-06", "이진건축사사무소", "", "숙박시설", "신축", 2000, "철근콘크리트구조"],
     [11, "서면 중방대리 106 일반숙박시설", "강원특별자치도 홍천군 서면 고두개길 112", "건축", "착공", "2024-08-08", "2026-06-30", "(주)세움건축사사무소", "", "숙박시설", "신축", 97, "일반목구조"],
    ]
    for r, row in enumerate(seed):
        for c, v in enumerate(row): p.cell(row=2 + r, column=c + 1, value=v)
widths(p, {"A": 5, "B": 30, "C": 34, "D": 6, "E": 6, "F": 12, "G": 12, "H": 22, "I": 14, "J": 10, "K": 6, "L": 9, "M": 14})

# ================= 설정 =================
s = sheet(ST)
s["A1"] = "설정 — 노란 칸만 고치십시오"; s["A1"].font = T
hdr(s, 3, ["항목", "산군 엑셀의 열 이름(1행 글자 그대로, 없으면 비움)", "찾은 열 번호", "상태"])
for f in fields:
    r = FROW[f]
    s.cell(row=r, column=1, value=f + (" ★필수" if f in REQUIRED else ""))
    c = s.cell(row=r, column=2, value=guess[f] or None); c.fill = IN
    s.cell(row=r, column=3, value=f'=IF(B{r}="","",IFERROR(MATCH(B{r},{P1}!$1:$1,0),""))').fill = AUTO
    s.cell(row=r, column=4, value=f'=IF(B{r}="","안 씀(비움)",IF(C{r}="","못 찾음 — B칸을 산군 엑셀 1행 글자와 똑같이","정상"))').fill = AUTO
s[f"A{NUM_START-1}"] = "숫자 입력칸"; s[f"A{NUM_START-1}"].font = B
for i, (k, label, v, note) in enumerate(num_rows):
    r = NUM_START + i
    s.cell(row=r, column=1, value=label); c = s.cell(row=r, column=2, value=v)
    c.fill = AUTO if k == "today" else IN
    s.cell(row=r, column=3, value=note).font = GREY
    if k in ("start", "today"): c.number_format = DATE
    if k in ("fee", "year", "minA", "star"): c.number_format = "#,##0"
s[f"A{WT_START-1}"] = "순위 점수 가중치 (5.이번주_갈곳 순서를 정함)"; s[f"A{WT_START-1}"].font = B
for i, (k, label, v) in enumerate(wt_rows):
    r = WT_START + i
    s.cell(row=r, column=1, value=label); s.cell(row=r, column=2, value=v).fill = IN
    s.cell(row=r, column=3, value="Claude 초안값 — 쓰시면서 고쳐 주십시오").font = GREY
s[f"A{KW_START-1}"] = f"대상 용도 키워드(주용도·현장명에 이 글자가 있으면 대상) — {KWN}칸"; s[f"A{KW_START-1}"].font = B
kw = ["숙박", "호텔", "리조트", "콘도", "기숙사", "수련", "연수", "노유자", "실버", "스테이"]
for i, k in enumerate(kw): s.cell(row=KW_START + i, column=1, value=k).fill = IN
s[f"A{STG_START-1}"] = f"공사단계 키워드 → 누구를 만나나(위에서부터 먼저 맞는 것) — {STN}칸. 안 맞으면 착공일 없음=설계사 / 있음=시공사"; s[f"A{STG_START-1}"].font = B
stg = [("착공전", "설계사"), ("설계", "설계사"), ("인허가", "설계사"), ("허가", "설계사"), ("착공", "시공사"), ("골조", "시공사"), ("마감", "늦음"), ("준공", "늦음")]
for i, (k, v) in enumerate(stg):
    s.cell(row=STG_START + i, column=1, value=k).fill = IN; s.cell(row=STG_START + i, column=2, value=v).fill = IN
widths(s, {"A": 56, "B": 30, "C": 60, "D": 40})
dv_stage = DataValidation(type="list", formula1='"설계사,시공사,늦음"', allow_blank=True); s.add_data_validation(dv_stage)
dv_stage.add(f"B{STG_START}:B{STG_START+STN-1}")

# ================= 2.현장 =================
h = sheet("2.현장")
cols = ["원본행", "주소", "현장명", "연면적(㎡)", "주용도", "허가구분", "설계사", "시공사", "감리", "건축주", "공사단계", "허가일", "착공일",
        "시·도", "유효", "키", "최신", "용도맞음", "면적맞음", "대상", "누구를 만나나", "설계사(정리)", "설계사번호", "허가일원문", "착공일원문", "대상키",
        "판정", "현장점수", "시공사(정리)", "시공사번호"]
hdr(h, 1, cols)
def raw(f):
    fr = FROW[f]
    return f'IF({ST}!$C${fr}="","",IFERROR(INDEX({P1}!$A$1:$AZ${N+1},ROW(),{ST}!$C${fr})&"",""))'
def parse_date(ref):
    t = f'SUBSTITUTE(SUBSTITUTE({ref},".","-"),"/","-")'
    return (f'=IF({ref}="","",IF(AND(LEN({ref})=8,ISNUMBER({ref}*1)),DATE(LEFT({ref},4),MID({ref},5,2),RIGHT({ref},2)),'
            f'IFERROR(DATEVALUE({t}),IFERROR({ref}*1,""))))')
def norm(ref):
    return f'SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE({ref},"(주)",""),"㈜",""),"주식회사","")," ","")'
kwor = ",".join(f'AND({ST}!$A${KW_START+i}<>"",ISNUMBER(SEARCH({ST}!$A${KW_START+i},E{{r}}&" "&C{{r}})))' for i in range(KWN))
def stage(r):
    f = f'IF(M{r}="","설계사","시공사")'
    for i in reversed(range(STN)):
        a = f"{ST}!$A${STG_START+i}"; b = f"{ST}!$B${STG_START+i}"
        f = f'IF(AND({a}<>"",ISNUMBER(SEARCH({a},K{r}))),{b},{f})'
    return f'=IF(T{r},{f},"")'
for r in range(2, R2 + 1):
    drift = f'OR(AND(L{r}<>"",M{r}<>"",N(M{r})-N(L{r})>{NUM["drift"]}),AND(L{r}<>"",M{r}="",{TD}-N(L{r})>{NUM["drift"]}))'
    row = {
     "A": "=ROW()",
     "B": "=" + raw("주소"), "C": "=" + raw("건물명"),
     "D": f'=IFERROR(VALUE(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE({raw("연면적")},",",""),"㎡",""),"m2","")),"")',
     "E": "=" + raw("주용도"), "F": "=" + raw("허가구분"), "G": "=" + raw("설계사"), "H": "=" + raw("시공사"),
     "I": "=" + raw("감리"), "J": "=" + raw("건축주"), "K": "=" + raw("공사단계"),
     "L": parse_date(f"X{r}"), "M": parse_date(f"Y{r}"),
     "N": f'=IFERROR(LEFT(B{r},FIND(" ",B{r})-1),B{r})',
     "O": f'=AND(B{r}<>"",B{r}<>{ST}!$B${FROW["주소"]},C{r}<>{ST}!$B${FROW["건물명"]})',
     "P": f'=IF(O{r},B{r}&"|"&C{r},"")',
     "Q": f'=AND(O{r},COUNTIF(P{r+1}:P${R2+1},P{r})=0)',
     "R": "=OR(" + kwor.format(r=r) + ")",
     "S": f'=OR(D{r}="",D{r}>={NUM["minA"]})',
     "T": f"=AND(Q{r},R{r},S{r})",
     "U": stage(r),
     "V": "=" + norm(f"G{r}"),
     "W": f'=IF(AND(T{r},V{r}<>""),IF(COUNTIFS(V$2:V{r},V{r},T$2:T{r},TRUE)=1,MAX(W$1:W{r-1})+1,""),"")',
     "X": "=" + raw("허가일"), "Y": "=" + raw("착공일"),
     "Z": f'=IF(T{r},V{r}&"|"&U{r},"")',
     "AA": f'=IF(NOT(Q{r}),"",IF(NOT(AND(R{r},S{r})),"✕",IF(OR({drift},AND(D{r}<>"",D{r}<{NUM["star"]})),"△","★")))',
     "AB": (f'=IF(NOT(T{r}),"",IF(U{r}="설계사",{NUM["w_design"]},IF(U{r}="시공사",IF(OR(M{r}="",{TD}-N(M{r})<={NUM["recent"]}*30.4),'
            f'{NUM["w_recent"]},{NUM["w_old"]}),{NUM["w_late"]}))+IF(AA{r}="★",{NUM["w_star"]},0))'),
     "AC": "=" + norm(f"H{r}"),
     "AD": f'=IF(AND(T{r},AC{r}<>""),IF(COUNTIFS(AC$2:AC{r},AC{r},T$2:T{r},TRUE)=1,MAX(AD$1:AD{r-1})+1,""),"")',
    }
    for k, v in row.items(): h[f"{k}{r}"] = v
    h[f"L{r}"].number_format = DATE; h[f"M{r}"].number_format = DATE; h[f"D{r}"].number_format = "#,##0"
h.freeze_panes = "C2"
widths(h, {"A": 6, "B": 30, "C": 28, "D": 10, "E": 12, "F": 6, "G": 22, "H": 16, "I": 8, "J": 10, "K": 8, "L": 11, "M": 11,
           "N": 14, "U": 10, "V": 18, "AA": 6, "AB": 8})
for c in ["O", "P", "Q", "R", "S", "T", "W", "X", "Y", "Z", "AC", "AD"]: h.column_dimensions[c].hidden = True
h.auto_filter.ref = f"A1:AD{R2}"
h.conditional_formatting.add(f"AA2:AA{R2}", FormulaRule(formula=['AA2="★"'], fill=PatternFill("solid", fgColor="FFE699")))
h.conditional_formatting.add(f"AA2:AA{R2}", FormulaRule(formula=['AA2="✕"'], font=Font(color="A6A6A6")))

# ================= 업체_메모 =================
m = sheet(MEMO)
m["A1"] = "업체 담당자 메모 — 설계사·시공사 모두. 이름은 3번·3b번 탭 글자와 똑같이(드롭다운은 설계사 목록)"; m["A1"].font = B
hdr(m, 2, ["업체명(시트 표기)", "담당자", "직함", "전화", "이메일", "주소·지역", "메모", "확신도", "출처"])
for r in range(3, 303):
    for c in range(1, 10): m.cell(row=r, column=c).fill = IN
widths(m, {"A": 26, "B": 10, "C": 8, "D": 15, "E": 22, "F": 24, "G": 44, "H": 8, "I": 40})

# ================= 현장_조사 =================
rs = sheet(RS)
rs["A1"] = "현장 조사노트 — 웹검색 결과(출처 포함). 현장명은 2.현장의 현장명과 똑같이"; rs["A1"].font = B
hdr(rs, 2, ["현장명(산군 표기)", "사업명·브랜드", "발주처(시행)", "시공사", "객실 수", "준공 예정", "최신 상황", "확신도", "출처", "조사일"])
for r in range(3, 203):
    for c in range(1, 11): rs.cell(row=r, column=c).fill = IN
widths(rs, {"A": 30, "B": 24, "C": 18, "D": 16, "E": 8, "F": 10, "G": 44, "H": 8, "I": 44, "J": 11})

# ================= 4.방문기록 =================
v = sheet("4.방문기록")
v["A1"] = "방문·통화 기록 — 다녀오시면 한 줄 (날짜·업체·결과만 넣어도 됩니다). 설계사·시공사 모두"; v["A1"].font = B
hdr(v, 2, ["날짜", "업체(설계사·시공사)", "관련 현장", "만난 사람", "들은 것", "다음 약속일", "다음 할 것", "결과", "키(자동)", "목록에없음(자동)"])
for r in range(3, 3 + VR):
    for c in range(1, 9): v.cell(row=r, column=c).fill = IN
    v[f"A{r}"].number_format = DATE; v[f"F{r}"].number_format = DATE
    v[f"I{r}"] = f'=IF(A{r}="","",B{r}&"|"&(A{r}*1))'
    v[f"J{r}"] = f'=IF(B{r}="",0,IF(COUNTIF({P3}!$B$4:$B${3+DS},B{r})+COUNTIF({P3B}!$B$4:$B${3+CS},B{r})=0,1,0))'
res = ["설계의뢰 받음", "도면 받음", "자료 요청받음", "재방문", "부재", "관심 없음"]
dv_res = DataValidation(type="list", formula1='"' + ",".join(res) + '"', allow_blank=True); v.add_data_validation(dv_res); dv_res.add(f"H3:H{2+VR}")
for ws_, rng in ((v, f"B3:B{2+VR}"), (m, "A3:A302")):
    dv = DataValidation(type="list", formula1=f"={P3}!$B$4:$B${3+DS}", allow_blank=True, showErrorMessage=False)
    ws_.add_data_validation(dv); dv.add(rng)
v.freeze_panes = "A3"
widths(v, {"A": 11, "B": 26, "C": 24, "D": 12, "E": 40, "F": 11, "G": 24, "H": 14})
v.column_dimensions["I"].hidden = True; v.column_dimensions["J"].hidden = True

# 방문 관련 공통 수식 (업체명 셀 nm)
def last_visit(nm): return f'IF(COUNTIF({P4}!$B$3:$B${2+VR},{nm})=0,"",_xlfn.MAXIFS({P4}!$A$3:$A${2+VR},{P4}!$B$3:$B${2+VR},{nm}))'
def next_appt(nm): return f'IF(_xlfn.MAXIFS({P4}!$F$3:$F${2+VR},{P4}!$B$3:$B${2+VR},{nm})=0,"",_xlfn.MAXIFS({P4}!$F$3:$F${2+VR},{P4}!$B$3:$B${2+VR},{nm}))'
def last_res(nm, lv): return f'IF({lv}="","",IFERROR(INDEX({P4}!$H$3:$H${2+VR},MATCH({nm}&"|"&({lv}*1),{P4}!$I$3:$I${2+VR},0))&"",""))'
def status(nm, lv, el, rs_): return f'IF({nm}="","",IF({rs_}="관심 없음","제외",IF({lv}="","안 가봄",IF({el}>={NUM["rev"]},{el}&"일 지남","최근 방문"))))'
def memo(nm, col): return f'IFERROR(VLOOKUP({nm},{MEMO}!$A$3:$I$302,{col},FALSE)&"","")'

# ================= 3.설계사 =================
g = sheet("3.설계사")
g["A1"] = "설계사무소별 모음 (자동) — 산군 대상 현장을 설계사무소 이름으로 묶음"; g["A1"].font = T
g["A2"] = f'="설계사무소 "&COUNTIF(B4:B{3+DS},"?*")&"곳 · 대상 현장 "&COUNTIF({P2}!T2:T{R2},TRUE)&"건 · ★ "&COUNTIFS({P2}!AA2:AA{R2},"★")&"건"'
hdr(g, 3, ["번호", "설계사무소", "대상 현장", "설계단계", "시공단계", "★ 수", "현장점수", "최대 연면적", "최근 허가일", "최근 착공일", "대표 현장", "지역",
           "담당자", "전화", "마지막 방문", "경과일", "다음 약속", "최근 결과", "상태", "우선점수", "순위"])
V2 = f"{P2}!$V$2:$V${R2}"; T2 = f"{P2}!$T$2:$T${R2}"; U2 = f"{P2}!$U$2:$U${R2}"; Z2 = f"{P2}!$Z$2:$Z${R2}"
for r in range(4, 4 + DS):
    k = r - 3; b = f"B{r}"
    g[f"A{r}"] = f'=IF({b}="","",{k})'
    g[b] = f'=IFERROR(INDEX({V2},MATCH({k},{P2}!$W$2:$W${R2},0)),"")'
    g[f"C{r}"] = f'=IF({b}="","",COUNTIFS({V2},{b},{T2},TRUE))'
    g[f"D{r}"] = f'=IF({b}="","",COUNTIFS({V2},{b},{U2},"설계사"))'
    g[f"E{r}"] = f'=IF({b}="","",COUNTIFS({V2},{b},{U2},"시공사"))'
    g[f"F{r}"] = f'=IF({b}="","",COUNTIFS({V2},{b},{T2},TRUE,{P2}!$AA$2:$AA${R2},"★"))'
    g[f"G{r}"] = f'=IF({b}="","",SUMIFS({P2}!$AB$2:$AB${R2},{V2},{b},{T2},TRUE))'
    g[f"H{r}"] = f'=IF({b}="","",_xlfn.MAXIFS({P2}!$D$2:$D${R2},{V2},{b},{T2},TRUE))'
    for col, src in (("I", "L"), ("J", "M")):
        mx = f'_xlfn.MAXIFS({P2}!${src}$2:${src}${R2},{V2},{b},{T2},TRUE)'
        g[f"{col}{r}"] = f'=IF({b}="","",IF({mx}=0,"",{mx}))'
    mi = f'IFERROR(MATCH({b}&"|설계사",{Z2},0),IFERROR(MATCH({b}&"|시공사",{Z2},0),MATCH({b}&"|늦음",{Z2},0)))'
    g[f"K{r}"] = f'=IF({b}="","",IFERROR(INDEX({P2}!$C$2:$C${R2},{mi}),""))'
    g[f"L{r}"] = f'=IF({b}="","",IFERROR(INDEX({P2}!$N$2:$N${R2},{mi}),""))'
    g[f"M{r}"] = f'=IF({b}="","",{memo(b, 2)})'
    g[f"N{r}"] = f'=IF({b}="","",{memo(b, 4)})'
    g[f"O{r}"] = f'=IF({b}="","",{last_visit(b)})'
    g[f"P{r}"] = f'=IF(O{r}="","",{TD}-O{r})'
    g[f"Q{r}"] = f'=IF({b}="","",{next_appt(b)})'
    g[f"R{r}"] = f'={last_res(b, f"O{r}")}'
    g[f"S{r}"] = f'={status(b, f"O{r}", f"P{r}", f"R{r}")}'
    appt = f'AND(Q{r}<>"",Q{r}>N(O{r}),Q{r}<={TD}+7)'
    g[f"T{r}"] = (f'=IF({b}="","",IF(S{r}="제외","",IF(AND({NUM["reg"]}<>"",L{r}<>{NUM["reg"]}),"",'
                  f'IF({appt},{NUM["w_appt"]}+7-(Q{r}-{TD}),IF(AND(G{r}>0,OR(S{r}="안 가봄",RIGHT(S{r},2)="지남")),'
                  f'G{r}+IF(O{r}="",{NUM["w_new"]},MIN(P{r},60)*{NUM["w_day"]})+MIN(N(H{r})/1000*{NUM["w_area"]},20),"")))))')
    g[f"U{r}"] = f'=IF(T{r}="","",COUNTIF(T$4:T${3+DS},">"&T{r})+COUNTIF(T$4:T{r},T{r}))'
    for c in "IJOQ": g[f"{c}{r}"].number_format = DATE
    g[f"H{r}"].number_format = "#,##0"; g[f"T{r}"].number_format = "0.0"; g[f"G{r}"].number_format = "0.#"
g.freeze_panes = "C4"
widths(g, {"A": 5, "B": 26, "C": 7, "D": 7, "E": 7, "F": 6, "G": 7, "H": 10, "I": 11, "J": 11, "K": 26, "L": 14, "M": 10, "N": 14,
           "O": 11, "P": 6, "Q": 11, "R": 12, "S": 10, "T": 8, "U": 6})

# ================= 3b.시공사 =================
c3 = sheet("3b.시공사")
c3["A1"] = "시공사(또는 발주처)별 모음 (자동) — 착공 현장 영업용. 산군 캡처의 「열J」 가 시공사인지 발주처인지는 미확인"; c3["A1"].font = T
c3["A2"] = f'="시공사 "&COUNTIF(B4:B{3+CS},"?*")&"곳"'
hdr(c3, 3, ["번호", "시공사", "대상 현장", "★ 수", "최대 연면적", "최근 착공일", "대표 현장", "지역", "설계사(같은 현장)", "담당자", "전화", "마지막 방문", "경과일", "최근 결과", "상태"])
AC2 = f"{P2}!$AC$2:$AC${R2}"
for r in range(4, 4 + CS):
    k = r - 3; b = f"B{r}"
    c3[f"A{r}"] = f'=IF({b}="","",{k})'
    c3[b] = f'=IFERROR(INDEX({AC2},MATCH({k},{P2}!$AD$2:$AD${R2},0)),"")'
    c3[f"C{r}"] = f'=IF({b}="","",COUNTIFS({AC2},{b},{T2},TRUE))'
    c3[f"D{r}"] = f'=IF({b}="","",COUNTIFS({AC2},{b},{T2},TRUE,{P2}!$AA$2:$AA${R2},"★"))'
    c3[f"E{r}"] = f'=IF({b}="","",_xlfn.MAXIFS({P2}!$D$2:$D${R2},{AC2},{b},{T2},TRUE))'
    mx = f'_xlfn.MAXIFS({P2}!$M$2:$M${R2},{AC2},{b},{T2},TRUE)'
    c3[f"F{r}"] = f'=IF({b}="","",IF({mx}=0,"",{mx}))'
    mi = f'MATCH({k},{P2}!$AD$2:$AD${R2},0)'
    c3[f"G{r}"] = f'=IF({b}="","",IFERROR(INDEX({P2}!$C$2:$C${R2},{mi}),""))'
    c3[f"H{r}"] = f'=IF({b}="","",IFERROR(INDEX({P2}!$N$2:$N${R2},{mi}),""))'
    c3[f"I{r}"] = f'=IF({b}="","",IFERROR(INDEX({P2}!$G$2:$G${R2},{mi}),""))'
    c3[f"J{r}"] = f'=IF({b}="","",{memo(b, 2)})'
    c3[f"K{r}"] = f'=IF({b}="","",{memo(b, 4)})'
    c3[f"L{r}"] = f'=IF({b}="","",{last_visit(b)})'
    c3[f"M{r}"] = f'=IF(L{r}="","",{TD}-L{r})'
    c3[f"N{r}"] = f'={last_res(b, f"L{r}")}'
    c3[f"O{r}"] = f'={status(b, f"L{r}", f"M{r}", f"N{r}")}'
    for c in "FL": c3[f"{c}{r}"].number_format = DATE
    c3[f"E{r}"].number_format = "#,##0"
c3.freeze_panes = "C4"
widths(c3, {"A": 5, "B": 22, "C": 7, "D": 6, "E": 10, "F": 11, "G": 28, "H": 14, "I": 22, "J": 10, "K": 14, "L": 11, "M": 6, "N": 12, "O": 10})

# ================= 5.이번주_갈곳 =================
w = sheet("5.이번주_갈곳")
w["A1"] = (f'="이번 주 찾아갈 설계사무소 "&COUNT({P3}!U4:U{3+DS})&"곳 · 지역: "&IF({NUM["reg"]}="","전국",{NUM["reg"]})'
           f'&" · 기준일 "&TEXT({TD},"yyyy-mm-dd")')
w["A1"].font = T
w["A2"] = "순서 : ①7일 안 약속 ②현장점수(설계단계>최근 착공>오래된 착공, ★ 가산) + 안 가본 곳 ③오래 안 간 곳. 관심 없음·최근 방문은 빠짐. 가중치·지역은 설정 탭."
hdr(w, 3, ["순위", "설계사무소", "왜 가야 하나", "대표 현장", "지역", "최대 연면적", "객실 수(조사)", "담당자", "전화", "전화 첫마디(복사)", "_행"])
for r in range(4, 34):
    k = r - 3
    w[f"K{r}"] = f'=IFERROR(MATCH({k},{P3}!$U$4:$U${3+DS},0),"")'
    def G(col): return f'INDEX({P3}!${col}$4:${col}${3+DS},K{r})'
    w[f"A{r}"] = f'=IF(K{r}="","",{k})'
    w[f"B{r}"] = f'=IF(K{r}="","",{G("B")})'
    w[f"C{r}"] = (f'=IF(K{r}="","",IF({G("T")}>={NUM["w_appt"]},"약속 "&TEXT({G("Q")},"m/d")&" · ",IF({G("O")}="","아직 안 가봄 · ","마지막 방문 "&{G("P")}&"일 전 · "))'
                  f'&IF({G("D")}>0,"설계단계 "&{G("D")}&"건 ","")&IF({G("E")}>0,"착공 "&{G("E")}&"건 ","")&IF({G("F")}>0,"★"&{G("F")},""))')
    w[f"D{r}"] = f'=IF(K{r}="","",{G("K")})'
    w[f"E{r}"] = f'=IF(K{r}="","",{G("L")})'
    w[f"F{r}"] = f'=IF(K{r}="","",{G("H")})'
    rv = lambda col: f'VLOOKUP(D{r},{RS}!$A$3:$J$202,{col},FALSE)'
    w[f"G{r}"] = f'=IF(K{r}="","",IFERROR(IF({rv(5)}&""="","",{rv(5)}&IF({rv(8)}="확정",""," ("&{rv(8)}&")")),""))'
    w[f"H{r}"] = f'=IF(K{r}="","",{G("M")})'
    w[f"I{r}"] = f'=IF(K{r}="","",{G("N")})'
    w[f"J{r}"] = (f'=IF(K{r}="","","한국마이크로닉 배성윤 차장입니다. "&D{r}&IF({G("D")}>0," 설계 진행하고 계신"," 설계하신")&" 걸로 알고 연락드렸습니다. '
                  f'호텔 객실관리(키센서·온도조절기·조명 제어) 도면과 특기시방서를 무상으로 지원해 드리고 있어서, 이번 주에 10분만 찾아뵈어도 될까요?")')
    w[f"F{r}"].number_format = "#,##0"; w[f"J{r}"].alignment = WRAP; w[f"C{r}"].alignment = WRAP
w.column_dimensions["K"].hidden = True
w.freeze_panes = "C4"
widths(w, {"A": 5, "B": 26, "C": 32, "D": 28, "E": 14, "F": 10, "G": 9, "H": 10, "I": 14, "J": 70})

# ================= 6.본전계산 =================
b6 = sheet("6.본전계산")
b6["A1"] = "산군 본전 계산 — 기준 숫자는 프로님이 넣으십시오(노란 칸)"; b6["A1"].font = T
hdr(b6, 3, ["항목", "값", "메모"])
rows6 = [
 ("산군 사용 시작일", f'=IF({NUM["start"]}="","",{NUM["start"]})', "설정 탭"),
 ("쓴 개월 수", f'=IF({NUM["start"]}="","시작일을 넣으십시오",DATEDIF({NUM["start"]},{TD},"m")+1)', ""),
 ("누적 비용(원)", f'=IF(ISNUMBER(B5),B5*{NUM["fee"]},"")', "월 요금 × 개월"),
 ("월결제 1년 vs 연결제 차액(원)", f'={NUM["fee"]}*12-{NUM["year"]}', "연결제로 바꾸면 아끼는 돈"),
 ("산군으로 잡힌 설계사무소(곳)", f'=COUNTIF({P3}!B4:B{3+DS},"?*")', "자동"),
 ("그중 한 번이라도 간 곳", f'=COUNTIF({P3}!O4:O{3+DS},">0")', "자동"),
 ("방문·통화 건수", f"=COUNT({P4}!A3:A{2+VR})", "자동"),
 ("설계의뢰 받음", f'=COUNTIF({P4}!H3:H{2+VR},"설계의뢰 받음")', "자동"),
 ("도면 받음", f'=COUNTIF({P4}!H3:H{2+VR},"도면 받음")', "자동"),
 ("방문 1건당 산군 비용(원)", f'=IF(AND(ISNUMBER(B6),B10>0),B6/B10,"")', ""),
 ("★ 현장 수(지금 붙여넣은 자료)", f'=COUNTIF({P2}!AA2:AA{R2},"★")', "자동"),
]
for i, (a, f, note) in enumerate(rows6):
    r = 4 + i; b6[f"A{r}"] = a; b6[f"B{r}"] = f; b6[f"B{r}"].fill = AUTO; b6[f"C{r}"] = note
b6["B4"].number_format = DATE
for r in (6, 7, 13): b6[f"B{r}"].number_format = "#,##0"
b6["A17"] = "유지 판단 기준(3개월 기준 권장)"; b6["A17"].font = B
hdr(b6, 18, ["기준", "넣으실 숫자", "실제", "통과"])
crit = [("간 설계사무소 최소(곳)", "B9"), ("설계의뢰 받음 최소(건)", "B11"), ("도면 받음 최소(건)", "B12")]
for i, (a, ref) in enumerate(crit):
    r = 19 + i; b6[f"A{r}"] = a; b6[f"B{r}"].fill = IN; b6[f"C{r}"] = f"={ref}"
    b6[f"D{r}"] = f'=IF(B{r}="","",IF(C{r}>=B{r},"통과","미달"))'
b6["A23"] = "판정"; b6["A23"].font = B
b6["B23"] = '=IF(COUNTA(B19:B21)=0,"기준 숫자를 넣으시면 판정합니다",IF(COUNTIF(D19:D21,"미달")=0,"유지","미달 있음 — 해지 또는 조건 변경 검토"))'
widths(b6, {"A": 34, "B": 30, "C": 26, "D": 10})

# ================= 7.문안 =================
t = sheet("7.문안")
t["A1"] = "보낼 문안 — B3 에서 업체를 고르면 채워집니다(비우면 5번 1순위). [ ] 칸만 고쳐서 복사"; t["A1"].font = T
t["A3"] = "업체 고르기"; t["B3"].fill = IN
dv_t = DataValidation(type="list", formula1=f"={P3}!$B$4:$B${3+DS}", allow_blank=True, showErrorMessage=False); t.add_data_validation(dv_t); dv_t.add("B3")
t["A5"] = "쓰는 이름"; t["B5"] = "=IF(B3<>\"\",B3,IFERROR('5.이번주_갈곳'!B4&\"\",\"\"))"
rowi = f"MATCH(B5,{P3}!$B$4:$B${3+DS},0)"
t["A6"] = "대표 현장"; t["B6"] = f'=IFERROR(INDEX({P3}!$K$4:$K${3+DS},{rowi})&"",IFERROR(INDEX({P3B}!$G$4:$G${3+CS},MATCH(B5,{P3B}!$B$4:$B${3+CS},0))&"","[현장명]"))'
t["A7"] = "담당자"; t["B7"] = f'=IF({memo("B5", 2)}="","[담당자]",{memo("B5", 2)})'
t["A8"] = "객실 수"; t["B8"] = f'=IFERROR(IF(AND(VLOOKUP(B6,{RS}!$A$3:$J$202,5,FALSE)&""<>"",VLOOKUP(B6,{RS}!$A$3:$J$202,8,FALSE)="확정"),VLOOKUP(B6,{RS}!$A$3:$J$202,5,FALSE)&"","[객실 수]"),"[객실 수]")'
t["C8"] = "현장_조사 확신도가 「확정」 일 때만 자동으로 넣음(추정 숫자가 대외 메일로 나가지 않게)"; t["C8"].font = GREY
t["A9"] = "전화(업체)"; t["B9"] = f'=IF({memo("B5", 4)}="","[업체 전화]",{memo("B5", 4)})'
for r in range(5, 10): t[f"B{r}"].fill = AUTO
NL = "CHAR(10)"
t["A11"] = "전화 첫마디"
t["B11"] = '="한국마이크로닉 배성윤 차장입니다. "&B6&" 설계하신 걸로 알고 연락드렸습니다. 호텔 객실관리(키센서·온도조절기·조명 제어) 도면과 특기시방서를 무상으로 지원해 드리고 있어서, [날짜]에 10분만 찾아뵈어도 될까요?"'
t["A12"] = "문자(짧게)"
t["B12"] = '="[한국마이크로닉 배성윤 차장] "&B7&"님, "&B6&" 객실관리시스템 설계(도면·특기시방서)를 무상 지원해 드립니다. [날짜] 잠시 찾아뵈어도 될지 여쭙습니다. [내 전화]"'
t["A13"] = "메일 제목"
t["B13"] = '=B6&" 객실관리시스템 설계 지원 안내 — 한국마이크로닉(주)"'
t["A14"] = "메일 본문"
t["B14"] = ('=B5&" "&B7&"님께"&' + NL + '&' + NL +
            '&"한국마이크로닉(주) 배성윤 차장입니다."&' + NL +
            '&"진행 중이신 "&B6&"("&B8&"실)의 객실관리시스템(RCU·CB함·키센서·온도조절기·조명스위치) 설계를 무상으로 지원해 드립니다."&' + NL + '&' + NL +
            '&"- 단위세대 평면도를 주시면 성급 기준으로 기구물 배치와 예상 수량표를 드립니다."&' + NL +
            '&"- 특기시방서(한글 파일)를 현장에 맞춰 드립니다."&' + NL +
            '&"- 전기설계업체용 설계 요청서도 함께 드립니다."&' + NL + '&' + NL +
            '&"[날짜] 중 편하신 시간에 찾아뵙겠습니다."&' + NL + '&' + NL +
            '&"한국마이크로닉(주) 배성윤 드림 / [내 전화]"')
t["A15"] = "방문 후 감사 문자"
t["B15"] = '="[한국마이크로닉 배성윤] 오늘 시간 내주셔서 감사합니다. 말씀하신 "&B6&" 자료는 [날짜]까지 보내 드리겠습니다."'
for r in range(11, 16): t[f"B{r}"].alignment = WRAP; t[f"B{r}"].fill = AUTO; t[f"A{r}"].font = B
t.row_dimensions[14].height = 230; t.row_dimensions[11].height = 60; t.row_dimensions[12].height = 45
widths(t, {"A": 18, "B": 100})

# ================= 0.자가진단 =================
d["A1"] = f"자가진단 — 전부 「통과」 면 정상 ({VER})"; d["A1"].font = T
hdr(d, 3, ["번호", "검사", "값", "판정"])
req_cnt = "+".join(f'N({ST}!$C${FROW[f]}<>"")' for f in REQUIRED)
wt_blank = f'COUNTBLANK({ST}!B{WT_START}:B{WT_START+len(wt_rows)-1})'
chk = [
 ("필수 열 6개(주소·현장명·연면적·용도·설계사·단계) 중 찾은 수", f"={req_cnt}", 'IF(C{r}=6,"통과","실패 — 설정 B4~B15 를 산군 엑셀 1행과 맞추기")'),
 ("붙여넣은 줄 수(1행 제외)", f"=MAX(COUNTA({P1}!B:B)-1,0)", f'IF(C{{r}}<={N},"통과","실패 — {N}줄 넘음, 아래는 안 읽힘")'),
 ("읽은 현장 줄", f"=COUNTIF({P2}!O2:O{R2},TRUE)", 'IF(C{r}=C5,"통과","확인 — 빈 줄·제목줄 섞임")'),
 ("중복 뺀 현장", f"=COUNTIF({P2}!Q2:Q{R2},TRUE)", '"참고"'),
 ("★ + △ + ✕ = 중복 뺀 현장", f'=COUNTIF({P2}!AA2:AA{R2},"★")+COUNTIF({P2}!AA2:AA{R2},"△")+COUNTIF({P2}!AA2:AA{R2},"✕")', 'IF(C{r}=C7,"통과","실패")'),
 ("대상(용도·면적 맞음) = ★ + △", f"=COUNTIF({P2}!T2:T{R2},TRUE)", f'IF(C{{r}}=COUNTIF({P2}!AA2:AA{R2},"★")+COUNTIF({P2}!AA2:AA{R2},"△"),"통과","실패")'),
 ("설계사행 + 시공사행 + 늦음 = 대상", f'=COUNTIF({P2}!U2:U{R2},"설계사")+COUNTIF({P2}!U2:U{R2},"시공사")+COUNTIF({P2}!U2:U{R2},"늦음")', 'IF(C{r}=C9,"통과","실패 — 설정 단계 행선지 확인")'),
 ("설계사무소 수 = 번호 최댓값", f'=COUNTIF({P3}!B4:B{3+DS},"?*")', f'IF(C{{r}}=MAX({P2}!W2:W{R2}),IF(C{{r}}<{DS},"통과","실패 — {DS}곳 넘음"),"실패")'),
 ("시공사 수 = 번호 최댓값", f'=COUNTIF({P3B}!B4:B{3+CS},"?*")', f'IF(C{{r}}=MAX({P2}!AD2:AD{R2}),IF(C{{r}}<{CS},"통과","실패 — {CS}곳 넘음"),"실패")'),
 ("설계사 이름 빈 대상 현장", f'=COUNTIFS({P2}!T2:T{R2},TRUE,{P2}!V2:V{R2},"")', 'IF(C{r}=0,"통과","확인 — 산군에 설계사 없는 현장(2.현장 필터)")'),
 ("방문기록 중 목록에 없는 업체", f"=SUM({P4}!J3:J{2+VR})", 'IF(C{r}=0,"통과","확인 — 이름이 3·3b번 탭과 다름")'),
 ("점수 가중치 빈칸", f"={wt_blank}", 'IF(C{r}=0,"통과","실패 — 설정 가중치 칸을 채우기")'),
 ("잘린 이름(…) 남음", f'=COUNTIF({P2}!G2:G{R2},"*…*")', 'IF(C{r}=0,"통과","알림 — 캡처에서 잘린 설계사 이름. 엑셀 다운로드로 바꾸면 사라짐")'),
]
for i, (name, val, judge) in enumerate(chk):
    r = 4 + i
    d[f"A{r}"] = i + 1; d[f"B{r}"] = name; d[f"C{r}"] = val; d[f"D{r}"] = "=" + judge.format(r=r)
LAST = 4 + len(chk) - 1
d.conditional_formatting.add(f"D4:D{LAST}", FormulaRule(formula=['LEFT(D4,2)="통과"'], fill=PatternFill("solid", fgColor="C6EFCE")))
d.conditional_formatting.add(f"D4:D{LAST}", FormulaRule(formula=['LEFT(D4,2)="실패"'], fill=PatternFill("solid", fgColor="FFC7CE")))
HR = LAST + 3
d[f"A{HR}"] = "변경 이력"; d[f"A{HR}"].font = B
hdr(d, HR + 1, ["날짜", "판", "내용", "이전 값"])
hist = [
 (TODAY, "v1", "최초 작성 — 11개 탭. 계획서 v1 구조 그대로. [예시] 가짜 7줄.", "없음"),
 (TODAY, "v2", "B: 1번 탭 예시 → 9/14 산군 캡처 실제 11줄 / 설정 열 이름 → 캡처 열 이름", "예시 열 이름(현장주소·건물명·주용도·설계사…)"),
 (TODAY, "v2", "B: 순위 점수 고정값 → 설정 탭 가중치 입력칸 9개 / 착공 현장도 설계사 점수(최근 착공 6, 오래된 착공 2)", "설계단계 수×10 + 안가봄 30 + 경과/2 + 면적/1000 (착공 현장은 0점)"),
 (TODAY, "v2", "A: ★△✕ 판정 · 3b.시공사 · 현장_조사 · 업체_메모(설계사_메모 확장, 확신도·출처 칸) · 자가진단 13항목", "—"),
 (TODAY, "v2", "B: 설계사_메모 → 업체_메모(시공사도 같이) / 방문기록 「설계사무소」 → 「업체」", "설계사_메모"),
]
for i, row in enumerate(hist):
    for c, vv in enumerate(row): d.cell(row=HR + 2 + i, column=c + 1, value=vv)
widths(d, {"A": 12, "B": 44, "C": 70, "D": 50})

# ================= 조사 결과 씨앗 =================
if SEED_JSON:
    J = json.load(open(SEED_JSON))
    for i, row in enumerate(J.get("memo", [])):
        for c, vv in enumerate(row): m.cell(row=3 + i, column=c + 1, value=vv)
    for i, row in enumerate(J.get("research", [])):
        for c, vv in enumerate(row): rs.cell(row=3 + i, column=c + 1, value=vv)

order = ["사용법", "0.자가진단", "5.이번주_갈곳", "1.산군_붙여넣기", "3.설계사", "3b.시공사", "4.방문기록", "7.문안", "업체_메모", "현장_조사",
         "2.현장", "6.본전계산", "설정"]
wb._sheets = [wb[n] for n in order]
wb.active = 2
wb.calculation = CalcProperties(fullCalcOnLoad=True)
wb.save(OUT); print("saved", OUT)
