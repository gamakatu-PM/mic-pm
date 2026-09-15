# -*- coding: utf-8 -*-
"""13. 증감(Rev) 비교 - 수량표 2개를 빼서 늘어난 것/줄어든 것/새로 생긴 것/빠진 것을 낸다.
입력 : CSV 또는 XLSX 2개. 품명열과 수량열만 있으면 됩니다.
금액은 넣지 않습니다 (요율/단가는 프로님이 정하십니다)."""
import os, sys, csv, io
from common import *

def read_rows(path):
    ext = os.path.splitext(path)[1].lower()
    if ext in ('.xlsx', '.xlsm'):
        try:
            import openpyxl
        except ImportError:
            print('openpyxl 이 없어 엑셀을 못 엽니다. CSV 로 주십시오.'); return []
        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb[wb.sheetnames[0]]
        return [[c if c is not None else '' for c in r] for r in ws.iter_rows(values_only=True)]
    for enc in ('cp949', 'utf-8-sig', 'utf-8'):
        try:
            with io.open(path, 'r', encoding=enc, newline='') as fp:
                return list(csv.reader(fp))
        except Exception:
            continue
    return []

def to_map(rows):
    """품명->수량. 숫자가 들어있는 마지막 칸을 수량으로 본다."""
    m = {}
    for r in rows:
        if not r:
            continue
        name = str(r[0]).strip()
        if not name or name.startswith('['):
            continue
        qty = None
        for c in r[1:]:
            s = str(c).replace(',', '').strip()
            try:
                qty = float(s)
            except Exception:
                continue
        if qty is not None:
            m[name] = m.get(name, 0) + qty
    return m

def _n(v):
    if v is None or v == '':
        return ''
    return int(v) if float(v) == int(v) else round(float(v), 2)

def diff(old, new):
    keys = sorted(set(old) | set(new))
    rows = []
    for k in keys:
        a, b = old.get(k), new.get(k)
        if a is None:
            rows.append([k, '', _n(b), _n(b), '신규(증)'])
        elif b is None:
            rows.append([k, _n(a), '', _n(-a), '삭제(감)'])
        elif a != b:
            rows.append([k, _n(a), _n(b), _n(b - a), '증' if b > a else '감'])
    return rows

def run(f1=None, f2=None):
    title('13. 증감(Rev) 비교')
    f1 = f1 or ask('이전 수량표 경로 > ')
    f2 = f2 or ask('새 수량표 경로 > ')
    if not (os.path.isfile(f1) and os.path.isfile(f2)):
        print('두 파일 경로를 확인해 주십시오.'); return
    rows = diff(to_map(read_rows(f1)), to_map(read_rows(f2)))
    if not rows:
        print('차이가 없습니다.'); return
    print('%-46s %8s %8s %8s %s' % ('품목', '이전', '이후', '증감', '구분'))
    print('-' * 86)
    for r in rows:
        print('%-46s %8s %8s %8s %s' % (r[0][:46], r[1], r[2], r[3], r[4]))
    o = os.path.join(outdir('증감'), '증감_%s.csv' % ymd6())
    write_csv(o, rows, ['품목', '이전', '이후', '증감', '구분'])
    print('')
    print('파일 : %s' % o)
    print('* 금액은 넣지 않았습니다. 단가/요율은 단가장에서 프로님이 넣으십니다.')

if __name__ == '__main__':
    a = sys.argv[1:]
    run(a[0] if a else None, a[1] if len(a) > 1 else None); pause()
