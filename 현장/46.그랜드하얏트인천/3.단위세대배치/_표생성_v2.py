import importlib.util, os
spec=importlib.util.spec_from_file_location('b','build_all.py')
src=open('build_all.py',encoding='utf-8').read().split('# ---- 이미지 ----')[0]
ns={}; exec(src,ns)
T=ns['T']; RULES=ns['RULES']; OUT=ns['OUT']; DATE=ns['DATE']
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
F=lambda **k: Font(**{'name':'맑은 고딕','size':10,**k})
thin=Side(style='thin',color='BBBBBB'); bd=Border(left=thin,right=thin,top=thin,bottom=thin)
Y=PatternFill('solid',fgColor='FFF2CC'); HB=PatternFill('solid',fgColor='DDEBF7'); GR=PatternFill('solid',fgColor='E2EFDA'); OR=PatternFill('solid',fgColor='FCE4D6')
ITEMS=['챠임벨','키센서 DM','BSP','거실 스위치','회의실 스위치','드레스룸 스위치','화장실 스위치','비상호출']
wb=Workbook()
# --- 1. 한눈표
m=wb.active; m.title='1.한눈표(타입×품목)'
m['A1']='그랜드 하얏트 인천 단위세대 기구물 — 1실당 개수 (2.상세 시트를 고치면 자동 반영)'; m['A1'].font=F(bold=True,size=12)
hdr=['No','타입','룸']+ITEMS+['1실 합계','실 수(입력)','총 개수','상태']
m.append([]); 
for i,h in enumerate(hdr,1):
    c=m.cell(row=2,column=i,value=h); c.font=F(bold=True); c.fill=HB; c.border=bd; c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
r=3; keys=[]
for ti,(src_,title,status,rooms) in enumerate(T,1):
    for rn,_,items in rooms:
        keys.append((title,rn or '-'))
        m.cell(row=r,column=1,value=ti); m.cell(row=r,column=2,value=title); m.cell(row=r,column=3,value=rn or '-')
        for j,it in enumerate(ITEMS):
            col=4+j
            m.cell(row=r,column=col,value=f"=SUMPRODUCT(('2.상세(수정)'!$A$2:$A$300=$B{r})*('2.상세(수정)'!$B$2:$B$300=$C{r})*('2.상세(수정)'!$D$2:$D$300={L(col)}$2)*'2.상세(수정)'!$E$2:$E$300)")
        sc=4+len(ITEMS)
        m.cell(row=r,column=sc,value=f'=SUM(D{r}:{L(sc-1)}{r})')
        m.cell(row=r,column=sc+1).fill=Y
        m.cell(row=r,column=sc+2,value=f'={L(sc)}{r}*{L(sc+1)}{r}')
        m.cell(row=r,column=sc+3,value=status)
        for c in m[r]:
            c.font=F(); c.border=bd; c.alignment=Alignment(vertical='center',wrap_text=True)
            if isinstance(c.value,str) and c.value.startswith('='): c.number_format='#,##0;-#,##0;""'
        m.cell(row=r,column=sc+3).fill=GR if status.startswith('확정') else OR
        r+=1
last=r-1
m.cell(row=r,column=2,value='합계 (실 수 넣으면 총 개수 자동)').font=F(bold=True)
for col in range(4,4+len(ITEMS)+3):
    if col==4+len(ITEMS)+1: continue
    c=m.cell(row=r,column=col,value=f'=SUMPRODUCT({L(col)}3:{L(col)}{last},$M$3:$M{last})' if col<4+len(ITEMS) else f'=SUM({L(col)}3:{L(col)}{last})')
    c.font=F(bold=True); c.number_format='#,##0'; c.border=bd
m.cell(row=r+1,column=2,value='※ 합계행 품목칸 = 품목별 총 개수(1실 개수 × 실 수). 중앙장비(PC 1, FIP 층당 1)는 층별 실 수 확정 후 추가.').font=F(color='808080')
for i,w in enumerate([4,30,20]+[9]*len(ITEMS)+[9,10,9,16],1): m.column_dimensions[L(i)].width=w
m.row_dimensions[2].height=32; m.freeze_panes='D3'
# --- 2. 상세
d=wb.create_sheet('2.상세(수정)')
d.append(['타입','룸','번호','품목','수량','위치','상태','비고(프로님 수정칸)'])
for c in d[1]: c.font=F(bold=True); c.fill=HB; c.border=bd
for ti,(src_,title,status,rooms) in enumerate(T,1):
    d.append([f'■ {ti:02d}. {title}   [{status}]']); hr=d.max_row
    d.cell(row=hr,column=1).font=F(bold=True); 
    for c in range(1,9): d.cell(row=hr,column=c).fill=GR if status.startswith('확정') else OR
    for rn,_,items in rooms:
        agg={}
        for n,name,_,_,loc in items: agg.setdefault((n,name),[]).append(loc)
        for (n,name),locs in agg.items():
            d.append([title,rn or '-',int(n),name,len(locs),' / '.join(locs),status,''])
            rr=d.max_row
            for c in d[rr]: c.font=F(); c.border=bd; c.alignment=Alignment(vertical='center',wrap_text=True)
            d.cell(row=rr,column=5).fill=Y; d.cell(row=rr,column=8).fill=Y; d.cell(row=rr,column=5).number_format='#,##0'
for i,w in enumerate([26,18,5,15,6,52,18,24],1): d.column_dimensions[L(i)].width=w
d.freeze_panes='A2'
# --- 3. 기준
g=wb.create_sheet('3.정정 기준(프로님 확정)')
g.append(['항목','기준 (2026-10-07)'])
for k,v in RULES: g.append([k,v])
for row in g.iter_rows():
    for c in row: c.font=F(); c.border=bd; c.alignment=Alignment(wrap_text=True,vertical='top')
for c in g[1]: c.font=F(bold=True); c.fill=HB
g.column_dimensions['A'].width=14; g.column_dimensions['B'].width=100
# --- 0. 자가진단 (맨 뒤)
z=wb.create_sheet('9.자가진단·변경이력')
z.append(['날짜','버전','내용'])
for l in [('2026-10-07','v1','12타입 번호판·표 최초 생성. 파일을 열면 자가진단 시트가 먼저 떠서 타입이 안 보였음.'),
          ('2026-10-07','v2','한눈표(타입×품목)를 첫 장으로. 이름에 ~ 가 든 타입(4~12F, K-3, DD-2)이 SUMIFS에서 0으로 잡히던 문제를 SUMPRODUCT로 바꿔 해결. 상세 시트에 타입별 구분 줄 추가. 한눈표는 상세를 자동 집계. 수량 값은 v1과 같음.'),
          ('2026-10-07','-','④ 미검토 8타입 위치는 Claude 판독. OS 회의·식당실, PS 다이닝·라이브러리 거실 스위치, KING&DD BSP 자리는 확신 낮음.'),
          ('2026-10-07','-','저장: 드라이브 _컴퓨터로 옮길 것/46.그랜드하얏트인천_3.단위세대배치 → PC 「3.단위세 대배치 해봄」으로 옮길 것.')]:
    z.append(list(l))
for row in z.iter_rows():
    for c in row: c.font=F(); c.border=bd; c.alignment=Alignment(wrap_text=True,vertical='top')
z.column_dimensions['C'].width=100
for ws in wb.worksheets:
    ws.page_setup.orientation='landscape'; ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0; ws.sheet_properties.pageSetUpPr.fitToPage=True
wb.active=0
p=os.path.join(OUT,f'그랜드하얏트인천_단위세대_기구물수량_v2_{DATE}.xlsx'); wb.save(p); print(p)
